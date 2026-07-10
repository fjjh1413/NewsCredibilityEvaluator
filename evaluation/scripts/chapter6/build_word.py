from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


HEADING_SIZES = {1: 16, 2: 14, 3: 12}
TABLE_MAX_DATA_ROWS = 8


def _plain_markdown(text: str) -> str:
    return text.replace("**", "").replace("`", "").strip()


def _set_run_font(run: Any, *, size: float, bold: bool = False, chinese: str = "宋体") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")


def _format_paragraph(paragraph: Any, *, kind: str = "body", level: int = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.5
    if kind == "body":
        fmt.first_line_indent = Pt(21)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            _set_run_font(run, size=10.5)
    elif kind == "heading":
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(6 if level > 1 else 10)
        fmt.space_after = Pt(3)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            _set_run_font(run, size=HEADING_SIZES[level], bold=True, chinese="黑体")
    elif kind == "caption":
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(3)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            _set_run_font(run, size=10.5)
    elif kind == "image":
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_cell_margins(cell: Any, top: int = 50, start: int = 70, bottom: int = 50, end: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def _format_table(table: Any) -> None:
    try:
        table.style = "Table Grid"
    except KeyError:
        table.style = "Normal Table"
    _set_table_borders(table)
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        _prevent_row_split(row)
        if row_index == 0:
            _repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    _set_run_font(run, size=8.5, bold=row_index == 0)


def _parse_markdown_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    table_lines: list[str] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1
    parsed = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
    if len(parsed) < 2:
        return parsed[0] if parsed else [], [], index
    return parsed[0], parsed[2:], index


def _add_table_blocks(
    document: DocumentObject,
    blocks: list[Any],
    headers: list[str],
    rows: list[list[str]],
    caption: str | None,
) -> None:
    chunks = [rows[index:index + TABLE_MAX_DATA_ROWS] for index in range(0, len(rows), TABLE_MAX_DATA_ROWS)] or [[]]
    for chunk_index, chunk in enumerate(chunks):
        if chunk_index > 0 and caption:
            continuation = document.add_paragraph(f"{caption}（续）", style="Normal")
            _format_paragraph(continuation, kind="caption")
            blocks.append(continuation._p)
        table = document.add_table(rows=1, cols=len(headers))
        for column, value in enumerate(headers):
            table.rows[0].cells[column].text = _plain_markdown(value)
        for source_row in chunk:
            cells = table.add_row().cells
            padded = source_row + [""] * (len(headers) - len(source_row))
            for column, value in enumerate(padded[:len(headers)]):
                cells[column].text = _plain_markdown(value)
        _format_table(table)
        blocks.append(table._tbl)


def _find_reference_paragraph(document: DocumentObject) -> Any:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().replace(" ", "") == "参考文献":
            return paragraph
    raise ValueError("未在论文中找到“参考文献”标题")


def _find_chapter_paragraph(document: DocumentObject) -> Any | None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().replace(" ", "") == "第6章系统测试与实验分析":
            return paragraph
    return None


def _remove_existing_chapter(document: DocumentObject, reference: Any) -> None:
    chapter = _find_chapter_paragraph(document)
    if chapter is None:
        return
    node = chapter._p
    while node is not None and node is not reference._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node


def _table_count_before_paragraph(document: DocumentObject, paragraph: Any) -> int:
    count = 0
    for child in document._element.body.iterchildren():
        if child is paragraph._p:
            break
        if child.tag == qn("w:tbl"):
            count += 1
    return count


def _field_code_count_before_paragraph(document: DocumentObject, paragraph: Any) -> int:
    count = 0
    for child in document._element.body.iterchildren():
        if child is paragraph._p:
            break
        count += child.xml.count("<w:instrText")
    return count


def _enable_field_updates(document: DocumentObject) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def build_document(source: str | Path, markdown: str | Path, output: str | Path) -> Path:
    source_path = Path(source)
    markdown_path = Path(markdown)
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)
    document = Document(output_path)
    reference = _find_reference_paragraph(document)
    _remove_existing_chapter(document, reference)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    blocks: list[Any] = []
    index = 0
    pending_table_caption: str | None = None
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", raw)
        image = re.match(r"^!\[(.*?)\]\((.+)\)$", raw)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(_plain_markdown(heading.group(2)), style=f"Heading {level}")
            _format_paragraph(paragraph, kind="heading", level=level)
            blocks.append(paragraph._p)
            pending_table_caption = None
            index += 1
            continue
        if image:
            image_path = Path(image.group(2))
            paragraph = document.add_paragraph()
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(5.9))
            _format_paragraph(paragraph, kind="image")
            blocks.append(paragraph._p)
            index += 1
            continue
        if raw.startswith("|"):
            headers, rows, index = _parse_markdown_table(lines, index)
            _add_table_blocks(document, blocks, headers, rows, pending_table_caption)
            pending_table_caption = None
            continue
        if re.match(r"^表6-\d+", raw):
            pending_table_caption = _plain_markdown(raw)
            paragraph = document.add_paragraph(pending_table_caption, style="Normal")
            _format_paragraph(paragraph, kind="caption")
            blocks.append(paragraph._p)
            index += 1
            continue
        if re.match(r"^图6-\d+", raw):
            paragraph = document.add_paragraph(_plain_markdown(raw), style="Normal")
            _format_paragraph(paragraph, kind="caption")
            blocks.append(paragraph._p)
            index += 1
            continue

        paragraph = document.add_paragraph(_plain_markdown(raw), style="Normal")
        _format_paragraph(paragraph, kind="body")
        blocks.append(paragraph._p)
        pending_table_caption = None
        index += 1

    for block in blocks:
        reference._p.addprevious(block)
    _enable_field_updates(document)
    document.save(output_path)
    return output_path


def _media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        }


def _count_xml(path: Path, token: bytes) -> int:
    with zipfile.ZipFile(path) as archive:
        return archive.read("word/document.xml").count(token)


def verify_document(source: str | Path, output: str | Path) -> dict[str, Any]:
    source_path = Path(source)
    output_path = Path(output)
    source_doc = Document(source_path)
    output_doc = Document(output_path)
    source_reference_index = next(index for index, paragraph in enumerate(source_doc.paragraphs) if paragraph.text.strip().replace(" ", "") == "参考文献")
    source_chapter_index = next(
        (
            index
            for index, paragraph in enumerate(source_doc.paragraphs)
            if paragraph.text.strip().replace(" ", "") == "第6章系统测试与实验分析"
        ),
        source_reference_index,
    )
    output_chapter_index = next(index for index, paragraph in enumerate(output_doc.paragraphs) if paragraph.text.strip().replace(" ", "") == "第6章系统测试与实验分析")
    output_reference_index = next(index for index, paragraph in enumerate(output_doc.paragraphs) if paragraph.text.strip().replace(" ", "") == "参考文献")
    source_prefix = [(paragraph.text, paragraph.style.name) for paragraph in source_doc.paragraphs[:source_chapter_index]]
    output_prefix = [(paragraph.text, paragraph.style.name) for paragraph in output_doc.paragraphs[:output_chapter_index]]
    source_boundary = source_doc.paragraphs[source_chapter_index]
    output_boundary = output_doc.paragraphs[output_chapter_index]
    source_prefix_table_count = _table_count_before_paragraph(source_doc, source_boundary)
    output_prefix_table_count = _table_count_before_paragraph(output_doc, output_boundary)
    source_prefix_field_count = _field_code_count_before_paragraph(source_doc, source_boundary)
    output_prefix_field_count = _field_code_count_before_paragraph(output_doc, output_boundary)
    source_tables = [[cell.text for row in table.rows for cell in row.cells] for table in source_doc.tables[:source_prefix_table_count]]
    output_original_tables = [[cell.text for row in table.rows for cell in row.cells] for table in output_doc.tables[:output_prefix_table_count]]
    source_media = _media_hashes(source_path)
    output_media = _media_hashes(output_path)
    with zipfile.ZipFile(output_path) as archive:
        bad_zip_member = archive.testzip()
    report = {
        "output_exists": output_path.is_file(),
        "output_size_bytes": output_path.stat().st_size,
        "zip_integrity": bad_zip_member is None,
        "docx_opened": True,
        "chapter_before_references": output_chapter_index < output_reference_index,
        "pre_chapter_paragraphs_unchanged": source_prefix == output_prefix,
        "original_tables_unchanged": source_tables == output_original_tables,
        "original_media_preserved": all(output_media.get(name) == digest for name, digest in source_media.items()),
        "source_paragraphs_before_chapter": len(source_prefix),
        "output_new_chapter_paragraphs": output_reference_index - output_chapter_index,
        "source_table_count": len(source_doc.tables),
        "output_table_count": len(output_doc.tables),
        "source_pre_chapter_table_count": source_prefix_table_count,
        "output_pre_chapter_table_count": output_prefix_table_count,
        "source_inline_shapes": len(source_doc.inline_shapes),
        "output_inline_shapes": len(output_doc.inline_shapes),
        "source_section_count": len(source_doc.sections),
        "output_section_count": len(output_doc.sections),
        "page_break_count_source": _count_xml(source_path, b'w:type="page"'),
        "page_break_count_output": _count_xml(output_path, b'w:type="page"'),
        "section_property_count_source": _count_xml(source_path, b"<w:sectPr"),
        "section_property_count_output": _count_xml(output_path, b"<w:sectPr"),
        "field_code_count_source": _count_xml(source_path, b"<w:instrText"),
        "field_code_count_output": _count_xml(output_path, b"<w:instrText"),
        "source_pre_chapter_field_code_count": source_prefix_field_count,
        "output_pre_chapter_field_code_count": output_prefix_field_count,
    }
    report["passed"] = all([
        report["zip_integrity"],
        report["chapter_before_references"],
        report["pre_chapter_paragraphs_unchanged"],
        report["original_tables_unchanged"],
        report["source_pre_chapter_table_count"] == report["output_pre_chapter_table_count"],
        report["original_media_preserved"],
        report["source_section_count"] == report["output_section_count"],
        report["page_break_count_source"] == report["page_break_count_output"],
        report["section_property_count_source"] == report["section_property_count_output"],
        report["source_pre_chapter_field_code_count"] == report["output_pre_chapter_field_code_count"],
    ])
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--verification", type=Path, required=True)
    args = parser.parse_args()
    build_document(args.source, args.markdown, args.output)
    report = verify_document(args.source, args.output)
    args.verification.parent.mkdir(parents=True, exist_ok=True)
    args.verification.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
