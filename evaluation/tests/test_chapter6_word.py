from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class Chapter6WordTests(unittest.TestCase):
    def test_build_document_inserts_markdown_before_references(self) -> None:
        from evaluation.scripts.chapter6.build_word import build_document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.docx"
            output = root / "output.docx"
            markdown = root / "chapter.md"
            document = Document()
            document.styles["Table Grid"].delete()
            document.add_heading("第5章 系统实现", level=1)
            document.add_paragraph("第五章正文保持不变。")
            document.add_heading("参考文献", level=1)
            document.add_paragraph("[1] Existing reference")
            document.save(source)
            markdown.write_text(
                "# 第6章系统测试与实验分析\n\n"
                "## 6.1测试目标\n\n"
                "新增正文。\n\n"
                "表6-1 示例表\n\n"
                "| 列A | 列B |\n| --- | --- |\n| 值1 | 值2 |\n",
                encoding="utf-8",
            )

            build_document(source, markdown, output)

            result = Document(output)
            texts = [paragraph.text for paragraph in result.paragraphs]
            self.assertLess(texts.index("第五章正文保持不变。"), texts.index("第6章系统测试与实验分析"))
            self.assertLess(texts.index("第6章系统测试与实验分析"), texts.index("参考文献"))
            self.assertEqual(texts[0:2], ["第5章 系统实现", "第五章正文保持不变。"])
            self.assertEqual(len(result.tables), 1)
            self.assertEqual(len(result.sections), 1)

    def test_verify_document_compares_prefix_before_existing_chapter_six(self) -> None:
        from evaluation.scripts.chapter6.build_word import build_document, verify_document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.docx"
            output = root / "output.docx"
            markdown = root / "chapter.md"
            document = Document()
            document.add_heading("第5章 系统实现", level=1)
            document.add_paragraph("第五章正文保持不变。")
            document.add_heading("第6章系统测试与实验分析", level=1)
            document.add_paragraph("需要替换的旧第六章。")
            field_paragraph = document.add_paragraph()
            field_run = field_paragraph.add_run()
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = "SEQ Figure"
            field_run._r.append(instruction)
            old_table = document.add_table(rows=1, cols=1)
            old_table.cell(0, 0).text = "旧第六章表格"
            document.add_heading("参考文献", level=1)
            document.add_paragraph("[1] Existing reference")
            document.save(source)
            markdown.write_text(
                "# 第6章系统测试与实验分析\n\n"
                "## 6.1测试目标\n\n"
                "替换后的新第六章。\n",
                encoding="utf-8",
            )

            build_document(source, markdown, output)
            report = verify_document(source, output)
            result = Document(output)
            texts = [paragraph.text for paragraph in result.paragraphs]

            self.assertTrue(report["pre_chapter_paragraphs_unchanged"])
            self.assertTrue(report["passed"])
            self.assertNotIn("需要替换的旧第六章。", texts)
            self.assertNotIn("旧第六章表格", [cell.text for table in result.tables for row in table.rows for cell in row.cells])


if __name__ == "__main__":
    unittest.main()
