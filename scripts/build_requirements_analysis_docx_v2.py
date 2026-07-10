from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import build_requirements_analysis_docx as source


PROJECT_TITLE = "基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计"
DOC_TITLE = "需求分析书"
DOC_NO = "NCE-SRS-001"
VERSION = "V1.1"
BASELINE = "SRS-BL-2026-06-20-R1"
DATE_TEXT = "2026 年 06 月 20 日"
OUTPUT_DIR = Path(r"E:\nan\《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》\开发管理")
OUTPUT_PATH = OUTPUT_DIR / f"02_《{PROJECT_TITLE}》{DOC_TITLE}.docx"


MODULE_INTROS = {
    "4.1": (
        "本模块负责建立系统的访问边界。游客只可使用公开浏览和注册能力；普通用户在登录后可提交检测、查看个人历史并生成报告；管理员在此基础上获得用户、知识库、模板、审核、日志和统计管理权限。",
        "认证信息采用令牌机制传递。需求重点不只是“能够登录”，还包括密码保护、禁用账号阻断、角色校验和当前用户信息回显，从而使后续所有业务操作都有可识别、可追责的主体。",
    ),
    "4.2": (
        "新闻检测入口既要兼容用户直接粘贴文本，也要支持从公开网页提取标题和正文。系统应在提交前完成输入校验、长度限制、链接安全检查和重复操作抑制，避免无效内容或危险地址进入分析链路。",
        "检测任务是整个系统的业务起点。前端应清楚呈现处理状态，后端应统一组织清洗、检索、模型分析、评分和持久化，并在外部服务不可用时返回可理解的降级结果。",
    ),
    "4.3": (
        "RAG 子系统用于把待鉴别新闻与本地可信知识、历史事实和联网证据联系起来。检索结果不是最终结论，而是供大语言模型和规则评分共同使用的证据材料。",
        "为了降低“检索到但不相关”造成的误判，系统需要记录证据来源、相似度和引用关系，并对多条证据的一致性进行仲裁。无有效证据时应明确标记，而不能把空结果包装成确定性判断。",
    ),
    "4.4": (
        "评分模块把模型判断、证据质量和规则特征转换为 0～100 分的可信度分值，并映射为可信、可疑、谣言和高风险四个等级。该分值用于辅助判断，不等同于行政或司法意义上的事实裁决。",
        "结果页面应同时给出结论、理由、证据和风险提示，使用户能够理解分数来源。若模型调用失败，系统应使用规则分数降级；若证据不足，应显式说明不确定性，避免制造虚假的精确感。",
    ),
    "4.5": (
        "检测历史用于保存用户已提交任务及其结论。普通用户只能访问自己的记录，管理员可在授权范围内查看全部记录。历史列表需要支持分页、筛选和进入详情。",
        "重新评估不是简单复制旧结论，而是基于当前知识库、模板和模型重新执行鉴别，并保留与原记录的关联，以便解释前后结果差异。",
    ),
    "4.6": (
        "报告功能把结构化检测结果整理为可归档、可阅读的 PDF 文档。报告应包含新闻摘要、各项分数、风险等级、分析说明、主要证据、生成时间和必要免责声明。",
        "报告文件属于检测记录的派生物，访问权限必须继承检测记录权限。生成失败时不得留下可下载的空文件或错误状态。",
    ),
    "4.7": (
        "公开高风险信息面向无需登录的访客，用于展示经过管理员审核且明确允许公开的内容。公开页面不能直接暴露未审核记录，也不能泄露提交用户信息和内部模型调用细节。",
        "公开范围由高风险、审核通过和允许公开三个条件共同决定，任何一个条件不满足都不得出现在公共接口中。",
    ),
    "4.8": (
        "管理员用户与检测管理用于维护账号状态和处理异常业务记录。相关操作具有较高权限，应记录操作者、对象、时间和结果，并对删除、禁用等动作进行确认。",
        "管理界面可以提供聚合统计，但不得返回密码散列、令牌或其他认证秘密。普通用户不得通过修改请求参数绕过权限边界。",
    ),
    "4.9": (
        "知识库是本地检索增强的基础。管理员可以维护新闻事实、来源、真实性标签、风险等级和正文，并触发向量化或重建索引。",
        "结构化记录与向量索引必须保持可追踪关系。索引失败不应破坏原始知识记录，系统应保留失败状态和重试入口。",
    ),
    "4.10": (
        "Prompt 模板用于约束不同分析场景下的大语言模型输入和输出。模板变化会直接影响鉴别结果，因此需要名称唯一、版本可辨、状态可控，并限制普通用户访问。",
        "系统应在启用模板前校验必要变量和格式；删除或停用正在使用的模板时应阻止操作或给出明确影响提示。",
    ),
    "4.11": (
        "高风险审核用于把自动检测结果转化为可公开的信息。系统首先生成待审核记录，管理员审阅新闻内容、分析说明和证据后，决定通过或拒绝，并单独决定是否公开。",
        "审核结论、审核人和审核时间应被保存。已公开记录若被撤销公开，应立即从公共查询结果中消失。",
    ),
    "4.12": (
        "统计分析帮助管理员掌握用户、检测、风险分布和知识库运行情况；系统日志用于定位错误、审计关键操作和判断外部服务健康状态。",
        "统计数据应来自一致的数据口径，日志应避免记录密码、密钥和完整令牌。查询量较大时可使用聚合或分页，但不能改变业务含义。",
    ),
    "4.13": (
        "定时采集是可配置的辅助能力，用于从指定来源补充候选新闻。采集任务默认不得绕过内容清洗、去重和来源校验，也不能自动把未审核内容标记为可信事实。",
        "部署方可以关闭调度器。关闭后系统不得创建后台采集任务，以避免测试环境和多实例部署中出现重复抓取。",
    ),
}


NFR_INTROS = {
    "7.1": "性能指标以受控试运行环境为测量边界，区分本地查询与依赖模型、搜索供应商的长耗时请求。对外部服务耗时应单独记录，避免把网络波动误判为数据库性能问题。",
    "7.2": "可靠性要求关注事务一致性、失败降级、任务恢复和数据备份。系统允许在外部智能服务不可用时降低能力，但不得返回伪造的成功结果。",
    "7.3": "安全要求覆盖身份认证、权限控制、输入校验、SSRF 防护、密钥管理和敏感日志治理，遵循最小权限和默认拒绝原则。",
    "7.4": "新闻正文可能包含个人信息或受版权保护的内容。系统需要明确处理目的、控制保存范围，并向用户说明可能发生的第三方模型或搜索服务传输。",
    "7.5": "易用性要求强调一致的术语、状态反馈、错误说明和键盘可操作性。风险颜色只能作为辅助提示，关键结论必须同时使用文字表达。",
    "7.6": "可维护性要求体现为清晰分层、集中配置、可替换外部供应商和稳定测试。需求文档中的接口与数据术语应与代码和后续设计文档保持一致。",
    "7.7": "可观测性用于回答任务何时开始、经过哪些阶段、在哪里失败以及由谁执行。日志和审计信息应足够定位问题，同时避免扩大敏感信息暴露面。",
}


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, size=8, color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="宋体", size=12, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(24)

    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Heading 1"].paragraph_format.page_break_before = False

    if "Requirement" not in doc.styles:
        req_style = doc.styles.add_style("Requirement", WD_STYLE_TYPE.PARAGRAPH)
    else:
        req_style = doc.styles["Requirement"]
    req_style.base_style = normal
    req_style.font.name = "Times New Roman"
    req_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    req_style.font.size = Pt(11)
    req_style.paragraph_format.first_line_indent = Pt(0)
    req_style.paragraph_format.left_indent = Cm(0.74)
    req_style.paragraph_format.first_line_indent = Cm(-0.74)
    req_style.paragraph_format.line_spacing = 1.35
    req_style.paragraph_format.space_after = Pt(3)


def configure_page(section, header_footer=True):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.4)
    if not header_footer:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""
        return
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.first_line_indent = Pt(0)
    r = hp.add_run(f"《{PROJECT_TITLE}》{DOC_TITLE}")
    set_font(r, "宋体", 9, False, "666666")
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Pt(0)
    run = fp.add_run("—  ")
    set_font(run, "宋体", 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    run = fp.add_run("  —")
    set_font(run, "宋体", 9)


def add_text(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, "宋体", 12, True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, "宋体", 12)
    else:
        r = p.add_run(text)
        set_font(r, "宋体", 12)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Cm(0.74)
        r = p.add_run(item)
        set_font(r, "宋体", 11.5)


def add_plain_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        if widths:
            cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(str(header))
        set_font(r, "黑体", font_size, True)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            r = p.add_run(str(value))
            set_font(r, "宋体", font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cover(doc: Document):
    section = doc.sections[0]
    configure_page(section, header_footer=False)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("浙江万里学院《移动应用开发管理》")
    set_font(r, "宋体", 22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("课程大作业封面")
    set_font(r, "宋体", 22)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("教师填写：")
    set_font(r, "黑体", 11, True)

    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [2.7, 3.8, 4.0, 5.5]
    for i, cell in enumerate(t.rows[0].cells):
        cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 180, 140, 180, 140)
        set_cell_border(cell, 10)
    t.rows[0].height = Cm(2.6)
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    for idx, text in ((0, "得分"), (2, "任课教师签名")):
        p = t.cell(0, idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(text)
        set_font(r, "宋体", 14)
    p = t.cell(0, 3).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(38)
    r = p.add_run("年  月  日")
    set_font(r, "宋体", 11)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(50)
    spacer.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("学生填写：")
    set_font(r, "黑体", 11, True)

    s = doc.add_table(rows=6, cols=4)
    s.alignment = WD_TABLE_ALIGNMENT.CENTER
    s.autofit = False
    widths = [2.6, 5.4, 2.6, 5.4]
    for row in s.rows:
        row.height = Cm(1.05)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 100, 120, 100, 120)
            set_cell_border(cell, 10)
    s.cell(2, 1).merge(s.cell(2, 3))
    s.cell(4, 1).merge(s.cell(4, 3))
    s.cell(5, 1).merge(s.cell(5, 3))
    values = {
        (0, 0): "姓名", (0, 2): "学号",
        (1, 0): "所在学院", (1, 1): "信息与智能工程学院", (1, 2): "班级",
        (2, 0): "课程名称", (2, 1): "移动应用开发管理",
        (3, 0): "任课教师", (3, 1): "熊波", (3, 2): "课程学分", (3, 3): "2",
        (4, 0): "上课时间", (4, 1): "2025 至 2026 学年第二学期",
        (5, 0): "递交时间", (5, 1): "2026 年 06 月 26 日",
    }
    for (ri, ci), value in values.items():
        p = s.cell(ri, ci).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(value)
        set_font(r, "宋体", 11.5)


def add_inner_title_and_control(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(50)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(f"《{PROJECT_TITLE}》")
    set_font(r, "黑体", 20, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run(DOC_TITLE)
    set_font(r, "黑体", 26, True)

    add_plain_table(doc, ["项目", "内容"], [
        ["文档编号", DOC_NO], ["版本号", VERSION], ["需求基线", BASELINE],
        ["编制依据", "GB/T 9385—2008、GB/T 8567—2006"],
        ["事实来源", "当前源码、数据库模型、接口定义、测试结果及 codebase 知识图谱"],
        ["文档状态", "修订正式版"], ["编制日期", DATE_TEXT],
    ], widths=[3.5, 12.0], font_size=10)

    doc.add_heading("修订说明", level=1)
    add_text(doc, "本版按照课程文档阅读习惯重新排版。具体需求改为分模块正文条款，仅在确有对照需要的位置使用表格；系统边界、业务规则和需求数量保持不变。后续设计、测试和验收文档应引用 V1.1 及需求基线 SRS-BL-2026-06-20-R1。")


def add_abstract_and_toc(doc: Document):
    doc.add_page_break()
    doc.add_heading("摘要", level=1)
    add_text(doc, "本需求分析书面向“基于 RAG 与大语言模型的网络新闻真伪鉴别系统”。系统接收用户输入的新闻标题、正文或公开网页链接，通过本地知识库检索、可选联网证据搜索、大语言模型分析和规则评分形成可信度结论，并提供检测历史、报告生成、公开高风险信息以及后台知识库、Prompt、审核、统计和日志管理能力。")
    add_text(doc, "文档以当前项目源码和数据库模型为事实基础，结合 codebase 知识图谱核对主要模块、实体和调用关系。需求分为 122 项功能需求与 45 项非功能需求。每项需求均保留稳定编号，并标注优先级、当前状态和建议验证方法，以支持后续设计、开发、测试和验收。")
    add_text(doc, "本系统提供的是辅助性可信度判断。输出分值和风险等级用于提示用户核查信息，不替代新闻主管部门、专业事实核查机构或司法机关的最终认定。")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("关键词：")
    set_font(r, "黑体", 12, True)
    r = p.add_run("网络新闻鉴别；检索增强生成；大语言模型；证据仲裁；可信度评分")
    set_font(r, "宋体", 12)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("目录")
    set_font(r, "黑体", 18, True)
    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.first_line_indent = Pt(0)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    toc_p._p.append(fld)


def add_requirement_paragraph(doc, req):
    req_id, text, priority, status, verify = req
    p = doc.add_paragraph(style="Requirement")
    r = p.add_run(f"{req_id}  ")
    set_font(r, "Times New Roman", 10.5, True)
    r = p.add_run(text)
    set_font(r, "宋体", 11)
    r = p.add_run(f"（优先级：{priority}；状态：{status}；验证：{verify}）")
    set_font(r, "宋体", 9, False, "666666")


def build_main_content(doc: Document):
    doc.add_page_break()
    doc.add_heading("1 引言", level=1)
    doc.add_heading("1.1 编写目的", level=2)
    add_text(doc, "本文档用于完整说明网络新闻真伪鉴别系统在当前需求基线下应具备的业务能力、外部接口、数据约束、质量属性和验收条件。它既是概要设计和详细设计的输入，也是开发实现、测试设计、项目验收和后续变更控制的共同依据。")
    add_text(doc, "需求描述强调系统必须实现的外部可观察行为，而不在本阶段固化不必要的内部实现细节。涉及现有代码事实的内容以工作区源码、数据库模型和已执行测试为准；涉及目标质量的内容作为验收约束，在部署或正式验收时进一步测量。")
    doc.add_heading("1.2 项目背景", level=2)
    add_text(doc, "互联网新闻传播速度快、来源复杂，普通用户很难在短时间内判断内容是否可信。单纯依赖关键词规则容易漏判，单纯依赖大语言模型又可能产生无依据的结论。因此，本项目采用本地 RAG 检索、可选联网证据、大语言模型分析和可解释评分相结合的方式，为用户提供辅助核查工具。")
    add_text(doc, "项目属于浙江万里学院《移动应用开发管理》课程大作业，当前形态为面向课程设计、毕业设计原型和小规模受控试运行的 Web 系统。后端使用 FastAPI，数据存储采用 MySQL，前端使用 Vue；模型、搜索和向量服务通过配置接入。")
    doc.add_heading("1.3 适用范围", level=2)
    add_text(doc, "系统适用于用户对公开网络新闻进行初步可信度分析，也适用于管理员维护核查知识、审核高风险记录和观察系统运行情况。系统不负责采集或处理国家秘密，不面向未经授权的内网资源，不承诺对所有新闻给出绝对正确结论。")
    doc.add_heading("1.4 术语和缩略语", level=2)
    add_plain_table(doc, ["术语", "说明"], [
        ["RAG", "Retrieval-Augmented Generation，检索增强生成。"],
        ["LLM", "Large Language Model，大语言模型。"],
        ["证据仲裁", "对多条检索证据的相关性、一致性和可用性进行综合判断。"],
        ["可信度分数", "0～100 的辅助性评分，分数越高表示当前证据下越可信。"],
        ["高风险记录", "最终分数落入高风险区间并进入管理员审核流程的检测记录。"],
        ["M/S/C", "Must/Should/Could，分别表示必须、应该和可选优先级。"],
        ["I/A/T", "Implemented/Acceptance target/TBD，分别表示已实现、验收目标和待定。"],
    ], widths=[3.2, 12.3], font_size=9.5)
    doc.add_heading("1.5 参考资料", level=2)
    add_plain_table(doc, ["文件", "用途"], [
        ["GB/T 9385—2008《计算机软件需求规格说明规范》", "规定需求规格说明的主要内容和表达要求。"],
        ["GB/T 8567—2006《计算机软件文档编制规范》", "规定软件文档的编制和组织要求。"],
        ["GB/T 11457—2006《信息技术 软件工程术语》", "统一软件工程术语。"],
        ["GB/T 25000.10—2016《系统与软件工程 系统与软件质量要求和评价》", "用于组织非功能质量要求。"],
        ["项目源码、接口、模型及测试资料", "用于核对系统当前事实和实现状态。"],
    ], widths=[7.0, 8.5], font_size=9.3)
    doc.add_heading("1.6 需求表述约定", level=2)
    add_text(doc, "正文使用“应”表示可验证的强制要求，使用“宜”表示推荐要求，使用“可”表示允许或可选能力。需求编号在版本演进中保持稳定；需求修改通过修订记录或变更单说明，不因文字调整而随意复用编号。")

    doc.add_heading("2 总体描述", level=1)
    doc.add_heading("2.1 产品定位与系统边界", level=2)
    add_text(doc, "本系统位于用户、公开新闻网页、本地知识库和外部智能服务之间。用户通过浏览器提交新闻；系统完成文本清洗、证据检索、模型分析、规则评分和结果保存；管理员通过后台维护知识、模板和审核状态。外部模型与搜索服务只作为受控依赖，不拥有本系统用户权限和业务数据管理权限。")
    add_text(doc, "系统边界内包括账号与权限、新闻检测、RAG 检索、证据仲裁、评分分级、历史记录、报告、高风险公开、知识库、Prompt、统计日志和可选采集任务。新闻网站本身、模型供应商能力、搜索引擎索引质量以及人工事实核查机构不属于系统控制范围。")
    doc.add_heading("2.2 建设目标", level=2)
    add_bullets(doc, [
        "为普通用户提供简单、可解释的新闻可信度检测入口。",
        "通过本地知识和联网证据减少模型无依据判断。",
        "形成可追踪的检测记录、证据引用、评分和报告。",
        "为管理员提供知识维护、高风险审核、统计和审计能力。",
        "在外部服务异常时保持可理解的降级行为，而不是静默失败。",
    ])
    doc.add_heading("2.3 用户及其特征", level=2)
    add_plain_table(doc, ["用户角色", "主要目标", "权限范围"], [
        ["游客", "了解系统并查看公开高风险信息", "注册、登录、查看已审核公开内容"],
        ["普通用户", "提交新闻并查看个人鉴别结果", "检测、历史、详情、重新评估、个人报告"],
        ["管理员", "维护系统资料并审核风险内容", "普通用户能力及全部后台管理能力"],
        ["维护人员", "保障部署、配置和故障恢复", "通过受控运维环境处理配置、日志与备份"],
    ], widths=[2.6, 6.0, 6.9], font_size=9.3)
    doc.add_heading("2.4 运行方式与降级模式", level=2)
    add_text(doc, "正常模式下，本地检索、可选联网搜索、证据仲裁和大语言模型均可用，系统输出完整评分和证据说明。若联网搜索关闭或失败，系统继续使用本地知识；若模型失败，系统使用规则分数降级；若知识检索无结果，系统说明证据不足。数据库不可写时，检测任务不得报告持久化成功。")
    doc.add_heading("2.5 假设、依赖和约束", level=2)
    add_text(doc, "系统依赖可用的数据库、向量检索组件和至少一种模型接入配置。用户提交的链接必须是可公开访问的 HTTP/HTTPS 地址。外部服务的响应速度、配额和内容政策可能变化，因此系统应通过配置和适配层隔离差异。")

    doc.add_heading("3 业务需求与用例分析", level=1)
    doc.add_heading("3.1 核心业务流程", level=2)
    add_text(doc, "普通用户登录后输入新闻文本或先从链接提取预览内容。系统校验输入并创建检测任务，随后清洗文本、检索本地知识、按配置搜索联网证据、调用模型分析并执行证据仲裁。评分模块根据模型分、证据质量分和规则分计算最终分数，形成风险等级、理由和建议，最后保存检测、证据和过程信息。")
    add_text(doc, "当结果属于高风险时，记录默认进入待审核状态且不公开。管理员审阅内容和证据后，可以批准、拒绝或调整公开状态。普通用户可在个人历史中查看结果、重新评估并生成报告。")
    doc.add_heading("3.2 主要用例", level=2)
    for title, body in [
        ("UC-01 注册与登录", "游客填写用户名、密码和可选邮箱完成注册，随后通过登录获取访问令牌。账号被禁用或凭据错误时，系统拒绝访问并给出明确提示。"),
        ("UC-02 提交新闻检测", "普通用户输入标题与正文，或从公开链接提取预览后提交。系统返回任务结果，展示分数、风险等级、分析理由和证据。"),
        ("UC-03 查看历史和重新评估", "用户在分页历史中选择记录查看详情；当知识库或模型配置更新后，可发起重新评估并获得新记录。"),
        ("UC-04 生成检测报告", "用户对有权限的检测记录生成 PDF 报告并下载。报告内容与当次检测结果一致。"),
        ("UC-05 维护知识与模板", "管理员维护知识条目、向量状态和 Prompt 模板，使检索和模型输出保持可控。"),
        ("UC-06 审核高风险信息", "管理员查看待审核高风险记录，填写审核意见，决定通过或拒绝，并控制是否公开。"),
        ("UC-07 查看统计与日志", "管理员查看用户、检测、风险和知识统计，并通过日志定位异常操作或外部服务故障。"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(title + "。")
        set_font(r, "黑体", 12, True)
        r = p.add_run(body)
        set_font(r, "宋体", 12)
    doc.add_heading("3.3 业务规则", level=2)
    add_plain_table(doc, ["规则编号", "规则内容"], [
        ["BR-01", "最终分数不低于 80 为可信，60～79 为可疑，40～59 为谣言，低于 40 为高风险。"],
        ["BR-02", "高风险记录创建后默认为待审核且不公开。"],
        ["BR-03", "公开记录必须同时满足高风险、审核通过和允许公开三个条件。"],
        ["BR-04", "普通用户只能访问自己的检测与报告；管理员按权限访问全部记录。"],
        ["BR-05", "模型失败时不得使用未完成仲裁的模型候选结论，应按规则分降级。"],
        ["BR-06", "任何分数和风险等级都必须附带辅助判断免责声明。"],
    ], widths=[2.4, 13.1], font_size=9.5)

    doc.add_heading("4 功能需求", level=1)
    add_text(doc, "本章定义系统必须提供的功能。状态 I 表示当前代码已实现，A 表示需要在验收或部署阶段确认，T 表示尚待决策；验证方式 T、D、I、A 分别表示测试、演示、检查和分析。元数据以较小字号附在需求条款后，不影响正文连续阅读。")
    for section_no, section_title, reqs in source.REQ_GROUPS:
        doc.add_heading(f"{section_no} {section_title}", level=2)
        for para in MODULE_INTROS.get(section_no, (section_title,)):
            add_text(doc, para)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run("具体需求")
        set_font(r, "黑体", 12, True)
        for req in reqs:
            add_requirement_paragraph(doc, req)

    doc.add_heading("5 外部接口需求", level=1)
    doc.add_heading("5.1 用户界面", level=2)
    add_text(doc, "前台界面至少包含首页、登录注册、新闻检测、检测结果、历史记录、报告和公开高风险信息。后台界面至少包含总览、用户、检测、知识库、Prompt、高风险审核、统计和日志管理。页面应使用统一导航和风险术语，表单错误显示在相关字段附近，长耗时检测展示明确的处理中状态。")
    doc.add_heading("5.2 软件接口", level=2)
    add_text(doc, "前端通过 JSON REST API 与 FastAPI 后端通信，下载类接口返回文件流。受保护接口使用 Bearer 令牌；后台接口在认证基础上继续校验管理员角色。主要接口按业务域概括如下。")
    add_plain_table(doc, ["业务域", "主要路径", "说明"], [
        ["认证", "/api/auth/*", "注册、登录和当前用户信息"],
        ["新闻检测", "/api/detect/*", "链接预览、提交检测、历史、详情和重新评估"],
        ["报告", "/api/reports/*", "生成、查询和下载检测报告"],
        ["公开信息", "/api/public/*", "查询审核通过且允许公开的高风险记录"],
        ["知识库", "/api/admin/knowledge/*", "知识条目和向量索引管理"],
        ["Prompt", "/api/admin/prompts/*", "模型提示模板管理"],
        ["后台管理", "/api/admin/*", "用户、检测、审核、统计和日志"],
    ], widths=[2.7, 5.3, 7.5], font_size=9.1)
    doc.add_heading("5.3 外部服务接口", level=2)
    add_text(doc, "模型接口接收清洗后的新闻文本、证据和 Prompt，返回结构化分析结果；搜索接口根据查询词返回标题、摘要、来源和链接；向量服务提供文本嵌入和相似度检索。系统应设置连接和读取超时，并把供应商错误转换为内部统一状态，不直接向用户暴露密钥和底层异常堆栈。")
    doc.add_heading("5.4 通信和文件接口", level=2)
    add_text(doc, "生产环境应使用 HTTPS。JSON 文本统一采用 UTF-8；时间使用带时区的标准格式或在接口文档中明确时区。PDF 报告下载应设置正确的内容类型和文件名，前端不得把错误 JSON 当作 PDF 保存。")

    doc.add_heading("6 数据需求", level=1)
    doc.add_heading("6.1 数据对象及关系", level=2)
    add_text(doc, "系统核心数据由用户、检测记录、证据匹配、知识条目、Prompt 模板、报告、系统日志和采集任务构成。用户与检测记录是一对多关系；检测记录与证据、报告是一对多关系；知识条目可被多个证据引用；管理员通过审核字段和日志记录对高风险内容实施治理。")
    add_plain_table(doc, ["数据对象", "主要内容", "关键约束"], [
        ["User", "账号、密码散列、邮箱、角色、状态", "用户名唯一；不得保存明文密码"],
        ["DetectionRecord", "新闻输入、各项分数、风险等级、分析结果、审核状态", "归属用户；分数范围 0～100"],
        ["EvidenceMatch", "证据来源、摘要、链接、相似度和引用信息", "必须关联检测记录"],
        ["KnowledgeItem", "事实文本、来源、标签、风险和向量状态", "原始记录与向量状态可追踪"],
        ["PromptTemplate", "名称、场景、模板、版本和启用状态", "名称/版本可区分；仅管理员维护"],
        ["Report", "报告路径、状态、生成时间和错误信息", "权限继承检测记录"],
        ["SystemLog", "时间、级别、模块、事件、对象和操作者", "不得记录密钥或完整密码"],
        ["CrawlTask", "来源、调度、执行状态和错误信息", "关闭采集时不得创建任务"],
    ], widths=[3.0, 7.2, 5.3], font_size=8.8)
    doc.add_heading("6.2 数据状态与完整性", level=2)
    add_text(doc, "检测记录创建后应保存输入快照和评分结果。高风险审核状态从 pending 转换为 approved 或 rejected；公开标志只有在审核通过后才有效。报告状态至少区分处理中、成功和失败。知识向量状态应区分未处理、处理中、成功和失败，并保存必要的错误原因。")
    add_text(doc, "数据库写操作应使用事务保证关联数据一致。删除用户、检测或知识项时，应根据业务规则采用限制删除、级联删除或逻辑删除，不能留下无法解释的孤立证据和报告。")
    doc.add_heading("6.3 数据输入与保留", level=2)
    add_text(doc, "用户名、密码、邮箱、URL、分页参数、新闻文本长度和枚举值应在接口边界校验。对于第三方网页提取内容，应保存必要来源信息并遵循版权和隐私要求。生产部署的数据保留期限、用户删除请求处理方式和备份周期在验收前由项目负责人确认。")

    doc.add_heading("7 非功能需求", level=1)
    add_text(doc, "非功能需求用于规定系统在性能、可靠性、安全、隐私、易用性、可维护性和可观测性方面应达到的质量水平。凡标记为 A 的指标均需在目标部署环境中补充测试证据。")
    for section_no, section_title, reqs in source.NFR_GROUPS:
        doc.add_heading(f"{section_no} {section_title}", level=2)
        add_text(doc, NFR_INTROS.get(section_no, section_title))
        for req in reqs:
            add_requirement_paragraph(doc, req)

    doc.add_heading("8 运行环境与资源需求", level=1)
    doc.add_heading("8.1 软件环境", level=2)
    add_text(doc, "服务端运行环境应支持 Python、FastAPI 及项目依赖，数据库采用 MySQL，前端采用现代浏览器可运行的 Vue 构建产物。开发和测试环境可在 Windows 上运行，受控部署宜使用 Linux 服务器或容器。Chrome、Edge 等主流浏览器应能完成核心业务。")
    doc.add_heading("8.2 硬件与网络", level=2)
    add_text(doc, "课程演示环境至少应具备 4 核处理器、8 GB 内存和足够保存数据库、日志及报告的磁盘空间。若使用本地大模型或本地嵌入模型，硬件配置需根据模型规模另行评估。系统需要访问配置的模型、搜索或网页来源；网络不可用时按降级模式工作。")
    doc.add_heading("8.3 配置资源", level=2)
    add_text(doc, "数据库连接、JWT 密钥、模型与搜索供应商密钥、超时时间、上传/文本限制、采集开关和日志级别应通过环境配置提供。生产密钥不得提交到源码仓库，模板示例值不得直接用于正式环境。")

    doc.add_heading("9 设计和实现约束", level=1)
    add_text(doc, "后端应保持 API、Service、CRUD、Model/Schema 分层，前端通过统一请求模块访问接口。系统以 REST/JSON 作为主要交互方式，数据库结构通过可追踪的迁移或初始化脚本维护。")
    add_text(doc, "模型输出必须经过结构校验和业务转换，不能把自由文本直接当作可信数据库字段。URL 抽取必须防止访问本机、内网和受限地址。高权限操作必须进行服务端鉴权，不能只依赖前端隐藏按钮。")
    add_text(doc, "课程项目允许采用外部模型和搜索服务，但文档、界面和报告应清楚说明服务依赖和辅助判断性质。若供应商接口发生变化，应优先修改适配层，不改变稳定业务接口和需求编号。")

    doc.add_heading("10 合格性规定", level=1)
    doc.add_heading("10.1 验收原则", level=2)
    add_text(doc, "验收以稳定需求编号为最小追踪单位。M 级需求原则上必须全部满足；S 级需求允许在不影响核心业务的前提下形成经批准的遗留项；C 级需求由项目负责人根据进度决定。标记为 I 不等于自动验收通过，仍需通过测试、演示、检查或分析提供证据。")
    doc.add_heading("10.2 验收活动", level=2)
    add_plain_table(doc, ["验收类别", "主要内容", "典型证据"], [
        ["功能测试", "注册登录、检测、历史、报告、后台管理和审核流程", "自动化测试结果、接口记录、演示截图"],
        ["接口测试", "状态码、参数校验、权限边界、文件下载和异常响应", "接口测试报告、OpenAPI 核对记录"],
        ["数据检查", "核心实体、关联完整性、状态迁移和事务回滚", "数据库检查记录、迁移脚本"],
        ["安全检查", "鉴权、越权、SSRF、密钥、敏感日志和输入校验", "安全测试清单和整改记录"],
        ["性能测试", "普通查询 P95、并发检测、外部服务耗时和资源占用", "性能脚本、监控和统计报告"],
        ["文档审查", "需求、设计、测试和验收文档的一致性", "审查意见和需求追踪记录"],
    ], widths=[2.8, 7.0, 5.7], font_size=9.2)
    doc.add_heading("10.3 当前验证基线", level=2)
    add_text(doc, "在本需求基线编制期间，后端测试结果为 392 项通过，前端测试结果为 14 项通过，前端生产构建成功并存在体积提示。上述结果用于证明当前实现具备可验证基础，不替代最终验收环境中的安全、性能和模型效果测试。")

    doc.add_heading("11 需求可追踪性", level=1)
    add_text(doc, "需求通过业务目标、需求编号、实现模块和验证证据建立追踪。后续文档应引用需求编号，而不是仅使用页面名称或代码文件名。需求发生变更时，应同步评估接口、数据、测试和验收影响。")
    add_plain_table(doc, ["需求域", "需求编号范围", "主要实现位置", "主要验证方式"], [
        ["认证与授权", "FR-AUTH-*", "auth API、用户服务、权限依赖", "认证、越权和账号状态测试"],
        ["新闻检测与评分", "FR-DET-*、FR-RAG-*、FR-SCORE-*", "detect API、detection service、RAG/LLM 服务", "检测链路、降级和分数边界测试"],
        ["历史与报告", "FR-HIS-*、FR-REP-*", "历史接口、报告服务", "归属权限、重评估和文件下载测试"],
        ["公开与审核", "FR-PUB-*、FR-HR-*", "公共接口、高风险审核服务", "三条件公开规则和状态迁移测试"],
        ["后台管理", "FR-ADM-*、FR-KB-*、FR-PRM-*", "管理员接口、知识和模板服务", "CRUD、角色和索引状态测试"],
        ["统计日志与采集", "FR-STAT-*、FR-CRAWL-*", "统计、日志和调度模块", "口径、审计和开关测试"],
        ["质量属性", "NFR-*", "全系统及部署配置", "测试、检查、演示和分析"],
    ], widths=[2.8, 4.1, 5.3, 3.8], font_size=8.6)

    doc.add_heading("12 尚未解决的问题", level=1)
    for title, body in [
        ("TBD-01 生产规模", "正式并发用户数、日检测量和数据库容量需由部署方确认，确认后据此调整性能验收参数。"),
        ("TBD-02 模型效果", "需确定正式验收数据集、样本标签来源以及准确率、召回率或一致性指标。"),
        ("TBD-03 数据保留", "用户新闻文本、检测记录、日志和报告的保留期限及删除流程需经项目负责人批准。"),
        ("TBD-04 第三方合规", "正式使用的模型与搜索供应商、数据处理区域和服务条款需在部署前确认。"),
        ("TBD-05 运维策略", "备份周期、恢复时间目标、告警接收人和密钥轮换周期需形成部署规定。"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(title + "。")
        set_font(r, "黑体", 12, True)
        r = p.add_run(body)
        set_font(r, "宋体", 12)

    doc.add_heading("附录 A 需求编号统计", level=1)
    rows = []
    for section_no, section_title, reqs in source.REQ_GROUPS + source.NFR_GROUPS:
        must = sum(1 for r in reqs if r[2] == "M")
        implemented = sum(1 for r in reqs if r[3] == "I")
        rows.append([f"{section_no} {section_title}", len(reqs), must, implemented, f"{reqs[0][0]}～{reqs[-1][0]}"])
    add_plain_table(doc, ["需求域", "数量", "M级", "已实现", "编号范围"], rows, widths=[5.2, 1.4, 1.4, 1.7, 5.3], font_size=8.2)
    add_text(doc, "统计结果：功能需求 122 项，非功能需求 45 项，共 167 项。编号用于后续概要设计、详细设计、测试报告和验收总结报告建立一一对应关系。")

    doc.add_heading("附录 B 需求基线声明", level=1)
    add_bullets(doc, [
        "本文件基于 2026 年 06 月 20 日工作区源码、设计资料、接口、数据库模型和 codebase 知识图谱编制。",
        "知识图谱用于核对核心实体、服务和调用关系；路由关联不足的部分采用定向源码检查补充。",
        "本版只调整封面、组织方式和表达形式，不减少已确认的功能与质量需求。",
        "后续需求变更必须记录原因、影响范围、批准人和生效版本。",
    ])


def enable_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_cover(doc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(body, header_footer=True)
    add_inner_title_and_control(doc)
    add_abstract_and_toc(doc)
    build_main_content(doc)
    enable_update_fields(doc)

    doc.core_properties.title = f"《{PROJECT_TITLE}》{DOC_TITLE}"
    doc.core_properties.subject = "软件需求规格说明（文字化修订版）"
    doc.core_properties.author = "浙江万里学院《移动应用开发管理》课程项目组"
    doc.core_properties.keywords = "需求分析, RAG, 大语言模型, 新闻真伪鉴别, GB/T 9385-2008"
    doc.core_properties.comments = "V1.1：按课程封面重排，正文改为论述和条款为主"
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
