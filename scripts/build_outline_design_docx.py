from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = Path(__file__).resolve().parents[1] / ".artifacts" / "docs"
OUTPUT_PATH = OUTPUT_DIR / "03_《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》概要设计书.docx"

PROJECT_TITLE = "基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计"
SYSTEM_NAME = "智闻辨真"
DOC_TITLE = "概要设计书"
DOC_NO = "NCE-HLD-001"
VERSION = "V1.0"
BASELINE = "HLD-BL-2026-06-20"
DATE_TEXT = "2026 年 06 月 20 日"

FIGURE_DIR = PROJECT_ROOT / "docs" / "thesis" / "figures" / "chapter4"

NAVY = "183B56"
BLUE = "176B87"
TEAL = "2A9D8F"
GOLD = "C58A18"
RED = "B42318"
GREEN = "187A5A"
GRAY = "667085"
DARK = "243746"
WHITE = "FFFFFF"
LIGHT_GRAY = "F6F8FA"
PALE_BLUE = "EEF6F8"
PALE_TEAL = "EAF8F5"
PALE_GOLD = "FFF8E8"
PALE_RED = "FFF0F0"


DB_TABLES: list[dict[str, object]] = [
    {
        "name": "users",
        "title": "用户表",
        "purpose": "保存普通用户与管理员的认证主体、角色及账户状态。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "用户主键"],
            ["username", "VARCHAR(50)", "否", "UQ", "登录用户名"],
            ["password_hash", "VARCHAR(255)", "否", "—", "bcrypt 密码散列，不保存明文"],
            ["email", "VARCHAR(100)", "是", "UQ", "邮箱地址"],
            ["role", "VARCHAR(20)", "否", "默认 user", "user/admin"],
            ["status", "VARCHAR(20)", "否", "默认 active", "active/disabled"],
            ["created_at", "DATETIME", "否", "默认当前时间", "创建时间"],
            ["updated_at", "DATETIME", "否", "自动更新", "最后更新时间"],
        ],
        "indexes": "PRIMARY(id)，UNIQUE(username)，UNIQUE(email)。",
    },
    {
        "name": "knowledge_items",
        "title": "知识条目表",
        "purpose": "保存可核验新闻、事实背景、辟谣材料及其向量同步状态，是本地 RAG 的业务事实源。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "知识条目主键"],
            ["title", "VARCHAR(255)", "否", "—", "新闻或事实标题"],
            ["content", "TEXT", "否", "—", "正文"],
            ["category", "VARCHAR(50)", "是", "IDX", "新闻类别"],
            ["truth_label", "VARCHAR(30)", "否", "IDX", "可信/虚假等事实标签"],
            ["source_name", "VARCHAR(100)", "是", "—", "来源名称"],
            ["source_url", "VARCHAR(500)", "是", "—", "来源链接"],
            ["publish_time", "DATETIME", "是", "—", "发布时间"],
            ["summary", "TEXT", "是", "—", "摘要"],
            ["keywords", "VARCHAR(500)", "是", "—", "逗号分隔关键词"],
            ["debunking_explanation", "TEXT", "是", "—", "辟谣或核查说明"],
            ["risk_level", "VARCHAR(30)", "是", "IDX", "知识条目的风险等级"],
            ["admin_note", "TEXT", "是", "—", "管理员内部备注"],
            ["vector_id", "VARCHAR(100)", "是", "—", "Chroma 文档标识，格式 knowledge:<id>"],
            ["vector_sync_status", "VARCHAR(20)", "否", "IDX，默认 pending", "pending/synced/failed/delete_failed"],
            ["vector_sync_error", "TEXT", "是", "—", "向量同步错误摘要"],
            ["created_at", "DATETIME", "否", "默认当前时间", "创建时间"],
            ["updated_at", "DATETIME", "否", "自动更新", "最后更新时间"],
        ],
        "indexes": "PRIMARY(id)，idx_category，idx_truth_label，idx_risk_level，idx_vector_sync_status。",
    },
    {
        "name": "detection_records",
        "title": "检测记录表",
        "purpose": "保存一次新闻鉴别的输入快照、评分分量、最终结论、分析载荷及高风险审核状态。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "检测记录主键"],
            ["user_id", "BIGINT", "是", "FK→users.id，SET NULL，IDX", "提交用户；游客为空"],
            ["input_title", "VARCHAR(255)", "否", "—", "规范化后的新闻标题"],
            ["input_content", "TEXT", "否", "—", "规范化后的新闻正文"],
            ["category", "VARCHAR(50)", "是", "—", "新闻类别"],
            ["keywords", "VARCHAR(500)", "是", "—", "提取与模型合并后的关键词"],
            ["final_score", "DECIMAL(5,2)", "否", "—", "0～100 最终可信度分"],
            ["evidence_score", "DECIMAL(5,2)", "否", "—", "有效证据相似度均值"],
            ["llm_score", "DECIMAL(5,2)", "否", "—", "大模型可信度分量"],
            ["rule_score", "DECIMAL(5,2)", "否", "—", "确定性规则评分"],
            ["risk_level", "VARCHAR(30)", "否", "IDX", "四级风险分类"],
            ["judgement_result", "VARCHAR(100)", "否", "—", "结论摘要"],
            ["reason", "TEXT", "是", "—", "判断理由"],
            ["risk_points", "TEXT", "是", "JSON 文本", "风险点列表"],
            ["suggestion", "TEXT", "是", "—", "核查与传播建议"],
            ["analysis_payload", "TEXT", "是", "JSON 文本", "候选/排除证据、契约版本、降级状态等审计载荷"],
            ["is_high_risk", "BOOLEAN", "否", "IDX，默认 0", "是否进入高风险复核范围"],
            ["review_status", "VARCHAR(20)", "否", "IDX，默认 pending", "pending/approved/rejected"],
            ["is_public", "BOOLEAN", "否", "IDX，默认 0", "是否允许公开展示"],
            ["admin_remark", "TEXT", "是", "—", "管理员复核备注"],
            ["reviewed_at", "DATETIME", "是", "—", "复核时间"],
            ["reviewed_by", "BIGINT", "是", "FK→users.id，SET NULL，IDX", "复核管理员"],
            ["report_url", "VARCHAR(500)", "是", "—", "报告下载地址或相对路径"],
            ["created_at", "DATETIME", "否", "IDX，默认当前时间", "检测时间"],
            ["updated_at", "DATETIME", "否", "自动更新", "最后更新时间"],
        ],
        "indexes": "PRIMARY(id)，idx_detection_user_id，idx_detection_risk_level，idx_detection_is_high_risk，idx_detection_review_status，idx_detection_is_public，idx_detection_reviewed_by，idx_detection_created_at。",
    },
    {
        "name": "evidence_matches",
        "title": "有效证据匹配表",
        "purpose": "保存通过证据仲裁并实际参与结果展示/评分的证据快照；完整候选池和排除理由保存在 detection_records.analysis_payload。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "证据主键"],
            ["detection_id", "BIGINT", "否", "FK→detection_records.id，CASCADE，IDX", "所属检测"],
            ["knowledge_id", "BIGINT", "是", "FK→knowledge_items.id，SET NULL，IDX", "对应本地知识；联网证据可为空"],
            ["title", "VARCHAR(255)", "否", "—", "证据标题快照"],
            ["summary", "TEXT", "是", "—", "证据摘要快照"],
            ["source_name", "VARCHAR(100)", "是", "—", "来源名称"],
            ["similarity_score", "DECIMAL(6,4)", "否", "—", "检索相似度"],
            ["rank_order", "INT", "否", "—", "仲裁后的展示顺序"],
            ["url", "VARCHAR(500)", "是", "迁移字段", "联网或知识来源地址"],
            ["origin", "VARCHAR(20)", "否", "IDX，默认 knowledge", "knowledge/web 来源标识"],
            ["created_at", "DATETIME", "否", "默认当前时间", "创建时间"],
            ["updated_at", "DATETIME", "否", "自动更新", "最后更新时间"],
        ],
        "indexes": "PRIMARY(id)，idx_evidence_detection_id，idx_evidence_knowledge_id，idx_evidence_origin。",
    },
    {
        "name": "prompt_templates",
        "title": "Prompt 模板表",
        "purpose": "保存大语言模型分析模板、启停状态和默认模板标识。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "模板主键"],
            ["name", "VARCHAR(100)", "否", "—", "模板名称"],
            ["type", "VARCHAR(50)", "否", "IDX", "模板业务类型"],
            ["content", "TEXT", "否", "—", "模板正文及占位符"],
            ["is_default", "BOOLEAN", "否", "联合 IDX，默认 0", "同类型默认模板标识"],
            ["status", "VARCHAR(20)", "否", "IDX，默认 enabled", "enabled/disabled"],
            ["created_by", "BIGINT", "是", "FK→users.id，SET NULL，IDX", "创建管理员"],
            ["created_at", "DATETIME", "否", "默认当前时间", "创建时间"],
            ["updated_at", "DATETIME", "否", "自动更新", "最后更新时间"],
        ],
        "indexes": "PRIMARY(id)，idx_prompt_type，idx_prompt_status，idx_prompt_type_default(type,is_default)，idx_prompt_created_by。",
    },
    {
        "name": "reports",
        "title": "检测报告表",
        "purpose": "保存检测报告元数据及 HTML/PDF 文件相对路径。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "报告主键"],
            ["detection_id", "BIGINT", "否", "FK→detection_records.id，CASCADE，UQ", "所属检测；一条检测最多一个报告"],
            ["user_id", "BIGINT", "是", "FK→users.id，SET NULL，IDX", "报告所有者"],
            ["report_title", "VARCHAR(255)", "否", "—", "报告标题"],
            ["html_path", "VARCHAR(500)", "是", "—", "HTML 相对路径"],
            ["pdf_path", "VARCHAR(500)", "是", "—", "PDF 相对路径"],
            ["created_at", "DATETIME", "否", "默认当前时间", "首次生成时间"],
            ["updated_at", "DATETIME", "否", "自动更新", "重新生成时间"],
        ],
        "indexes": "PRIMARY(id)，UNIQUE(detection_id)，idx_report_user_id。",
    },
    {
        "name": "system_logs",
        "title": "系统审计日志表",
        "purpose": "记录管理员关键业务操作，用于追责、复核和问题定位。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "日志主键"],
            ["user_id", "BIGINT", "是", "FK→users.id，SET NULL，IDX", "操作用户"],
            ["action", "VARCHAR(100)", "否", "IDX", "动作代码"],
            ["module", "VARCHAR(100)", "否", "IDX", "业务模块"],
            ["description", "TEXT", "是", "—", "操作摘要，不记录秘密"],
            ["ip_address", "VARCHAR(50)", "是", "—", "请求来源地址"],
            ["created_at", "DATETIME", "否", "IDX，默认当前时间", "发生时间"],
        ],
        "indexes": "PRIMARY(id)，idx_system_log_user_id，idx_system_log_module，idx_system_log_action，idx_system_log_created_at。",
    },
    {
        "name": "crawl_tasks",
        "title": "定时采集任务执行表",
        "purpose": "记录每次后台新闻采集任务的输入、统计结果、错误摘要和执行状态。",
        "fields": [
            ["id", "BIGINT", "否", "PK，自增", "任务执行主键"],
            ["job_name", "VARCHAR(100)", "否", "IDX", "任务名称"],
            ["search_query", "VARCHAR(500)", "否", "—", "实际搜索词"],
            ["freshness", "VARCHAR(50)", "否", "—", "搜索时间范围"],
            ["total_found", "INT", "否", "默认 0", "搜索返回数"],
            ["new_added", "INT", "否", "默认 0", "新增知识数"],
            ["duplicates", "INT", "否", "默认 0", "去重跳过数"],
            ["fetch_failed", "INT", "否", "默认 0", "全文抓取失败数"],
            ["errors", "TEXT", "是", "—", "错误摘要"],
            ["started_at", "DATETIME", "否", "IDX，默认当前时间", "开始时间"],
            ["finished_at", "DATETIME", "是", "—", "完成时间"],
            ["status", "VARCHAR(20)", "否", "默认 running", "running/success/partial/failed"],
        ],
        "indexes": "PRIMARY(id)，idx_crawl_tasks_job，idx_crawl_tasks_time。",
    },
]


API_GROUPS = [
    ["健康检查", "GET /api/health", "公开", "服务存活与基本状态"],
    ["认证", "POST /api/auth/register；POST /api/auth/login；GET /api/auth/me", "注册/公开；me 登录", "注册、登录、当前用户"],
    ["新闻检测", "POST /api/detect/news；POST /api/detect/extract-preview", "游客/用户", "文本检测、URL 安全提取预览"],
    ["检测记录", "GET /api/detect/history；GET /api/detect/{id}；POST /api/detect/{id}/re-evaluate", "登录且本人；管理员可扩展", "历史、详情、重新评估"],
    ["独立 RAG", "POST /api/rag/search", "登录", "知识库语义检索"],
    ["公开高风险", "GET /api/high-risk/public；/ranking；/keywords；/category-distribution", "公开", "仅已审核且允许公开的数据"],
    ["报告", "POST /api/report/generate/{detection_id}；GET /api/report/download/{report_id}", "登录且有权访问记录", "生成与下载 PDF"],
    ["管理员检测", "GET/DELETE /api/admin/detections[/{id}]", "管理员", "全局检测查询、详情与删除"],
    ["管理员知识库", "GET/POST/PUT/DELETE /api/admin/knowledge；rebuild-index；vectorize", "管理员", "知识 CRUD 与向量维护"],
    ["管理员 Prompt", "GET/POST/PUT/DELETE /api/admin/prompts；enable/disable/set-default", "管理员", "模板生命周期管理"],
    ["管理员高风险", "GET /api/admin/high-risk；PUT review/public/remark", "管理员", "复核、公开控制、备注"],
    ["管理员用户", "GET /api/admin/users；enable/disable/role；用户检测记录", "管理员", "用户查询、启停与角色维护"],
    ["管理员报告", "GET /api/admin/reports；详情；download", "管理员", "全局报告管理"],
    ["统计", "GET /api/admin/statistics/*", "管理员", "概览、趋势、风险/类别分布、关键词、活跃度、知识概览"],
    ["审计日志", "GET /api/admin/logs", "管理员", "分页筛选系统审计日志"],
]


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=12, bold=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def set_page_number_format(section, fmt="decimal", start=1):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        sect_pr.append(node)
    node.set(qn("w:fmt"), fmt)
    node.set(qn("w:start"), str(start))


def configure_section(section, header_footer=True, page_format="decimal", page_start=1):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.45)
    section.bottom_margin = Cm(2.35)
    section.left_margin = Cm(2.85)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.05)
    section.footer_distance = Cm(1.05)
    if not header_footer:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""
        return
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    r = hp.add_run(f"{DOC_NO}  |  {DOC_TITLE}")
    set_run_font(r, size=8.5, color=GRAY)
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Cm(0)
    r = fp.add_run("受控文档    •    ")
    set_run_font(r, size=8.5, color=GRAY)
    add_field(fp, "PAGE", str(page_start))
    set_page_number_format(section, page_format, page_start)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.8)
    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 14, BLUE, 14, 7),
        ("Heading 3", 12, "24556B", 10, 5),
        ("Heading 4", 11, "365F70", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(11.2)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, indent=True, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.8) if indent else Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, east_asia="黑体", bold=True, color=NAVY)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        r = p.add_run(text)
        set_run_font(r, size=11.2)


def add_numbered(doc, items):
    for index, text in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        r = p.add_run(f"{index}. {text}")
        set_run_font(r, size=11.2)


def add_note(doc, title, text, kind="info"):
    fill, accent = {
        "info": (PALE_BLUE, BLUE),
        "success": (PALE_TEAL, GREEN),
        "warning": (PALE_GOLD, GOLD),
        "danger": (PALE_RED, RED),
    }[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{title}｜")
    set_run_font(r, east_asia="黑体", size=10.5, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, font_size=9.2, header_fill=BLUE, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        if widths:
            cell.width = Cm(widths[i])
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, east_asia="黑体", size=font_size, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        if row_index % 2 == 1:
            for c in row.cells:
                set_cell_shading(c, LIGHT_GRAY)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, east_asia="黑体" if first_col_bold and i == 0 else "宋体", size=font_size, bold=(first_col_bold and i == 0), color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, east_asia="黑体", size=10.5, bold=True, color=NAVY)
    return p


def add_figure(doc, filename, caption, max_width=15.0, max_height=18.0):
    path = FIGURE_DIR / filename
    if not path.exists():
        add_note(doc, "图形缺失", f"未找到 {path}", "warning")
        return
    with Image.open(path) as img:
        ratio = img.width / img.height
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width), height=Cm(height))
    add_caption(doc, caption)


def add_toc(doc):
    add_heading(doc, "目录", 1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "右键选择‘更新域’以生成目录")
    add_note(doc, "目录更新", "文档已设置为打开时更新域。若目录未自动刷新，请在 Word 中按 Ctrl+A 后按 F9。", "info")


def add_cover(doc):
    section = doc.sections[0]
    configure_section(section, header_footer=False)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SOFTWARE ENGINEERING DOCUMENT")
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=9, bold=True, color=TEAL)
    bar = doc.add_table(rows=1, cols=1)
    bar.cell(0, 0).height = Cm(0.16)
    set_cell_shading(bar.cell(0, 0), TEAL)
    bar.autofit = False
    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(f"《{PROJECT_TITLE}》")
    set_run_font(r, east_asia="黑体", latin="Arial", size=20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(DOC_TITLE)
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=30, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("HIGH-LEVEL DESIGN DESCRIPTION")
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=11, bold=True, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(42)
    info = doc.add_table(rows=6, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.style = "Table Grid"
    for i, (label, value) in enumerate([
        ("文档编号", DOC_NO), ("版本号", VERSION), ("配置基线", BASELINE),
        ("文档状态", "正式版"), ("编制单位", "项目开发组"), ("发布日期", DATE_TEXT),
    ]):
        set_cell_shading(info.cell(i, 0), PALE_BLUE)
        for j, text in enumerate((label, value)):
            c = info.cell(i, j)
            c.width = Cm(4.2 if j == 0 else 8.4)
            set_cell_margins(c, top=130, bottom=130)
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run_font(r, east_asia="黑体" if j == 0 else "宋体", size=10.5, bold=(j == 0), color=NAVY if j == 0 else DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("受控文档 · 未经批准不得擅自修改")
    set_run_font(r, size=9.5, color=GRAY)


def add_front_matter(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, header_footer=True, page_format="lowerRoman", page_start=1)
    heading = add_heading(doc, "文档控制信息", 1)
    heading.paragraph_format.page_break_before = False
    add_body(doc, "本文件为项目概要设计阶段的受控技术文件，用于统一系统总体架构、模块边界、接口、运行机制、数据存储和安全设计。设计变更应同步修改本文档、数据库迁移和接口说明，并履行评审与版本控制。")
    add_heading(doc, "文档审批", 2)
    add_table(doc, ["角色", "姓名/单位", "职责", "签字", "日期"], [
        ["编制", "项目开发组", "撰写、核对代码事实", "", DATE_TEXT],
        ["审核", "", "架构、数据库与安全设计审查", "", ""],
        ["批准", "", "批准进入详细设计/实现基线", "", ""],
    ], [2.2, 3.2, 5.2, 2.4, 2.7])
    add_heading(doc, "修订记录", 2)
    add_table(doc, ["版本", "日期", "修订内容", "修订人", "状态"], [
        [VERSION, DATE_TEXT, "首次发布；依据当前代码、Alembic 迁移、前后端路由与项目设计材料形成概要设计基线。", "项目开发组", "正式"],
    ], [2.0, 2.8, 8.3, 2.5, 2.1])
    add_heading(doc, "分发范围", 2)
    add_table(doc, ["序号", "接收方", "用途", "介质"], [
        ["1", "项目负责人/指导教师", "设计评审与阶段验收", "电子版"],
        ["2", "开发与测试人员", "详细设计、实现、测试依据", "电子版"],
        ["3", "配置管理员", "基线归档与变更控制", "电子版"],
    ], [1.6, 4.4, 7.6, 3.2])
    heading = add_heading(doc, "摘要", 1)
    heading.paragraph_format.page_break_before = False
    add_body(doc, f"{SYSTEM_NAME}采用 B/S 架构和前后端分离模式，以 Vue 3 构建用户端与管理端，以 FastAPI 提供 REST 服务；MySQL 保存业务事实，Chroma 保存知识条目的派生向量索引，报告文件保存于受控目录。新闻鉴别主流程由文本规范化、本地 RAG 召回、按需联网补充、DeepSeek 证据仲裁、规则评分、综合评分、风险分级和结果持久化组成。")
    add_body(doc, "本文档依据 GB/T 8567—2006 的概要设计文档思想组织，覆盖引言、总体设计、系统结构、模块、接口、运行、数据结构、异常处理、安全、性能、部署和维护设计；数据库章节给出 8 张关系表、Chroma 集合、文件存储以及一致性与备份策略。")
    add_note(doc, "适用边界", "系统输出是新闻可信度辅助评估，不构成行政、司法或权威事实裁决；外部模型与搜索服务的可用性和内容质量具有不确定性。", "warning")
    add_toc(doc)


def add_main_content(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, header_footer=True, page_format="decimal", page_start=1)

    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 标识", 2)
    add_table(doc, ["项目", "内容"], [
        ["软件名称", f"{SYSTEM_NAME}——{PROJECT_TITLE}"],
        ["文档名称", DOC_TITLE], ["文档编号", DOC_NO], ["版本/基线", f"{VERSION} / {BASELINE}"],
        ["软件形态", "Web 应用（Vue 3 前端 + FastAPI 后端 + MySQL/Chroma 存储）"],
        ["适用阶段", "概要设计评审、详细设计、编码、测试、部署和验收"],
    ], [4.1, 13.0], first_col_bold=True)
    add_heading(doc, "1.2 编写目的", 2)
    add_body(doc, "本文档把已确认的软件需求转化为可实现、可测试、可维护的总体技术方案，明确各子系统的职责边界、依赖关系、数据组织、外部接口、运行流程、异常恢复和安全控制。文档既作为开发人员编写详细设计和代码的上位依据，也作为测试人员设计集成测试、数据库测试和安全测试的输入。")
    add_heading(doc, "1.3 项目背景与范围", 2)
    add_body(doc, "网络新闻传播速度快、来源复杂、标题与正文可能包含夸张或误导表达，单靠关键词或一次大模型问答难以给出可追溯结论。本系统通过检索增强生成（RAG）把待鉴别新闻与本地知识和联网候选证据联系起来，再由大语言模型进行受约束的证据仲裁，并与确定性规则共同形成可信度评分。")
    add_body(doc, "本期范围包括游客/用户新闻检测、URL 内容预览、注册登录、个人历史、重新评估、PDF 报告、公开高风险信息，以及管理员用户、检测、知识、Prompt、高风险、报告、统计和日志管理；后台定时采集作为配置化运行能力。系统不承诺自动给出绝对真伪裁决，不包含跨平台舆情监控、人工事实核查工作台、模型训练和多租户计费。")
    add_heading(doc, "1.4 读者对象", 2)
    add_bullets(doc, ["项目负责人、指导教师和评审人员；", "前端、后端、数据库和测试人员；", "部署、运维和配置管理人员；", "后续维护与二次开发人员。"])
    add_heading(doc, "1.5 术语与缩略语", 2)
    add_table(doc, ["术语", "英文/缩写", "说明"], [
        ["检索增强生成", "RAG", "先检索外部知识，再以证据约束大模型生成或分析。"],
        ["大语言模型", "LLM", "用于语义分析、证据仲裁和结构化结果生成的模型。"],
        ["向量嵌入", "Embedding", "把文本映射为固定维度数值向量，用于语义相似检索。"],
        ["证据仲裁", "Evidence Arbitration", "模型对候选证据逐条判定相关性、质量、立场和取舍。"],
        ["业务事实源", "Source of Truth", "对业务状态具有最终解释权的数据存储；本系统为 MySQL。"],
        ["派生索引", "Derived Index", "由知识条目生成、可重建的 Chroma 向量数据。"],
        ["访问令牌", "JWT", "登录后用于身份与角色校验的 Bearer 令牌。"],
        ["服务端请求伪造", "SSRF", "攻击者诱导服务端访问内网或受限地址的风险。"],
    ], [3.2, 3.3, 10.6])
    add_heading(doc, "1.6 参考资料与编制依据", 2)
    add_table(doc, ["编号", "资料", "用途"], [
        ["[1]", "GB/T 8567—2006《计算机软件文档编制规范》", "文档内容组织与配置管理参考"],
        ["[2]", "GB/T 11457—2006《信息技术 软件工程术语》", "术语使用参考"],
        ["[3]", "GB/T 25000.10—2016《系统与软件工程 系统与软件质量要求和评价（SQuaRE）第10部分》", "质量特性参考"],
        ["[4]", "项目需求分析书（NCE-SRS-001）", "需求基线"],
        ["[5]", "Alembic 0001～0007 迁移、SQLAlchemy 模型", "数据库物理设计事实依据"],
        ["[6]", "FastAPI 路由、Service/CRUD、Vue Router 与页面代码", "接口、模块与运行设计事实依据"],
    ], [1.5, 10.8, 4.8])

    add_heading(doc, "2 总体设计", 1)
    add_heading(doc, "2.1 需求规定概述", 2)
    add_body(doc, "系统围绕‘输入新闻—检索证据—仲裁与评分—解释与报告—治理与复核’形成闭环。游客可执行一次性检测和浏览已公开高风险信息；注册用户可保存个人历史、重新评估和生成报告；管理员负责业务数据治理和系统监督。")
    add_table(doc, ["角色", "主要能力", "数据范围"], [
        ["游客", "浏览公开页、提交文本/URL 预览与检测", "检测记录 user_id 为空；无个人历史与报告权限"],
        ["注册用户", "检测、历史、详情、重新评估、报告、独立 RAG", "仅本人数据"],
        ["管理员", "普通用户能力 + 用户/检测/知识/Prompt/高风险/报告/统计/日志管理", "经管理员接口访问全局数据"],
        ["调度器", "按 Cron 触发新闻搜索、抓取、去重、入库与向量化", "受配置开关控制的后台任务"],
    ], [2.8, 7.0, 7.3])
    add_heading(doc, "2.2 设计目标", 2)
    add_bullets(doc, [
        "正确性：评分、风险分级和公开条件由服务端统一执行，前端不复制决定性规则。",
        "可解释性：保存候选证据、有效证据、排除证据、仲裁状态、评分分量和理由。",
        "可靠性：外部搜索失败不阻断本地路径；LLM 失败时退化为确定性规则评分。",
        "安全性：后端强制鉴权、SSRF 默认拒绝、秘密配置外置、报告目录越界检查。",
        "可维护性：API/Schema/Service/CRUD/Model 分层，外部服务通过专用适配器封装。",
        "可恢复性：MySQL 为事实源，Chroma 可全量重建，迁移由 Alembic 顺序管理。",
    ])
    add_heading(doc, "2.3 运行环境", 2)
    add_table(doc, ["类别", "建议/基线环境", "说明"], [
        ["客户端", "现代 Chromium/Edge/Firefox 浏览器", "支持 ES 模块、Fetch/XHR、CSS Grid/Flex"],
        ["前端", "Node.js + Vite 6 + Vue 3.5", "开发端口默认 5173；生产构建为静态资源"],
        ["后端", "Python 3.11+、FastAPI 0.111、Uvicorn 0.30", "默认端口 8000"],
        ["关系数据库", "MySQL 8.x，InnoDB，utf8mb4", "业务事实源；开发测试可用 SQLite 变体"],
        ["向量数据库", "Chroma 1.x，本地持久化", "knowledge_items 集合，余弦空间"],
        ["外部服务", "DeepSeek Chat、DashScope text-embedding-v4、博查搜索", "API Key 由环境变量提供"],
        ["报告", "Jinja2 + xhtml2pdf", "生成 HTML/PDF 到源码目录外"],
    ], [3.0, 7.0, 7.1])
    add_heading(doc, "2.4 设计约束与假设", 2)
    add_bullets(doc, [
        "部署环境能够访问 MySQL，并按需访问 DeepSeek、DashScope 和博查 API；无外网时智能能力可能降级。",
        "MySQL 与 Chroma 不提供跨存储分布式事务，必须通过同步状态、补偿和重建处理一致性。",
        "新闻网页结构不可控，URL 提取结果须由用户确认后再启动正式检测。",
        "外部证据和模型结论可能有偏差，系统必须保留免责声明及人工复核路径。",
        "生产环境不得使用占位 SECRET_KEY、宽泛 CORS 或内网抓取开关。",
    ])
    add_heading(doc, "2.5 设计原则", 2)
    add_table(doc, ["原则", "设计落实"], [
        ["关注点分离", "展示、接口、业务、持久化、外部服务和基础能力分层。"],
        ["业务事实单一来源", "MySQL 保存业务记录；Chroma 与报告文件均可由事实数据追踪。"],
        ["默认拒绝与最小权限", "公开、登录、管理员三类后端依赖明确划分。"],
        ["先仲裁后计分", "未经契约校验的候选证据不得参与有效证据评分。"],
        ["故障显式化", "记录 failed/delete_failed、arbitration_status、quality_status 等状态。"],
        ["配置与代码分离", "密钥、连接、模型、目录、限流和调度均通过环境变量。"],
    ], [4.2, 12.9])

    add_heading(doc, "3 系统总体架构设计", 1)
    add_heading(doc, "3.1 架构风格", 2)
    add_body(doc, "系统采用浏览器/服务器（B/S）架构、前后端分离和面向服务的分层设计。浏览器只承担交互与展示；FastAPI 是业务与安全边界；MySQL、Chroma 和报告目录分别承担结构化事实、向量索引和派生文件存储；外部智能服务通过专用服务类接入。")
    add_figure(doc, "fig4-1-system-architecture.png", "图 3-1 系统总体架构图", max_width=15.2, max_height=12.0)
    add_heading(doc, "3.2 前端架构", 2)
    add_body(doc, "前端以 Vue 3 单页应用实现，Vue Router 负责公开、登录和管理员路由，Pinia 保存用户会话，Axios 统一访问 /api 接口，Element Plus 提供交互组件，ECharts 展示统计数据。路由守卫用于改善体验，但不作为最终授权依据；所有敏感操作由后端再次校验。")
    add_heading(doc, "3.3 后端分层架构", 2)
    add_figure(doc, "fig4-2-backend-layered-architecture.png", "图 3-2 后端分层架构图", max_width=13.2, max_height=13.0)
    add_table(doc, ["层次", "目录/组件", "职责", "禁止事项"], [
        ["API", "app/api/v1", "路由、依赖注入、参数解释、HTTP 响应", "不编写复杂 SQL 和模型调用流程"],
        ["Schema", "app/schemas", "Pydantic 请求/响应契约", "不包含持久化副作用"],
        ["Service", "app/services", "检测编排、RAG、LLM、评分、报告、统计", "不依赖具体页面"],
        ["CRUD", "app/crud", "SQLAlchemy 查询、分页与事务落库", "不决定业务风险等级"],
        ["Model", "app/models", "表、主外键、关系和索引映射", "不调用外部 API"],
        ["Core/DB", "app/core、app/db", "配置、安全、限流、调度、会话、迁移", "不承载页面业务"],
    ], [2.4, 3.4, 6.6, 5.0])
    add_heading(doc, "3.4 数据与外部服务架构", 2)
    add_body(doc, "MySQL 是业务事实源；Chroma 中的向量、文档和元数据由 knowledge_items 派生；报告文件由 detection_records 与 reports 派生。DeepSeek 负责受证据约束的语义分析，DashScope 负责正式语义嵌入，博查只在证据不足且允许联网时补充候选。")
    add_note(doc, "关键边界", "Chroma 命中后必须依据 knowledge_id 回查 MySQL；模型输出必须经过服务端契约校验和分值归一化，不能直接作为最终业务状态。", "success")
    add_heading(doc, "3.5 功能模块划分", 2)
    add_figure(doc, "fig4-3-functional-modules.png", "图 3-3 系统功能模块图", max_width=15.5, max_height=7.0)
    add_table(doc, ["模块", "输入", "核心处理", "输出/存储"], [
        ["认证与个人中心", "账号、密码、JWT", "注册、散列、登录、角色/状态校验", "用户信息、令牌、会话状态"],
        ["新闻可信度评估", "标题、正文、来源、联网开关", "清洗、检索、仲裁、规则/综合评分", "分数、等级、理由、证据、记录"],
        ["记录与报告", "记录 ID、筛选条件", "历史查询、详情、重评、模板渲染", "列表、详情、PDF"],
        ["知识与采集", "知识条目、采集配置", "CRUD、Embedding、同步、重建、定时抓取", "MySQL 条目、Chroma 向量、任务日志"],
        ["管理员治理", "管理查询与操作", "用户启停、模板、高风险复核、报告管理", "管理结果、审计日志"],
        ["统计与日志", "时间范围、维度", "聚合查询、分页审计", "ECharts 数据、日志列表"],
    ], [3.2, 4.7, 6.2, 4.7])
    add_heading(doc, "3.6 部署拓扑", 2)
    add_body(doc, "开发环境由 Vite 开发服务器代理 /api 到 Uvicorn；生产环境可由同一反向代理提供前端静态资源并转发后端请求。MySQL、Chroma 持久化目录和报告目录应位于受控存储，外部服务经 HTTPS 访问。若部署多个后端实例，应避免每个实例同时启动相同 APScheduler 任务，并将当前进程内限流替换为共享限流存储。")

    add_heading(doc, "4 模块概要设计", 1)
    modules = [
        ("4.1 用户认证与权限模块", "完成注册、登录和当前用户查询。密码使用 bcrypt 散列；JWT 包含用户标识、角色和到期时间；受保护接口在每次请求中回查账户状态，禁用账号即使令牌未到期也会被拒绝。", ["公开依赖：健康检查、注册、登录、公开高风险、游客检测。", "登录依赖：个人历史、详情、重评、RAG、报告。", "管理员依赖：/api/admin 下全部管理接口。"]),
        ("4.2 新闻输入与 URL 预览模块", "支持手工输入标题/正文，或粘贴 HTTP(S) URL 提取预览。抓取器拒绝本机、私网、链路本地和保留地址，并逐跳检查重定向；正文提取后由用户确认和修订，预览动作不创建检测记录。", ["标题规范化后最长 255 字符；正文最长 12000 字符。", "抓取设置超时与最大响应体，拒绝非 HTTP(S) 协议。", "正式检测入口应用滑动窗口限流，默认每 60 秒 3 次。"]),
        ("4.3 RAG 与证据治理模块", "把标题和正文构造成查询文本，调用 Embedding 后在 Chroma knowledge_items 集合按余弦距离召回 Top-10；本地证据不足且本次请求允许联网时，调用博查补充候选。候选被赋予稳定 candidate_id，并以来源中立的确定性随机顺序送入模型。", ["模型必须把每个 candidate_id 放入 ranked_evidence 或 rejected_evidence。", "服务端校验唯一性、完整性、0～100 分值、stance 枚举和非空理由。", "只有通过校验的 ranked 证据进入有效集合；原始候选与排除理由写入 analysis_payload。"]),
        ("4.4 评分与风险分级模块", "规则引擎从缺少来源、强情绪、绝对化表达和高相似冲突证据等特征扣分；模型给出可信度分量和证据质量；服务端按证据可用性选择公式并映射四级风险。", ["有有效证据：LLM 50% + 证据质量 30% + 规则 20%。", "无有效证据但 LLM 正常：LLM 60% + 规则 40%。", "LLM 失败：未仲裁候选不得计分，最终分退化为规则分。"]),
        ("4.5 检测记录与报告模块", "检测成功后保存输入快照、评分分量、结论、有效证据和分析载荷。登录用户可查询本人历史和重新评估；报告由已保存记录渲染 HTML 并转换为 PDF，不重复调用模型。", ["reports.detection_id 唯一，重复生成替换派生文件。", "普通用户只能操作本人记录；管理员通过独立接口访问全局数据。", "报告下载前解析并校验目标路径位于 REPORT_DIR 内。"]),
        ("4.6 知识库与采集模块", "管理员维护知识条目并触发向量 upsert、单条重试或全量重建。定时采集根据配置调用搜索、抓取正文、去重、写 MySQL、按需同步向量，并把一次执行的统计与错误写入 crawl_tasks。", ["知识原文落库成功但向量失败时保留记录并标记 failed。", "删除采用应用层补偿：向量与关系数据任一失败时保留可恢复状态。", "切换 Embedding 提供商或维度后必须重建全部索引。"]),
        ("4.7 管理与统计模块", "管理员可维护用户状态、检测记录、知识、Prompt、高风险记录和报告，查看统计与审计日志。高风险公开必须同时满足 is_high_risk=true、review_status=approved、is_public=true。", ["不能禁用或降级当前管理员，也不能移除最后一个可用管理员。", "Prompt 启用前校验必需占位符和结构；同类型仅一个默认模板。", "统计数据直接聚合真实业务表，日期范围最大 366 天。"]),
    ]
    for heading, body, bullets in modules:
        add_heading(doc, heading, 2)
        add_body(doc, body)
        add_bullets(doc, bullets)

    add_heading(doc, "5 接口设计", 1)
    add_heading(doc, "5.1 用户界面接口", 2)
    add_body(doc, "前端分为用户布局与管理员布局。公开页面包括首页、登录、注册、检测、结果和高风险展示；历史、个人中心要求登录；/admin 下页面要求管理员角色。结果页支持游客承接当次检测结果，但持久化详情与报告仍受后端权限控制。")
    add_table(doc, ["界面域", "路由", "说明"], [
        ["公开与用户端", "/、/login、/register、/detect、/result/:id、/high-risk", "公开浏览、输入、结果和公开高风险"],
        ["登录用户", "/history、/profile", "个人历史与账号信息"],
        ["管理员", "/admin/dashboard、users、detections、knowledge、prompts、high-risk、statistics、reports、logs", "后台治理与统计"],
    ], [3.2, 6.6, 8.5])
    add_heading(doc, "5.2 REST API 总体约定", 2)
    add_bullets(doc, [
        "统一前缀为 /api；JSON 请求使用 UTF-8；报告下载返回文件流。",
        "成功响应统一包含 code、message、data；异常响应保持相同外层语义。",
        "分页参数通常为 page、page_size，page_size 最大 100；日期筛选使用 ISO 日期/时间。",
        "登录请求在 Authorization: Bearer <JWT> 中传递令牌。",
        "HTTP 401 表示未认证或令牌无效，403 表示身份有效但权限不足，409 表示业务状态冲突，422 表示参数校验失败。",
    ])
    add_heading(doc, "5.3 API 接口分组", 2)
    add_table(doc, ["接口域", "方法与路径", "权限", "职责"], API_GROUPS, [2.6, 8.4, 3.1, 4.2], font_size=8.3)
    add_heading(doc, "5.4 外部接口", 2)
    add_table(doc, ["外部系统", "调用方向", "输入", "输出", "失败策略"], [
        ["DeepSeek Chat", "后端→HTTPS API", "系统指令、新闻、候选证据、Prompt", "契约版本 2.0 JSON", "校验失败可重试；最终退化为规则评分"],
        ["DashScope Embedding", "后端→HTTPS API", "规范化查询/知识文本", "默认 1024 维向量", "检索失败返回明确错误；知识同步标记 failed"],
        ["博查搜索", "后端→HTTPS API", "查询词、freshness、count", "候选标题、摘要、URL", "保留本地证据继续分析"],
        ["新闻网页", "后端→HTTP(S)", "经安全校验的 URL", "标题、正文、发布时间", "超时、体积或地址不合法时拒绝并提示"],
        ["MySQL", "后端→SQL", "事务查询/写入", "结构化业务数据", "回滚当前事务并返回统一错误"],
        ["Chroma", "后端→本地客户端", "向量、过滤条件、Top-K", "文档、元数据、距离", "清理缓存句柄并重试一次；可重建"],
    ], [3.2, 3.2, 4.7, 4.8, 4.6], font_size=8.6)
    add_heading(doc, "5.5 关键数据契约", 2)
    add_body(doc, "检测请求至少包含 title 与 content，可选 category、source_name、source_url、publish_time 和 enable_web_search。检测响应包含 detection_id、四类分值、risk_level、judgement_result、reason、risk_points、keywords、候选/有效/排除证据、similar_news、suggestion、agent_steps、web_search 状态、evidence_quality、arbitration_status 和免责声明。")
    add_note(doc, "接口演进", "新增响应字段应保持向后兼容；删除/改名字段必须先修改需求与接口基线并提供迁移说明。数据库字段不应直接无筛选地序列化给公开接口。", "warning")

    add_heading(doc, "6 运行设计", 1)
    add_heading(doc, "6.1 系统启动与关闭", 2)
    add_numbered(doc, [
        "加载 backend/.env，解析数据库、密钥、外部服务、目录、限流与调度配置。",
        "校验 SECRET_KEY：禁止为空或占位值，生产环境要求足够长度和字符多样性。",
        "执行 Alembic 迁移到 head，确保 0001～0007 顺序生效。",
        "启动 FastAPI/Uvicorn，注册统一异常处理、CORS、路由和应用生命周期。",
        "在 CRAWL_ENABLED 且配置满足条件时启动 APScheduler 周期任务。",
        "前端开发环境启动 Vite；生产环境加载构建后的静态资源。",
        "关闭时停止调度器，完成在途请求，并释放数据库会话和外部连接句柄。",
    ])
    add_heading(doc, "6.2 新闻检测运行流程", 2)
    add_figure(doc, "fig4-4-detection-overall-flow.png", "图 6-1 新闻可信度评估总体流程", max_width=9.4, max_height=17.8)
    add_numbered(doc, [
        "校验限流、可选身份、标题和正文；执行文本清洗与长度约束。",
        "提取关键词并构造 RAG 查询文本，生成向量后召回本地 Top-10。",
        "根据 Top-1 相似度和有效候选数量判断证据是否不足；满足条件且用户允许时联网搜索。",
        "合并本地/联网候选，分配 candidate_id，并按与来源、检索排名无关的稳定散列顺序排列。",
        "调用 DeepSeek 输出证据仲裁和新闻分析 JSON；服务端校验所有候选被完整覆盖。",
        "对有效证据重算质量、执行规则评分并按分支公式得到最终分和风险等级。",
        "在同一关系数据库事务中保存检测记录及有效证据；完整分析审计信息写入 analysis_payload。",
        "返回结构化结果；需要报告时再根据已保存记录生成，不重复分析。",
    ])
    add_heading(doc, "6.3 本地检索与联网补充", 2)
    add_figure(doc, "fig4-5-local-retrieval-web-supplement.png", "图 6-2 本地检索与联网补充流程", max_width=9.8, max_height=16.6)
    add_body(doc, "联网补充的默认触发规则为：本地无证据；或 Top-1 相似度低于 0.45；或 Top-1 低于 0.60 且相似度不低于 0.30 的有效候选少于 3 条。该规则用于决定是否扩大证据面，不直接决定新闻真伪。")
    add_heading(doc, "6.4 知识同步运行流程", 2)
    add_numbered(doc, [
        "先在 MySQL 创建/更新知识记录并标记 pending。",
        "拼接标题、正文、类别、摘要、关键词、事实标签与辟谣说明作为向量文本。",
        "调用 Embedding，向 Chroma upsert knowledge:<id>，成功后回写 vector_id 与 synced。",
        "失败时保留 MySQL 记录，写入 failed 与错误摘要；管理员可单条 vectorize 或全量 rebuild-index。",
        "删除时协调 Chroma 与 MySQL；若任一步失败，记录 delete_failed 并支持补偿恢复。",
    ])
    add_heading(doc, "6.5 报告生成流程", 2)
    add_body(doc, "服务校验当前用户对 detection_id 的访问权，读取检测记录、有效证据和用户信息，以 Jinja2 渲染 HTML，再由 xhtml2pdf 生成 PDF。文件写入 REPORT_DIR，数据库只保存相对路径；同一 detection_id 重复生成时替换派生文件并更新时间。")
    add_heading(doc, "6.6 并发、事务与幂等", 2)
    add_table(doc, ["场景", "设计"], [
        ["请求级数据库会话", "每个 API 请求获得独立 SQLAlchemy Session；写失败 rollback，完成后关闭。"],
        ["检测记录 + 有效证据", "通过 ORM 关系在一次提交中持久化；删除检测时证据和报告级联删除。"],
        ["报告重复生成", "detection_id 唯一约束保证一对一，服务更新既有记录和文件。"],
        ["知识向量 upsert", "vector_id 使用 knowledge:<id>，重复执行覆盖同一文档，具备幂等语义。"],
        ["定时采集", "按来源/内容去重并记录一次执行统计；多实例部署需单独选举调度实例。"],
        ["限流", "当前为进程内滑动窗口；单实例有效，多实例生产需共享存储。"],
    ], [4.2, 12.9])

    add_heading(doc, "7 数据库与数据存储设计", 1)
    add_heading(doc, "7.1 设计目标与选型", 2)
    add_body(doc, "关系数据库选用 MySQL，存储用户、知识、检测、证据、模板、报告、日志和采集任务等强关系业务数据；InnoDB 提供事务和外键，utf8mb4 支持中文、Emoji 和多语言文本。Chroma 用于余弦语义检索，不承担业务状态；报告目录保存可重新生成的 HTML/PDF 文件。")
    add_heading(doc, "7.2 概念结构设计（E-R）", 2)
    add_figure(doc, "fig4-7-database-er.png", "图 7-1 数据库实体关系图", max_width=16.0, max_height=9.2)
    add_table(doc, ["关系", "基数", "外键/约束", "删除策略"], [
        ["用户—检测记录", "1:N（检测可无用户）", "detection_records.user_id", "用户删除置空"],
        ["管理员—复核记录", "1:N（可空）", "detection_records.reviewed_by", "管理员删除置空"],
        ["检测—有效证据", "1:N", "evidence_matches.detection_id", "检测删除级联"],
        ["知识—有效证据", "1:N（证据可为联网来源）", "evidence_matches.knowledge_id", "知识删除置空并保留快照"],
        ["检测—报告", "1:0..1", "reports.detection_id UNIQUE", "检测删除级联"],
        ["用户—报告", "1:N（可空）", "reports.user_id", "用户删除置空"],
        ["用户—Prompt/日志", "1:N（可空）", "created_by/user_id", "用户删除置空"],
    ], [3.1, 3.5, 6.0, 4.3])
    add_heading(doc, "7.3 逻辑数据模型", 2)
    add_table(doc, ["实体", "主键", "核心业务字段", "主要关系"], [
        ["users", "id", "username、password_hash、role、status", "提交检测、创建模板、复核、拥有报告、产生日志"],
        ["knowledge_items", "id", "title、content、truth_label、vector_sync_status", "被有效证据可选引用，映射 Chroma 向量"],
        ["detection_records", "id", "输入、四类分值、结论、analysis_payload、审核字段", "从属用户；拥有证据和可选报告"],
        ["evidence_matches", "id", "证据快照、similarity_score、rank_order", "从属检测；可引用知识"],
        ["prompt_templates", "id", "type、content、status、is_default", "可选关联创建管理员"],
        ["reports", "id", "report_title、html_path、pdf_path", "一对一关联检测，可选关联用户"],
        ["system_logs", "id", "action、module、description、ip_address", "可选关联用户"],
        ["crawl_tasks", "id", "query、统计字段、status、时间", "独立执行日志"],
    ], [3.0, 2.2, 7.6, 4.2], font_size=8.8)
    add_heading(doc, "7.4 MySQL 物理表设计", 2)
    add_body(doc, "以下字段以 Alembic 0001～0007 迁移和当前模型为设计基线。BIGINT 主键用于满足长期增长；SQLite 测试环境对部分主键使用 Integer 变体以保持自增行为。")
    for index, spec in enumerate(DB_TABLES, 1):
        add_heading(doc, f"7.4.{index} {spec['title']}（{spec['name']}）", 3)
        add_body(doc, str(spec["purpose"]))
        add_table(doc, ["字段名", "数据类型", "可空", "键/约束", "说明"], spec["fields"], [3.0, 3.0, 1.4, 4.4, 6.1], font_size=7.9)
        add_body(doc, f"索引设计：{spec['indexes']}", indent=False, bold_prefix="索引设计：")
    add_heading(doc, "7.5 索引、约束与数据完整性", 2)
    add_bullets(doc, [
        "实体完整性：所有业务表使用自增主键；reports.detection_id 使用唯一约束保证一对一。",
        "参照完整性：历史主体删除时多采用 SET NULL 保留业务记录；检测删除时证据和报告 CASCADE。",
        "域完整性：角色、状态、风险等级、审核状态、向量状态由 Schema/Service 枚举校验；分值归一化到 0～100。",
        "检索性能：对用户、风险、公开/审核状态、创建时间、知识分类与同步状态建立索引。",
        "审计完整性：候选证据与模型契约状态以 JSON 文本快照保存，避免后续知识变更抹除当时上下文。",
    ])
    add_heading(doc, "7.6 Chroma 向量库设计", 2)
    add_table(doc, ["项目", "设计"], [
        ["集合名称", "knowledge_items"], ["距离空间", "cosine（余弦距离）"],
        ["文档 ID", "knowledge:<knowledge_items.id>"],
        ["Embedding", "正式默认 DashScope text-embedding-v4，1024 维；hash 仅为本地演示回退"],
        ["document", "标题、正文、类别、摘要、关键词、truth_label、辟谣说明的规范化拼接"],
        ["metadata", "knowledge_id、title、category、truth_label、source_name、risk_level"],
        ["查询", "Top-K 1～50，检测默认 10；可按 category/truth_label/risk_level 过滤"],
        ["一致性", "命中后回查 MySQL；无对应条目则跳过；向量集合可全量重建"],
    ], [4.1, 13.0])
    add_heading(doc, "7.7 MySQL、Chroma 与文件协作", 2)
    add_figure(doc, "fig4-8-mysql-chroma-collaboration.png", "图 7-2 MySQL—Chroma—报告文件协作与恢复流程", max_width=10.0, max_height=17.0)
    add_heading(doc, "7.8 数据量、保留与归档", 2)
    add_table(doc, ["数据", "增长特征", "建议保留/归档"], [
        ["用户", "低速增长", "长期保留；账号停用优先于物理删除"],
        ["检测与 analysis_payload", "主要增长源，正文和 JSON 较大", "按业务需要长期保留；超过周期可匿名化后归档"],
        ["有效证据", "每次检测通常 0～10 条", "随检测级联；归档时保持关联"],
        ["知识与向量", "管理员/采集持续增长", "MySQL 长期保留；Chroma 可重建"],
        ["报告文件", "按用户请求生成", "与报告记录一致；定期检查孤儿文件"],
        ["审计日志/采集任务", "持续追加", "按月/季度归档，保留必要审计周期"],
    ], [3.2, 5.7, 8.2])
    add_heading(doc, "7.9 备份、恢复与迁移", 2)
    add_numbered(doc, [
        "每日备份 MySQL，生产环境至少保留全量备份和增量/二进制日志；定期执行恢复演练。",
        "备份 Chroma 持久化目录前停止写入或使用一致性快照；若不可用，可从 synced 知识记录全量重建。",
        "报告目录与 reports 表应在同一备份窗口内保存，恢复后执行路径存在性与越界检查。",
        "数据库结构只通过 Alembic 迁移；部署前备份，升级执行 upgrade head，回滚须验证 downgrade 的数据损失影响。",
        "切换 Embedding 提供商/维度时创建维护窗口，清空旧集合并全量生成，不混用不同维度向量。",
    ])

    add_heading(doc, "8 关键算法与数据结构设计", 1)
    add_heading(doc, "8.1 证据候选与仲裁结构", 2)
    add_body(doc, "候选证据的 candidate_id 使用 kb:<knowledge_id> 或 web:<序号/稳定标识>。送入模型前移除 similarity_score、检索排名等可能造成位置偏见的元数据，并用 SHA-256(seed_material:candidate_id) 排序，使相同输入可复现且顺序不由来源类型决定。")
    add_table(doc, ["仲裁字段", "类型/范围", "语义"], [
        ["candidate_id", "string，必须存在", "精确引用候选"],
        ["relevance_score", "number，0～100", "与新闻核心事实的相关程度"],
        ["quality_score", "number，0～100", "来源、完整性、时间与可验证性质量"],
        ["stance", "support/contradict/neutral", "支持、质疑或中性背景"],
        ["reason", "非空字符串", "仲裁理由"],
        ["ranked/rejected", "两数组完全覆盖候选", "有效与排除集合，不得重复或遗漏"],
    ], [3.2, 4.6, 9.1])
    add_heading(doc, "8.2 证据评分", 2)
    add_body(doc, "有效证据分 evidence_score 为前 10 条有效证据 similarity_score 转换到 0～100 后的算术平均。若无有效证据则为 0。该分数用于展示检索匹配程度；最终有证据公式采用模型输出并经服务端验证的 evidence_quality.score，而不是直接使用检索相似度替代证据质量。")
    add_heading(doc, "8.3 综合评分公式", 2)
    add_table(doc, ["分支", "公式", "说明"], [
        ["LLM 正常且有有效证据", "final = 0.50×LLM + 0.30×EvidenceQuality + 0.20×Rule", "证据质量参与最终分"],
        ["LLM 正常但无有效证据", "final = 0.60×LLM + 0.40×Rule", "理由中追加证据不足提示"],
        ["LLM 调用失败", "final = Rule", "未仲裁候选不计分；显式标记降级"],
    ], [4.3, 7.0, 5.6])
    add_heading(doc, "8.4 风险等级映射", 2)
    add_table(doc, ["分数区间", "风险等级", "结论摘要"], [
        ["80 ≤ score ≤ 100", "可信新闻", "整体可信度较高"],
        ["60 ≤ score < 80", "存疑信息", "存在一定疑点，建议进一步核查"],
        ["40 ≤ score < 60", "疑似谣言", "疑似存在谣言风险"],
        ["0 ≤ score < 40", "高风险谣言", "存在较高谣言风险"],
    ], [4.3, 4.7, 7.9])
    add_figure(doc, "fig4-6-scoring-risk-classification.png", "图 8-1 评分与风险分类流程", max_width=10.2, max_height=16.5)
    add_heading(doc, "8.5 规则评分", 2)
    add_body(doc, "规则分以 100 为基准，对缺少明确来源、强情绪词、绝对化词和与高相似负面证据冲突等命中项扣分，最终限制在 0～100。命中规则以结构化数据合并到风险点，既用于 LLM 故障降级，也用于正常流程中的确定性校正。")

    add_heading(doc, "9 出错处理与恢复设计", 1)
    add_heading(doc, "9.1 错误分类与处理", 2)
    add_table(doc, ["故障场景", "检测方式", "处理/降级", "用户可见结果", "恢复"], [
        ["输入不合法", "Pydantic/文本清洗", "拒绝进入主流程", "422 与字段提示", "修正后重试"],
        ["鉴权失败/越权", "JWT、角色、资源所有权", "拒绝访问", "401/403", "重新登录或申请权限"],
        ["检测限流", "滑动窗口", "拒绝高频请求", "429", "窗口结束后重试"],
        ["URL/重定向不安全", "协议、DNS/IP、每跳校验", "停止抓取", "安全错误提示", "使用公开地址或手工输入"],
        ["Chroma 检索失败", "客户端异常", "清缓存句柄后重试一次；仍失败则终止检索", "明确检索失败", "修复目录/依赖后重试"],
        ["联网搜索失败", "Bocha 异常/超时", "保留本地候选继续", "web_search 状态显示失败/未补充", "外部服务恢复后重评"],
        ["LLM/契约失败", "HTTP/JSON/完整性校验", "重试仲裁；最终只用规则分", "显示降级与不确定性", "模型恢复后重新评估"],
        ["MySQL 写失败", "SQLAlchemy 异常", "rollback，不返回伪成功", "统一 500", "修复数据库后重试"],
        ["向量同步失败", "Embedding/Chroma 异常", "保留知识，标记 failed", "管理端显示错误", "单条重试/全量重建"],
        ["PDF 生成失败", "模板/转换异常", "不改变检测结论，不留下空文件", "500 与生成失败", "修复依赖后重新生成"],
    ], [3.1, 3.4, 4.6, 4.0, 3.2], font_size=7.7)
    add_heading(doc, "9.2 一致性补偿", 2)
    add_body(doc, "知识库写操作跨越 MySQL 与 Chroma，不能宣称原子事务。系统用 vector_sync_status 表示 pending、synced、failed、delete_failed；服务在删除失败时尝试恢复向量并保留错误上下文，管理员可重试或重建。检测记录及其有效证据都位于 MySQL，可在单一事务中保存。")
    add_heading(doc, "9.3 恢复优先级", 2)
    add_numbered(doc, ["保证 MySQL 可用并恢复业务事实；", "校验迁移版本与外键完整性；", "恢复/重建 Chroma 派生索引；", "恢复报告目录并核对相对路径；", "恢复外部 API 配置和调度；", "执行健康检查、抽样检测和权限回归。"])

    add_heading(doc, "10 安全与保密设计", 1)
    add_heading(doc, "10.1 身份认证与授权", 2)
    add_body(doc, "密码只保存 bcrypt 散列；JWT 默认使用 HS256，过期时间由 ACCESS_TOKEN_EXPIRE_MINUTES 配置。后端依赖分别实现可选用户、当前用户和当前管理员校验，并在访问个人记录/报告时检查资源所有权。前端路由守卫不替代后端鉴权。")
    add_heading(doc, "10.2 输入与网络安全", 2)
    add_bullets(doc, [
        "Pydantic 与文本清洗限制类型、枚举和长度；数据库访问使用 ORM 参数绑定。",
        "URL 抓取默认拒绝私网、本机、链路本地和保留地址，重定向每跳重验，限制超时和响应体。",
        "生产 CORS 配置为明确前端域名，不使用 *；仅开放必要 HTTP 方法和头。",
        "检测接口按用户/来源限流，避免滥用外部模型和搜索配额。",
    ])
    add_heading(doc, "10.3 秘密与文件安全", 2)
    add_bullets(doc, [
        "SECRET_KEY、数据库口令、DeepSeek/DashScope/Bocha API Key 只从 .env/安全配置注入，不提交代码库。",
        "启动时拒绝空值和占位 SECRET_KEY；生产密钥至少 32 字符且具有足够字符多样性。",
        "REPORT_DIR 必须位于后端源码目录外；下载时解析路径并确认仍位于允许根目录。",
        "日志不记录密码、完整令牌、API Key 或完整外部响应中的敏感内容。",
    ])
    add_heading(doc, "10.4 LLM 与内容安全", 2)
    add_body(doc, "新闻标题和正文被当作不可信数据，以明确标签边界包裹；系统消息声明新闻中的指令不得改变评分、输出格式、证据或流程。送入模型的候选使用稳定标识，服务端拒绝不存在、重复或遗漏的 candidate_id，并重算关键质量数据。")
    add_heading(doc, "10.5 公开与隐私边界", 2)
    add_body(doc, "公开高风险接口只返回审核通过且允许公开的摘要字段，不返回完整正文、提交用户、管理员备注、内部审核字段或模型调用细节。用户历史与报告按所有权隔离；管理员操作写审计日志。部署方应在隐私政策中说明新闻正文可能传输给第三方模型/搜索服务的目的和范围。")
    add_heading(doc, "10.6 安全责任矩阵", 2)
    add_table(doc, ["控制对象", "前端", "后端", "数据库/存储", "部署方"], [
        ["认证", "保存并附加令牌、退出清理", "签发、校验、回查状态", "保存散列与角色", "保护 SECRET_KEY"],
        ["授权", "隐藏无权入口", "角色与资源所有权强制校验", "外键和最小数据库权限", "账号与权限审计"],
        ["内容", "输入提示与确认", "清洗、长度、Prompt 防注入", "保留必要数据", "隐私告知与保留策略"],
        ["网络", "仅访问配置 API", "SSRF、超时、限流、CORS", "限制网络暴露", "TLS、反向代理、防火墙"],
    ], [2.8, 4.0, 4.9, 4.3, 3.5], font_size=8.5)

    add_heading(doc, "11 性能与可扩展性设计", 1)
    add_heading(doc, "11.1 性能控制", 2)
    add_table(doc, ["环节", "设计"], [
        ["前端", "路由组件懒加载；统计图按页面请求；构建产物压缩。"],
        ["API", "分页默认 20、最大 100；日期范围限制 366 天；统一超时与错误返回。"],
        ["数据库", "按常用筛选字段建立索引；避免返回正文到列表接口；级联关系减少孤儿清理。"],
        ["RAG", "查询文本最长约 8000 字符，检测 Top-K=10，Chroma Top-K 上限 50。"],
        ["外部服务", "按证据不足条件才联网；设置搜索/抓取/模型超时；失败可降级。"],
        ["报告", "按需生成并复用一对一报告记录，不在检测主请求内默认生成。"],
    ], [3.5, 13.6])
    add_heading(doc, "11.2 扩展策略", 2)
    add_bullets(doc, [
        "将前端静态资源部署到 CDN/对象存储，后端保持无页面状态。",
        "后端多实例化前，将限流、任务锁和会话辅助状态迁移到 Redis 等共享基础设施。",
        "把长耗时检测/报告任务迁移为任务队列时，新增 task_id 与状态查询而不破坏现有同步接口。",
        "知识规模增长后可把 Chroma 迁移到服务化向量数据库，保持 Service 层接口和 knowledge_id 回查规则。",
        "统计量增长后可增加汇总表或离线聚合，但指标口径仍以业务表定义为准。",
    ])
    add_heading(doc, "11.3 容量初步估算", 2)
    add_body(doc, "课程设计/小规模试运行场景可按 10 万检测、100 万有效证据上限进行初步容量验证。正文与 analysis_payload 是主要空间来源；若每条检测平均 15～30 KB，则 10 万条约 1.5～3 GB（不含索引、备份和报告）。上线前应使用真实样本测量平均行长、向量占用和报告大小，再修订容量计划。")

    add_heading(doc, "12 部署、运行与运维设计", 1)
    add_heading(doc, "12.1 关键配置项", 2)
    add_table(doc, ["类别", "环境变量", "设计要求"], [
        ["应用", "APP_ENV、PROJECT_NAME、API_PREFIX", "生产明确设置环境名"],
        ["数据库", "DATABASE_URL 或 DATABASE_HOST/PORT/USER/PASSWORD/NAME", "使用最小权限账号和 utf8mb4"],
        ["认证", "SECRET_KEY、ALGORITHM、ACCESS_TOKEN_EXPIRE_MINUTES", "强随机密钥，定期轮换"],
        ["模型", "DEEPSEEK_API_KEY/BASE_URL/MODEL/TIMEOUT", "密钥保密，设置超时"],
        ["向量", "EMBEDDING_PROVIDER、EMBEDDING_DIMENSION、DASHSCOPE_API_KEY、CHROMA_PATH", "维度变更后重建"],
        ["联网", "WEB_SEARCH_ENABLED、BOCHA_API_KEY、FRESHNESS、COUNT、TIMEOUT", "可全局关闭"],
        ["抓取", "CRAWL_ENABLED、CRAWL_SCHEDULE、并发/超时/大小/私网开关", "生产私网开关保持 false"],
        ["文件", "REPORT_DIR", "源码目录外、应用账号可写"],
        ["安全", "BACKEND_CORS_ORIGINS、DETECT_RATE_LIMIT_*", "生产使用明确域名与合理配额"],
    ], [2.6, 7.3, 10.4], font_size=8.3)
    add_heading(doc, "12.2 部署步骤", 2)
    add_numbered(doc, [
        "准备 Python、Node.js、MySQL 和持久化目录；创建 utf8mb4 数据库与最小权限账号。",
        "从模板创建 .env，填写强 SECRET_KEY 与必要外部服务 Key，检查 REPORT_DIR 和 CHROMA_PATH。",
        "安装 backend/requirements.txt，执行 python -m app.db.migrate 或 alembic upgrade head。",
        "按需执行 init_db/seed_demo_data；若切换向量配置，重建 Chroma 索引。",
        "启动 Uvicorn，访问 GET /api/health 和 Swagger；再执行认证、检测、报告与管理员权限冒烟测试。",
        "安装前端依赖并执行 npm run build，将 dist 部署到静态服务器/反向代理。",
        "配置 TLS、CORS、日志轮转、数据库与目录备份，确认调度器只在指定实例运行。",
    ])
    add_heading(doc, "12.3 运行监测", 2)
    add_table(doc, ["监测对象", "最低检查项", "异常信号"], [
        ["API", "健康检查、HTTP 5xx/429、响应时间", "持续 5xx、延迟突增、限流异常"],
        ["MySQL", "连接数、慢查询、容量、备份成功率", "连接耗尽、复制/备份失败"],
        ["Chroma", "集合数量、查询/写入失败、目录空间", "维度不匹配、孤立/缺失向量"],
        ["外部服务", "调用量、超时、错误率、配额", "401/429/5xx、持续超时"],
        ["业务", "检测量、降级率、仲裁失败率、报告失败率", "模型契约错误或质量显著变化"],
        ["安全", "登录失败、SSRF 拒绝、管理员操作日志", "异常来源、高频失败、越权尝试"],
    ], [3.0, 7.2, 10.1])

    add_heading(doc, "13 可维护性与配置管理设计", 1)
    add_heading(doc, "13.1 代码组织与依赖方向", 2)
    add_body(doc, "前端按 api、components、layouts、router、stores、utils、views 划分；后端按 api、schemas、services、crud、models、core、db 划分。上层依赖下层稳定接口，外部供应商逻辑集中在 llm_service、embedding_service 和 services/web，避免散落到路由与页面。")
    add_heading(doc, "13.2 数据库变更控制", 2)
    add_bullets(doc, [
        "任何表结构变更必须新增 Alembic 迁移，禁止只改 ORM 或手工执行未归档 SQL。",
        "迁移文件包含 upgrade/downgrade、约束名和索引；发布前在空库与已有数据副本上测试。",
        "模型、Schema、CRUD、接口文档和本概要设计的数据字典应在同一变更中同步。",
        "重大数据结构决策应记录 ADR，旧决策不删除，只标记被替代。",
    ])
    add_heading(doc, "13.3 测试与追踪", 2)
    add_body(doc, "项目已有认证、检测、RAG、LLM、Embedding、Chroma、报告、高风险、统计、日志、迁移、CORS、限流相关单元/接口测试。概要设计评审后，应建立需求—模块—接口—表—测试的追踪矩阵，确保设计项能够被验证。")
    add_heading(doc, "13.4 已知边界与后续演进", 2)
    add_table(doc, ["边界", "影响", "建议"], [
        ["进程内限流", "多实例配额不共享", "生产横向扩容时改为 Redis/网关限流"],
        ["进程内 APScheduler", "多实例可能重复采集", "单独调度实例或分布式任务锁"],
        ["本地 Chroma/报告目录", "多实例共享与备份复杂", "迁移服务化向量库和对象存储"],
        ["外部模型与搜索依赖", "延迟、配额和内容质量波动", "监测降级率，支持供应商适配器与人工复核"],
        ["JSON 文本分析载荷", "数据库难以对内部字段查询", "稳定后可迁移为 MySQL JSON 或拆分审计表"],
    ], [4.3, 5.6, 10.4])

    add_heading(doc, "14 设计追踪与结论", 1)
    add_heading(doc, "14.1 需求—设计追踪矩阵", 2)
    add_table(doc, ["需求域", "设计模块", "主要接口", "数据对象", "验证重点"], [
        ["认证与权限", "4.1、10.1", "/auth/*、/admin/*", "users、system_logs", "密码散列、401/403、禁用账号"],
        ["新闻检测", "4.2～4.4、6.2", "/detect/news、extract-preview", "detection_records、evidence_matches", "主流程、分支公式、降级"],
        ["RAG 与联网", "4.3、6.3、7.6", "/rag/search", "knowledge_items、Chroma", "召回、回查、触发阈值、失败补偿"],
        ["历史与报告", "4.5、6.5", "/detect/*、/report/*", "detection_records、reports", "所有权、一对一、路径安全"],
        ["知识与采集", "4.6、6.4", "/admin/knowledge/*", "knowledge_items、crawl_tasks", "同步状态、重建、去重"],
        ["高风险与管理", "4.7、10.5", "/high-risk/*、/admin/high-risk/*", "detection_records、system_logs", "三条件公开、审计"],
        ["统计与运维", "4.7、12", "/admin/statistics/*、/health", "全业务表", "口径、日期边界、运行监测"],
    ], [3.0, 3.2, 4.8, 4.3, 5.0], font_size=8.0)
    add_heading(doc, "14.2 设计评审检查表", 2)
    add_table(doc, ["检查项", "结论", "评审备注"], [
        ["架构覆盖需求且模块职责无明显重叠", "通过/待确认", ""],
        ["接口权限、错误码和资源所有权清晰", "通过/待确认", ""],
        ["8 张关系表、外键、索引和向量集合可实现", "通过/待确认", ""],
        ["评分、仲裁、降级和高风险公开规则可测试", "通过/待确认", ""],
        ["SSRF、秘密、报告路径和审计控制充分", "通过/待确认", ""],
        ["备份、迁移、索引重建和故障恢复可执行", "通过/待确认", ""],
    ], [9.2, 3.2, 7.9])
    add_heading(doc, "14.3 结论", 2)
    add_body(doc, "本概要设计把需求基线落实为前后端分离、分层后端、关系事实库与向量派生索引协作的可实施方案。核心鉴别流程以证据候选治理为中心，通过服务端契约校验、分支式评分、风险分级和确定性降级控制大模型不确定性；数据库设计覆盖业务、审计、采集和派生文件追踪。经评审批准后，可据此开展详细设计、接口联调、数据库迁移验证、系统测试和部署验收。")

    add_heading(doc, "附录 A 状态与枚举数据字典", 1)
    add_table(doc, ["数据项", "取值", "说明"], [
        ["users.role", "user / admin", "普通用户/管理员"],
        ["users.status", "active / disabled", "可用/禁用"],
        ["risk_level", "可信新闻 / 存疑信息 / 疑似谣言 / 高风险谣言", "由最终分统一映射"],
        ["review_status", "pending / approved / rejected", "待审/通过/驳回"],
        ["vector_sync_status", "pending / synced / failed / delete_failed", "向量同步状态"],
        ["prompt status", "enabled / disabled", "模板启停"],
        ["crawl_tasks.status", "running / success / partial / failed", "采集执行状态"],
        ["evidence stance", "support / contradict / neutral", "证据立场"],
        ["arbitration_status", "ok / failed / degraded（按响应载荷）", "证据仲裁结果状态"],
    ], [4.2, 7.1, 8.9])

    add_heading(doc, "附录 B 配置与目录基线", 1)
    add_table(doc, ["对象", "默认/示例", "注意事项"], [
        ["API 前缀", "/api", "前端代理与反向代理保持一致"],
        ["后端服务", "127.0.0.1:8000", "生产经反向代理和 TLS 暴露"],
        ["前端开发服务", "127.0.0.1:5173", "仅开发环境"],
        ["数据库名", "zhiyun_bianzhen", "utf8mb4，InnoDB"],
        ["Chroma 集合", "knowledge_items", "余弦空间，可重建"],
        ["Embedding", "dashscope / text-embedding-v4 / 1024", "变更后全量重建"],
        ["检测限流", "3 次 / 60 秒", "可配置；多实例需共享限流"],
        ["报告目录", "REPORT_DIR", "必须位于后端源码目录外"],
        ["采集计划", "0 */6 * * *", "默认每 6 小时；部署时确认单实例执行"],
    ], [4.1, 6.5, 9.6])


def set_document_metadata(doc):
    props = doc.core_properties
    props.title = f"《{PROJECT_TITLE}》{DOC_TITLE}"
    props.subject = "软件工程概要设计、数据库设计、接口与运行设计"
    props.author = "项目开发组"
    props.keywords = "概要设计, RAG, 大语言模型, 新闻真伪鉴别, FastAPI, Vue3, MySQL, Chroma"
    props.comments = f"文档编号 {DOC_NO}；版本 {VERSION}；基线 {BASELINE}"
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_main_content(doc)
    set_document_metadata(doc)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(path)
