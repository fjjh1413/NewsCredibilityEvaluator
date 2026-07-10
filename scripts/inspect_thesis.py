from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("unsupported parent")

    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def table_text(table: Table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def main() -> None:
    path = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    doc = Document(path)
    blocks = []
    for i, block in enumerate(iter_block_items(doc)):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            drawings = len(block._p.xpath(".//w:drawing")) + len(block._p.xpath(".//w:pict"))
            blocks.append(
                {
                    "index": i,
                    "type": "paragraph",
                    "style": block.style.name if block.style else None,
                    "text": text,
                    "drawings": drawings,
                }
            )
        else:
            blocks.append(
                {
                    "index": i,
                    "type": "table",
                    "rows": len(block.rows),
                    "cols": len(block.columns),
                    "text": table_text(block),
                }
            )

    summary = {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "blocks": blocks,
    }
    payload = json.dumps(summary, ensure_ascii=False, indent=2)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(payload, encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
