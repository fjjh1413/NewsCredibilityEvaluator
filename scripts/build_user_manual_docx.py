from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_ROOT = Path(r"E:\nan\《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》\报告")
SCREENSHOT_ROOT = REPORT_ROOT / "使用说明书截图"
OUTPUT_PATH = REPORT_ROOT / "《基于 RAG 与大语言模型的网络新闻真伪鉴别系统》使用说明书.docx"

TITLE = "《基于 RAG 与大语言模型的网络新闻真伪鉴别系统》"
SHORT_TITLE = "网络新闻真伪鉴别系统使用说明书"
BLUE = "0B6E99"
LIGHT_BLUE = "EAF5FA"
PALE_BLUE = "F5FAFD"
ORANGE = "F59E0B"
RED = "D92D20"
GREEN = "16A368"
GRAY = "667085"


def set_run_font(run, east_asia: str = "宋体", latin: str = "Times New Roman", size: float = 10.5,
                 bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def set_page_number_start(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def configure_section(section, with_header_footer: bool = True) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    if not with_header_footer:
        return

    section.header.is_linked_to_previous = False
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(2)
    run = header_p.add_run(SHORT_TITLE)
    set_run_font(run, east_asia="宋体", size=8.5, color=GRAY)

    section.footer.is_linked_to_previous = False
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("—  ")
    set_run_font(run, size=9, color=GRAY)
    add_field(footer_p, "PAGE", "1")
    run = footer_p.add_run("  —")
    set_run_font(run, size=9, color=GRAY)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name, size, color, before, after in (
        ("Title", 28, "101828", 0, 20),
        ("Heading 1", 18, BLUE, 18, 10),
        ("Heading 2", 14, "12334A", 14, 7),
        ("Heading 3", 12, "1D4F68", 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(3)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None, indent: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, east_asia="黑体", size=10.5, bold=True, color="12334A")
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_steps(doc: Document, steps: list[str]) -> None:
    for index, text in enumerate(steps, 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        run = p.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        run = p.add_run(text)
        set_run_font(run)


def add_note(doc: Document, title: str, text: str, *, kind: str = "info") -> None:
    fill, accent = {
        "info": (LIGHT_BLUE, BLUE),
        "warning": ("FFF7E8", ORANGE),
        "danger": ("FFF1F0", RED),
        "success": ("ECFDF3", GREEN),
    }[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(f"{title}：")
    set_run_font(r1, east_asia="黑体", size=10, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, size=10, color="344054")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for i, header in enumerate(headers):
        cell = header_row.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(header)
        set_run_font(r, east_asia="黑体", size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, PALE_BLUE)
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.2
            r = p.add_run(value)
            set_run_font(r, size=9.5)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, filename: str, caption: str, *, width_inches: float = 6.35) -> None:
    image_path = SCREENSHOT_ROOT / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width_inches))

    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.first_line_indent = Cm(0)
    caption_p.paragraph_format.space_after = Pt(8)
    caption_p.paragraph_format.keep_with_next = False
    run = caption_p.add_run(caption)
    set_run_font(run, size=9, color=GRAY)


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def build_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    mark = doc.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark.paragraph_format.first_line_indent = Cm(0)
    run = mark.add_run("真")
    set_run_font(run, east_asia="黑体", size=30, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("《基于 RAG 与大语言模型的网络新闻\n真伪鉴别系统》")
    set_run_font(run, east_asia="黑体", size=25, bold=True, color="101828")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("使 用 说 明 书")
    set_run_font(run, east_asia="黑体", size=28, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("User Manual")
    set_run_font(run, east_asia="Arial", latin="Arial", size=13, bold=True, color=GRAY)

    for _ in range(3):
        doc.add_paragraph()

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    info = [
        ("文档版本", "V1.0"),
        ("适用系统", "智闻辨真 V0.1.0"),
        ("适用角色", "普通用户、系统管理员"),
        ("编制日期", "2026年6月"),
    ]
    for row, (key, value) in zip(table.rows, info):
        row.cells[0].width = Cm(4.2)
        row.cells[1].width = Cm(8.2)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.first_line_indent = Cm(0)
        r1 = row.cells[0].paragraphs[0].add_run(key)
        set_run_font(r1, east_asia="黑体", size=10.5, bold=True, color=BLUE)
        r2 = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(r2, size=10.5)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("基于实际运行界面编制")
    set_run_font(run, size=10, color=GRAY)


def build_toc_and_preface(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("目录")
    set_run_font(r, east_asia="黑体", size=22, bold=True, color=BLUE)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "打开文档后按 F9 更新目录")
    add_note(doc, "目录提示", "若目录页码未自动显示，请在 Word 中选中目录并按 F9 更新域。", kind="info")

    add_page_break(doc)
    doc.add_heading("编制说明", level=1)
    add_body(
        doc,
        "本说明书参照课程样例的编排方式，以系统真实运行界面为依据，说明《基于 RAG 与大语言模型的网络新闻真伪鉴别系统》的访问、登录和主要功能操作。文档将普通用户与管理员功能分章描述，截图均来自 2026 年 6 月 20 日在本地演示环境中的实际操作。",
    )
    add_body(
        doc,
        "系统以“智闻辨真”为界面名称，采用 RAG 知识库检索、DeepSeek 大语言模型分析与规则评分相结合的方式，对输入新闻形成可信度评分、风险等级、判断理由、风险点、有效证据和辟谣建议。系统结论用于辅助核验，不能替代权威机构通报、专业事实核查或人工判断。",
    )
    add_note(doc, "适用范围", "本文档适用于课程演示环境和同版本部署环境。页面数据量、用户名、时间和评分会随实际运行数据变化。", kind="warning")


def build_chapter_1(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("1 系统概述", level=1)
    doc.add_heading("1.1 系统简介", level=2)
    add_body(
        doc,
        "本系统面向需要快速核验网络新闻、社交平台消息和转述内容的用户。用户提交新闻标题、正文及可选来源信息后，系统依次执行文本预处理、知识库证据检索、联网补充检索、大语言模型分析、规则评分和综合评估，并将结果保存为可追溯记录。",
    )
    add_bullets(doc, [
        "支持手工输入新闻，也支持粘贴网页链接后提取标题与正文。",
        "通过 RAG 检索知识库中的相似新闻与核查材料，并展示有效证据。",
        "将证据质量、大模型判断和来源/规则评分组合为综合可信度评分。",
        "支持普通用户查看个人历史、公开高风险案例及生成 PDF 报告。",
        "支持管理员维护用户、检测记录、知识库、Prompt、高风险审核、统计、报告和系统日志。",
    ])

    doc.add_heading("1.2 角色与权限", level=2)
    add_table(
        doc,
        ["角色", "主要入口", "主要权限", "限制"],
        [
            ["游客", "首页、新闻检测、高风险新闻", "浏览公开页面，可提交临时检测", "不能保存个人历史，不能进入后台"],
            ["普通用户", "用户端顶部导航", "检测新闻、查看结果、生成报告、查看本人历史和个人中心", "只能查看自己的历史与报告"],
            ["管理员", "管理员后台侧边栏", "管理全站用户、检测、知识库、Prompt、高风险、统计、报告和日志", "管理操作应遵循最小权限和审计要求"],
        ],
        widths=[2.3, 3.5, 6.8, 4.2],
    )

    doc.add_heading("1.3 运行环境与访问方式", level=2)
    add_table(
        doc,
        ["项目", "建议配置或说明"],
        [
            ["浏览器", "Chrome、Edge 等现代浏览器，建议分辨率不低于 1280×720"],
            ["前端地址", "本地演示环境：http://127.0.0.1:5173"],
            ["后端服务", "FastAPI 服务默认运行于 http://127.0.0.1:8000"],
            ["网络", "启用联网搜索或真实大模型分析时，需要可用网络与相应 API 配置"],
            ["演示账号", "普通用户 user_demo；管理员 admin_demo；密码由部署方或 seed 控制台提供"],
        ],
        widths=[3.3, 13.5],
    )
    add_note(doc, "安全提示", "不要在公开文档中写入生产账号密码、API Key 或数据库口令。演示账号仅用于课程展示。", kind="danger")

    doc.add_heading("1.4 评分与风险等级", level=2)
    add_table(
        doc,
        ["风险等级", "评分区间", "含义", "建议"],
        [
            ["可信新闻", "80–100", "现有证据和来源整体支持新闻内容", "仍应保留基本核查意识"],
            ["存疑信息", "60–79", "存在信息缺口或部分线索不一致", "补充权威来源后再判断"],
            ["疑似谣言", "40–59", "风险特征较明显或模型判断存在冲突", "谨慎传播并进行人工核验"],
            ["高风险谣言", "0–39", "高危特征突出，可信度较低", "不建议转发，必要时向平台举报"],
        ],
        widths=[3.0, 2.5, 6.5, 4.8],
    )


def build_chapter_2(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("2 系统访问、注册与登录", level=1)
    doc.add_heading("2.1 访问首页", level=2)
    add_body(doc, "在浏览器地址栏输入系统访问地址并回车，系统进入用户端首页。首页展示系统能力、检测流程和四级风险说明，顶部导航可进入新闻检测、历史记录、高风险新闻和个人中心。")
    add_figure(doc, "01-home.png", "图2-1 系统用户端首页")

    doc.add_heading("2.2 注册普通用户", level=2)
    add_body(doc, "未拥有账号的用户可单击页面右上角“注册”或登录页中的“立即注册”进入注册页面。管理员账号不能由前台注册产生，应由系统初始化或管理员在受控环境中创建。")
    add_steps(doc, [
        "输入 2～30 个字符的用户名，建议使用便于识别且不包含敏感信息的名称。",
        "按需输入邮箱；邮箱用于识别账号，若填写则不能与已有账号重复。",
        "输入密码并在确认密码框再次输入相同内容。",
        "单击“注册账号”。注册成功后，系统跳转登录页，使用新账号登录。",
    ])
    add_figure(doc, "03-register.png", "图2-2 普通用户注册界面")
    add_note(doc, "注册校验", "若用户名或邮箱已存在、两次密码不一致或字段长度不符合要求，系统会在表单处提示并阻止提交。", kind="warning")

    doc.add_heading("2.3 登录系统", level=2)
    add_steps(doc, [
        "单击页面右上角“登录”，进入登录页面。",
        "在“账号/用户名”输入框填写用户名，在“密码”输入框填写对应密码。",
        "单击“登录并进入检测”。普通用户登录后进入新闻检测页；管理员登录后进入后台首页。",
        "若提示用户名或密码错误，请检查大小写、账号状态及演示账号是否已完成初始化。",
    ])
    add_figure(doc, "02-login.png", "图2-3 系统登录界面")

    doc.add_heading("2.4 退出登录", level=2)
    add_body(doc, "普通用户可单击顶部右侧电源图标，或在个人中心单击“退出登录”；管理员可单击后台右上角电源图标。系统弹出确认框后，单击“退出登录”清除当前会话。")
    add_note(doc, "公共设备", "在教室、实验室等公共设备上使用后必须退出登录，并关闭浏览器窗口，避免本地会话被继续使用。", kind="danger")


def build_chapter_3(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("3 普通用户使用方法", level=1)
    doc.add_heading("3.1 功能概述", level=2)
    add_table(
        doc,
        ["功能", "入口", "用途"],
        [
            ["新闻检测", "顶部导航“新闻检测”", "提交新闻，执行证据检索、模型分析和规则评分"],
            ["检测结果", "提交成功自动进入", "查看评分、风险等级、理由、风险点、证据和建议"],
            ["历史记录", "顶部导航“历史记录”", "查看、筛选本人检测记录及下载已生成报告"],
            ["高风险新闻", "顶部导航“高风险新闻”", "浏览管理员审核并公开的高风险案例"],
            ["个人中心", "顶部导航“个人中心”", "查看账号信息和常用功能入口"],
        ],
        widths=[3.0, 5.0, 8.8],
    )

    doc.add_heading("3.2 提交新闻检测", level=2)
    add_body(doc, "进入“新闻检测”页面后，可使用链接提取或手工输入两种方式。页面右侧展示五步检测流程，底部可控制是否启用联网搜索。")
    add_steps(doc, [
        "可选：在“新闻链接”框粘贴网页地址，单击“提取”，系统将尝试提取标题、正文、来源和发布时间。若提取失败，可改为手工输入。",
        "填写新闻标题。标题用于检索和结果展示，应尽量完整。",
        "选择新闻类别，并可填写来源媒体或平台。来源越清晰，规则评分越有参考价值。",
        "可选填写新闻发布日期和具体时间，用于判断时效性。",
        "在“新闻正文”中粘贴需要核验的完整内容。正文过短会降低分析信息量。",
        "按需打开“启用联网搜索”。打开后系统会在本地知识库证据不足时补充网络检索，但检测耗时可能增加。",
        "单击“提交可信度检测”，等待系统完成证据检索、DeepSeek 分析、规则评分和综合评估。",
    ])
    add_figure(doc, "04-user-detect.png", "图3-1 普通用户新闻检测界面")
    add_note(doc, "示例填充", "课程演示时可单击“填充示例新闻”自动填入示例内容，再根据需要修改后提交。", kind="success")

    doc.add_heading("3.3 查看检测结果", level=2)
    add_body(doc, "检测完成后系统自动进入结果页。页面顶部显示新闻标题、检测时间、综合可信度评分和风险等级；下方依次展示证据质量分、大模型判断分、来源/规则评分及可解释分析。")
    add_figure(doc, "05-user-result.png", "图3-2 检测结果总览")
    add_bullets(doc, [
        "综合可信度评分：0～100 分，分值越高表示当前证据下可信度越高。",
        "证据质量：反映有效证据的覆盖度和相互一致性。",
        "大模型判断：基于新闻文本与证据形成的结构化判断分。",
        "来源/规则评分：检查来源完整性、夸张或绝对化表述、风险关键词等。",
        "判断理由与辟谣建议：解释结论，并给出进一步核验或传播建议。",
    ])
    add_figure(doc, "05-user-result-details.png", "图3-3 证据质量、判断理由与辟谣建议")

    doc.add_heading("3.4 查看风险点、关键词与有效证据", level=2)
    add_body(doc, "继续向下滚动结果页，可查看模型识别的风险点、关键词和有效证据列表。有效证据标明来源类型、标题和摘要，用于解释为什么系统得到当前结论。若证据不足或相互冲突，应优先进行人工核验。")
    add_figure(doc, "05-user-result-evidence.png", "图3-4 风险点、关键词与有效证据")

    doc.add_heading("3.5 生成与下载 PDF 报告", level=2)
    add_steps(doc, [
        "在结果页右上角单击“生成 PDF 报告”。",
        "按钮显示“报告生成中”时请勿重复单击，等待服务端完成生成。",
        "生成成功后按钮变为“下载 PDF 报告”，单击即可下载；需要覆盖旧报告时可单击“重新生成”。",
        "也可进入“历史记录”，在对应记录的“报告状态”列下载已生成报告。",
    ])
    add_figure(doc, "05-user-result-report-ready.png", "图3-5 PDF 报告已生成状态")

    doc.add_heading("3.6 查看历史记录", level=2)
    add_body(doc, "单击顶部导航“历史记录”，页面显示当前账号提交的检测记录。可按新闻标题关键词和风险等级筛选；列表展示检测时间、评分、风险等级、高风险标记和报告状态。")
    add_steps(doc, [
        "在关键词输入框中输入新闻标题片段，或选择风险等级。",
        "单击“查询”显示符合条件的记录；单击“重置”恢复全部记录。",
        "单击“详情”返回完整检测结果；若报告已生成，可直接单击“下载报告”。",
    ])
    add_figure(doc, "06-user-history.png", "图3-6 普通用户历史记录")

    doc.add_heading("3.7 浏览公开高风险新闻", level=2)
    add_body(doc, "“高风险新闻”页面仅展示经管理员审核且允许公开的记录。可按关键词、类别、风险等级和时间范围筛选，并查看公开记录数、排行榜记录数、风险关键词和涉及类别。未审核记录不会出现在用户端。")
    add_figure(doc, "07-user-high-risk.png", "图3-7 用户端高风险新闻页面")

    doc.add_heading("3.8 使用个人中心", level=2)
    add_body(doc, "个人中心展示当前用户名、邮箱、角色和注册时间，并提供新闻检测、历史记录和高风险新闻的快捷入口。普通用户不能在此修改角色，也不能进入管理员后台。")
    add_figure(doc, "08-user-profile.png", "图3-8 普通用户个人中心")


def build_chapter_4(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("4 管理员使用方法", level=1)
    doc.add_heading("4.1 登录与后台布局", level=2)
    add_body(doc, "使用管理员账号登录后，系统自动跳转至“管理员后台”。后台左侧为固定功能菜单，右上角显示当前管理员账号和退出按钮，页面底部提供“返回用户端”。普通用户即使直接访问后台地址，也会被权限守卫拦截。")
    add_note(doc, "管理原则", "涉及禁用用户、删除记录、重建向量、修改 Prompt 或公开高风险新闻的操作会影响系统运行或对外展示，执行前应核对对象和影响范围。", kind="danger")

    doc.add_heading("4.2 后台首页", level=2)
    add_body(doc, "后台首页汇总总检测数、高风险新闻数、知识库数量和用户总数，并展示近 7 日检测趋势及风险等级分布。单击“刷新”获取最新数据，单击“完整数据统计”进入统计页。")
    add_figure(doc, "09-admin-dashboard.png", "图4-1 管理员后台首页")

    doc.add_heading("4.3 用户管理", level=2)
    add_body(doc, "用户管理页支持按用户名或邮箱、角色、状态筛选账号。列表展示用户名、邮箱、角色、状态和注册时间。管理员可查看用户详情与检测记录，并可禁用或启用账号。")
    add_steps(doc, [
        "输入用户名或邮箱，或选择角色、状态后单击“查询”。",
        "单击“详情”查看账号资料；单击“检测记录”查看该用户提交的数据。",
        "需要限制账号访问时单击“禁用”并确认；恢复账号时使用“启用”。",
    ])
    add_figure(doc, "10-admin-users.png", "图4-2 管理员用户管理")

    doc.add_heading("4.4 检测记录管理", level=2)
    add_body(doc, "检测记录管理页可查看全站检测数据，并按关键词、风险等级、时间范围和用户 ID 检索。管理员可查看详情或删除记录。删除检测记录可能关联报告和高风险数据，应谨慎操作。")
    add_figure(doc, "11-admin-detections.png", "图4-3 管理员检测记录管理")

    doc.add_heading("4.5 知识库管理", level=2)
    add_body(doc, "知识库是 RAG 检索的主要证据来源。页面支持按关键词、类别、真实性标签、风险等级和向量状态筛选，可新增、编辑、删除知识项，查看原文，并执行单条重新向量化或全量重建。")
    add_steps(doc, [
        "新增知识：单击“新增知识”，填写标题、正文、摘要、来源、标签等信息后保存。",
        "编辑知识：在目标行单击“编辑”，修改后保存；正文或核心字段变化后建议重新向量化。",
        "重新向量化：单击“重新向量化”同步单条记录到 Chroma 向量库。",
        "全量重建：仅在 embedding 模型/维度切换或索引异常时使用，并确保运行期间不重复触发。",
        "删除知识：确认该材料不再作为证据后再删除。",
    ])
    add_figure(doc, "12-admin-knowledge.png", "图4-4 管理员知识库管理")
    add_note(doc, "向量一致性", "切换 embedding provider 或向量维度后必须全量重建索引，否则可能出现维度不匹配或检索结果失真。", kind="warning")

    doc.add_heading("4.6 Prompt 模板管理", level=2)
    add_body(doc, "Prompt 模板用于约束新闻可信度分析的输入结构和输出字段。管理员可按名称、类型或内容搜索模板，新增、编辑、启用/停用、删除模板，并维护同类型的默认模板。")
    add_steps(doc, [
        "单击“新增模板”填写模板名称、类型和 Prompt 内容。",
        "模板应保留系统所需占位符，并明确输出字段与风险等级取值。",
        "修改后先在演示新闻上验证输出结构，再将模板启用或设为默认。",
        "停用或删除前确认系统仍存在可用的同类型默认模板。",
    ])
    add_figure(doc, "13-admin-prompts.png", "图4-5 Prompt 模板管理")

    doc.add_heading("4.7 高风险新闻管理", level=2)
    add_body(doc, "该页面用于审核高风险检测记录并维护公开状态。可按标题关键词、风险等级、新闻类别、审核状态、公开状态和时间范围筛选。只有“审核通过”且设置公开的记录才会出现在用户端。")
    add_steps(doc, [
        "单击“详情”核对新闻正文、评分、证据和风险理由。",
        "单击“审核”选择通过或驳回，并填写必要的审核备注。",
        "审核通过后按需设置公开；不再适合公开时单击“取消公开”。",
        "需要补充内部说明时单击“备注”。审核通过不等于自动公开。",
    ])
    add_figure(doc, "14-admin-high-risk.png", "图4-6 高风险新闻审核与公开管理")

    doc.add_heading("4.8 数据统计", level=2)
    add_body(doc, "数据统计页支持按起止日期查询，并提供近 7 日、近 30 日和近 90 日快捷范围。页面展示今日检测数、总检测数、高风险新闻数、知识库总数以及趋势和分布图表。")
    add_figure(doc, "15-admin-statistics.png", "图4-7 管理员数据统计")

    doc.add_heading("4.9 报告管理", level=2)
    add_body(doc, "报告管理页汇总检测报告、关联用户、检测结果、风险等级、可信度分值和 PDF 文件状态。可按关键词、用户 ID、检测 ID、报告状态和时间范围筛选，并查看详情或下载 PDF。")
    add_figure(doc, "16-admin-reports.png", "图4-8 管理员报告管理")

    doc.add_heading("4.10 系统日志", level=2)
    add_body(doc, "系统日志记录登录、新闻检测和关键管理操作。管理员可按用户名/描述/IP、模块、操作以及起止时间筛选，用于演示审计、问题排查和异常追踪。")
    add_steps(doc, [
        "输入关键词或选择模块，必要时填写操作名，如 login、create、review。",
        "设置时间范围后单击“查询”；单击“重置”清空筛选条件。",
        "结合用户、模块、描述、IP 和时间判断操作链路，不应随意删除审计信息。",
    ])
    add_figure(doc, "17-admin-logs.png", "图4-9 管理员系统日志")

    doc.add_heading("4.11 返回用户端与退出", level=2)
    add_body(doc, "单击左侧底部“返回用户端”可回到普通用户界面，管理员会话仍然保留；需要结束会话时，应单击后台右上角电源图标并确认退出。")


def build_chapter_5(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("5 常见问题与注意事项", level=1)

    doc.add_heading("5.1 无法登录", level=2)
    add_bullets(doc, [
        "确认用户名与密码大小写正确，输入框前后没有多余空格。",
        "演示账号应先执行数据初始化，密码以 DEMO_PASSWORD、ADMIN_DEMO_PASSWORD 或 seed 控制台输出为准。",
        "若账号被管理员禁用，需要管理员在“用户管理”中恢复启用。",
        "确认后端服务和数据库可访问；若服务未启动，前端会显示请求失败。",
    ])

    doc.add_heading("5.2 检测耗时较长或失败", level=2)
    add_bullets(doc, [
        "启用联网搜索后，系统需要额外请求网页和搜索服务，耗时通常更长。",
        "检查 DeepSeek、DashScope 和联网搜索 API 配置及网络连接。",
        "正文过长时可保留完整事实主张、来源和时间信息，删除与核验无关的广告或重复段落。",
        "检测过程中不要连续重复提交；如页面长时间无响应，可查看后端日志和系统日志。",
    ])

    doc.add_heading("5.3 报告无法生成或下载", level=2)
    add_bullets(doc, [
        "确认当前用户已登录且有权访问该检测记录。",
        "确认报告目录可写，后端 PDF 生成依赖已正确安装。",
        "报告生成中请等待按钮恢复；已生成报告可从历史记录或管理员报告管理页下载。",
    ])

    doc.add_heading("5.4 结果理解与使用边界", level=2)
    add_note(doc, "重要声明", "系统是新闻可信度辅助评估工具，评分和风险等级仅供参考。涉及公共安全、医疗健康、法律、金融等重要信息时，应以主管部门、权威媒体和专业事实核查结果为准。", kind="danger")
    add_bullets(doc, [
        "不要只看总分，应同时查看有效证据、判断理由、风险点和来源/规则评分。",
        "证据覆盖不足时，即使某一项评分较高，也应继续核验原始来源。",
        "管理员公开高风险案例前，应确认内容不会侵犯隐私、泄露敏感信息或造成二次传播风险。",
    ])

    doc.add_heading("5.5 数据与账号安全", level=2)
    add_bullets(doc, [
        "不得在新闻正文中提交身份证号、联系方式、未公开业务数据等敏感信息。",
        "管理员不得共享账号；重要操作应通过系统日志保留审计轨迹。",
        "生产环境应使用强密码、固定 CORS 白名单、独立密钥和受控数据库权限。",
        "演示环境中的账号、新闻和统计数据仅用于课程展示，不应视为真实业务数据。",
    ])


def build_appendices(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("附录A 功能权限对照表", level=1)
    add_table(
        doc,
        ["功能", "游客", "普通用户", "管理员"],
        [
            ["浏览首页", "允许", "允许", "允许"],
            ["提交新闻检测", "允许（临时结果）", "允许并保存历史", "允许"],
            ["查看个人历史", "不允许", "仅本人", "可在后台查看全站记录"],
            ["生成/下载本人报告", "不允许", "允许", "允许查看全站报告"],
            ["浏览公开高风险新闻", "允许", "允许", "允许"],
            ["用户、知识库、Prompt 管理", "不允许", "不允许", "允许"],
            ["高风险审核与公开", "不允许", "不允许", "允许"],
            ["数据统计与系统日志", "不允许", "不允许", "允许"],
        ],
        widths=[6.3, 3.4, 3.6, 3.6],
    )

    doc.add_heading("附录B 演示操作速查", level=1)
    add_table(
        doc,
        ["角色", "推荐演示路径"],
        [
            ["普通用户", "登录 user_demo → 新闻检测 → 填充示例新闻 → 提交检测 → 查看结果 → 生成 PDF → 历史记录 → 高风险新闻 → 个人中心"],
            ["管理员", "登录 admin_demo → 后台首页 → 用户管理 → 检测记录 → 知识库 → Prompt → 高风险审核 → 数据统计 → 报告管理 → 系统日志"],
        ],
        widths=[3.0, 13.9],
    )
    add_note(doc, "演示口令", "演示密码不写入本说明书，请以部署方提供或 seed 初始化时控制台输出的口令为准。", kind="info")


def enable_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def build_document() -> Path:
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], with_header_footer=False)
    build_cover(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section, with_header_footer=True)
    set_page_number_start(body_section, 1)

    build_toc_and_preface(doc)
    build_chapter_1(doc)
    build_chapter_2(doc)
    build_chapter_3(doc)
    build_chapter_4(doc)
    build_chapter_5(doc)
    build_appendices(doc)
    enable_update_fields(doc)

    doc.core_properties.title = f"{TITLE}使用说明书"
    doc.core_properties.subject = "普通用户与管理员操作说明"
    doc.core_properties.author = "《基于 RAG 与大语言模型的网络新闻真伪鉴别系统》项目组"
    doc.core_properties.keywords = "RAG, 大语言模型, 新闻真伪鉴别, 使用说明书"
    doc.core_properties.comments = "依据本地实际运行界面与截图编制"

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
