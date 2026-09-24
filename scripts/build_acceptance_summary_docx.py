from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_outline_design_docx as base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = Path(__file__).resolve().parents[1] / ".artifacts" / "docs"
OUTPUT_PATH = OUTPUT_DIR / "05_《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》项目验收总结报告.docx"

PROJECT_TITLE = "基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计"
SYSTEM_NAME = "智闻辨真"
DOC_TITLE = "项目验收总结报告"
DOC_NO = "NCE-ACR-001"
VERSION = "V1.0"
BASELINE = "ACC-BL-2026-06-20"
DATE_TEXT = "2026 年 06 月 20 日"


# Keep the five-document family visually and semantically consistent.
base.PROJECT_TITLE = PROJECT_TITLE
base.SYSTEM_NAME = SYSTEM_NAME
base.DOC_TITLE = DOC_TITLE
base.DOC_NO = DOC_NO
base.VERSION = VERSION
base.BASELINE = BASELINE
base.DATE_TEXT = DATE_TEXT

add_heading = base.add_heading
add_body = base.add_body
add_bullets = base.add_bullets
add_numbered = base.add_numbered
add_note = base.add_note
add_table = base.add_table
add_toc = base.add_toc
add_figure = base.add_figure
configure_section = base.configure_section
configure_styles = base.configure_styles
set_run_font = base.set_run_font
set_cell_shading = base.set_cell_shading
set_cell_margins = base.set_cell_margins


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    configure_section(section, header_footer=False)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SOFTWARE ENGINEERING DOCUMENT")
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=9, bold=True, color=base.TEAL)

    bar = doc.add_table(rows=1, cols=1)
    bar.autofit = False
    set_cell_shading(bar.cell(0, 0), base.TEAL)
    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(f"《{PROJECT_TITLE}》")
    set_run_font(r, east_asia="黑体", latin="Arial", size=20, bold=True, color=base.NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(DOC_TITLE)
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=29, bold=True, color=base.BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("PROJECT ACCEPTANCE SUMMARY REPORT")
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=11, bold=True, color=base.GRAY)

    status = doc.add_table(rows=1, cols=1)
    status.alignment = WD_TABLE_ALIGNMENT.CENTER
    status.autofit = False
    cell = status.cell(0, 0)
    cell.width = Cm(7.6)
    set_cell_shading(cell, base.PALE_TEAL)
    set_cell_margins(cell, top=150, start=220, bottom=150, end=220)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("建议验收结论：通过")
    set_run_font(r, east_asia="黑体", size=13, bold=True, color=base.GREEN)

    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    info = doc.add_table(rows=7, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.style = "Table Grid"
    rows = [
        ("文档编号", DOC_NO),
        ("版本号", VERSION),
        ("配置基线", BASELINE),
        ("验收类型", "项目最终验收"),
        ("文档状态", "正式版（待签署）"),
        ("编制单位", "项目开发组"),
        ("编制日期", DATE_TEXT),
    ]
    for i, (label, value) in enumerate(rows):
        set_cell_shading(info.cell(i, 0), base.PALE_BLUE)
        for j, text in enumerate((label, value)):
            c = info.cell(i, j)
            c.width = Cm(4.2 if j == 0 else 8.4)
            set_cell_margins(c, top=125, bottom=125)
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run_font(
                r,
                east_asia="黑体" if j == 0 else "宋体",
                size=10.5,
                bold=(j == 0),
                color=base.NAVY if j == 0 else base.DARK,
            )

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("受控文档 · 验收结论以签署页为准")
    set_run_font(r, size=9.5, color=base.GRAY)


def add_front_matter(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, header_footer=True, page_format="lowerRoman", page_start=1)

    heading = add_heading(doc, "文档控制信息", 1)
    heading.paragraph_format.page_break_before = False
    add_body(
        doc,
        "本文件为项目最终验收阶段受控文档，用于记录验收依据、范围、环境、方法、执行结果、遗留问题、交付状态和验收意见。文中“通过”表示当前软件快照在已执行的验证范围内满足约定要求；正式验收效力以验收各方签署为准。",
    )

    add_heading(doc, "文档审批", 2)
    add_table(
        doc,
        ["角色", "姓名/单位", "职责", "签字", "日期"],
        [
            ["编制", "项目开发组", "汇总项目事实与验证证据", "", DATE_TEXT],
            ["复核", "", "复核需求、设计、测试与交付物", "", ""],
            ["验收组长", "", "组织评审并形成验收意见", "", ""],
            ["批准", "", "批准结项、移交与归档", "", ""],
        ],
        [2.0, 3.2, 5.0, 2.4, 2.6],
    )

    add_heading(doc, "修订记录", 2)
    add_table(
        doc,
        ["版本", "日期", "修订内容", "修订人", "状态"],
        [
            [
                VERSION,
                DATE_TEXT,
                "首次发布；依据当前代码知识图谱、自动化测试、生产构建、数据库迁移和前四类软件工程文档形成验收总结基线。",
                "项目开发组",
                "正式",
            ]
        ],
        [1.7, 2.7, 8.0, 2.3, 1.8],
    )

    add_heading(doc, "分发范围", 2)
    add_table(
        doc,
        ["序号", "接收方", "用途", "介质"],
        [
            ["1", "项目负责人/指导教师", "最终验收、答辩与结项", "电子版"],
            ["2", "开发、测试与维护人员", "缺陷跟踪、移交与维护", "电子版"],
            ["3", "配置管理员/档案管理人员", "基线固化和归档", "电子版"],
        ],
        [1.4, 4.2, 7.3, 2.5],
    )

    heading = add_heading(doc, "摘要", 1)
    heading.paragraph_format.page_break_before = False
    add_body(
        doc,
        f"{SYSTEM_NAME}是一套采用 Vue 3、FastAPI、MySQL、Chroma、RAG 与大语言模型构建的网络新闻可信度辅助评估系统。系统覆盖游客/用户新闻检测、URL 内容提取预览、证据检索与仲裁、综合评分、风险分级、历史记录、重新评估、PDF 报告、公开高风险信息，以及管理员用户、检测、知识、Prompt、高风险、报告、统计和日志管理。",
    )
    add_body(
        doc,
        "本次验收以当前工作区快照为对象。知识图谱增量索引完成后包含 3,645 个节点和 10,892 条关系，确认 8 个关系数据模型、37 个后端测试文件以及完整的检测与报告调用链。实际执行结果为：后端 392 项自动化测试全部通过；前端 14 项自动化测试全部通过；Vite 生产构建成功并转换 1,650 个模块；核心 Python 模块语法编译通过；Git 空白错误检查通过。",
    )
    add_note(
        doc,
        "验收结论摘要",
        "已验证范围内未发现阻断性问题，建议项目通过验收。DeepSeek、DashScope 与博查等第三方服务的实网质量和配额应在正式部署环境使用有效密钥复核；当前工作区的待提交变更应在归档前冻结为可追踪版本。",
        "success",
    )
    add_toc(doc)


def add_chapter_1(doc: Document) -> None:
    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 编写目的", 2)
    add_body(
        doc,
        "本文档对项目建设过程和最终成果进行系统总结，给出验收对象、依据、组织、环境、方法、执行证据、问题处置、结论与移交要求，为结项审批、配置归档、运行维护和后续改进提供统一依据。",
    )
    add_heading(doc, "1.2 项目背景", 2)
    add_body(
        doc,
        "面对网络新闻来源复杂、信息更新快、标题夸张和证据分散等问题，本项目以“证据可追踪、判断可解释、结果可复核”为目标，将本地知识库语义检索、按需联网补证、大语言模型分析、证据仲裁和确定性规则评分组合为受约束的新闻可信度辅助评估流程。",
    )
    add_heading(doc, "1.3 验收对象与范围", 2)
    add_table(
        doc,
        ["对象", "纳入范围", "不纳入范围"],
        [
            ["软件产品", "Vue 3 前端、FastAPI 后端、MySQL 数据模型、Chroma 向量索引、报告模板", "外部模型和搜索服务本身的服务等级"],
            ["业务功能", "认证、检测、RAG、补证、评分、历史、报告、高风险公开、管理后台", "人工事实核查、模型训练、舆情全网监测"],
            ["工程资产", "源代码、迁移脚本、测试、配置模板、设计与验收文档", "生产硬件采购和第三方商业合同"],
            ["质量属性", "正确性、安全基线、可构建性、可维护性、兼容性与可恢复性", "未经专项压测确认的高并发 SLA"],
        ],
        [3.0, 7.1, 5.4],
        font_size=8.8,
    )
    add_heading(doc, "1.4 术语和缩略语", 2)
    add_table(
        doc,
        ["术语", "缩写", "说明"],
        [
            ["检索增强生成", "RAG", "先检索证据，再以检索结果约束模型分析。"],
            ["大语言模型", "LLM", "用于新闻语义分析和证据仲裁的外部模型。"],
            ["向量嵌入", "Embedding", "将文本映射为固定维数向量以进行相似检索。"],
            ["验收测试", "AT", "依据验收准则对交付软件实施的确认性测试。"],
            ["配置基线", "Baseline", "经确认并纳入变更控制的配置项集合。"],
            ["证据仲裁", "—", "对候选证据的相关性、质量、立场和取舍进行结构化判定。"],
        ],
        [3.2, 2.4, 9.9],
    )
    add_heading(doc, "1.5 参考和采用的规范", 2)
    add_table(
        doc,
        ["序号", "规范/文件", "本报告中的应用"],
        [
            ["1", "GB/T 8567—2006《计算机软件文档编制规范》", "文档内容要素、编号、修订、审批和归档组织。"],
            ["2", "GB/T 28035—2011《软件系统验收规范》", "验收准备、实施、评价、结论和移交框架。"],
            ["3", "GB/T 9386—2008《计算机软件测试文档编制规范》", "测试记录、结果汇总、问题和结论表达。"],
            ["4", "GB/T 25000.51—2016《系统与软件工程 系统与软件质量要求和评价（SQuaRE） 第51部分：就绪可用软件产品（RUSP）的质量要求和测试细则》", "功能适合性、性能效率、兼容性、易用性、可靠性、安全性和维护性评价。"],
            ["5", "GB/T 11457—2006《信息技术 软件工程术语》", "术语使用和定义一致性。"],
            ["6", "GB/T 22239—2019《信息安全技术 网络安全等级保护基本要求》", "身份鉴别、访问控制、审计、数据和通信安全参考。"],
            ["7", "项目可行性研究报告、需求分析书、概要设计书、详细设计书", "需求、设计、数据和程序实现的验收基线。"],
        ],
        [1.2, 8.0, 6.3],
        font_size=8.5,
    )


def add_chapter_2(doc: Document) -> None:
    add_heading(doc, "2 项目概况与建设总结", 1)
    add_heading(doc, "2.1 项目标识", 2)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["项目名称", PROJECT_TITLE],
            ["软件名称", SYSTEM_NAME],
            ["软件形态", "B/S 架构、前后端分离的 Web 应用"],
            ["当前版本", VERSION],
            ["验收基线", BASELINE],
            ["代码分支/参考提交", "main / bf80101（验收工作区另含待冻结变更）"],
            ["验收日期", DATE_TEXT],
        ],
        [4.2, 11.3],
        first_col_bold=True,
    )
    add_heading(doc, "2.2 建设目标达成情况", 2)
    add_table(
        doc,
        ["建设目标", "实现情况", "评价"],
        [
            ["形成可追踪的新闻可信度评估流程", "候选证据带稳定标识，仲裁结果、排除理由、评分和风险点可保存。", "达成"],
            ["避免单纯依赖一次模型回答", "融合 RAG、条件联网补证、证据质量、规则评分和大模型评分；模型失败可确定性降级。", "达成"],
            ["支持用户完整业务闭环", "支持输入/链接预览、检测、结果、历史、重评、报告和公开高风险浏览。", "达成"],
            ["支持管理和治理", "提供用户、检测、知识、向量、Prompt、高风险、报告、统计和日志管理。", "达成"],
            ["形成规范工程交付物", "具备迁移脚本、自动化测试、配置模板、部署说明及五类软件工程文档。", "达成"],
        ],
        [5.2, 8.0, 2.3],
        font_size=8.7,
    )
    add_heading(doc, "2.3 系统总体构成", 2)
    add_body(
        doc,
        "系统采用分层结构。浏览器端负责页面交互、路由和状态管理；FastAPI 路由层执行参数、身份与权限校验；Service 层编排检测、知识、模型、报告和统计业务；CRUD/ORM 层访问 MySQL；Chroma 保存可重建的派生向量；DeepSeek、DashScope 和博查作为可配置外部服务。",
    )
    add_figure(doc, "fig4-1-system-architecture.png", "图 2-1 系统总体架构", max_width=14.8, max_height=12.5)
    add_heading(doc, "2.4 核心处理流程", 2)
    add_numbered(
        doc,
        [
            "清洗并校验标题、正文、来源、发布时间和联网开关，抽取关键词。",
            "以标题和正文构造检索文本，从 Chroma 召回本地知识候选；知识检索失败时给出明确服务错误。",
            "在用户允许且本地证据不足时调用博查补充网页证据，失败则回退到本地证据。",
            "为候选证据分配稳定 candidate_id，并以来源中立、可复现的顺序提交 DeepSeek。",
            "校验证据仲裁契约；首次失败时执行一次聚焦重试，仍失败则未仲裁候选不得参与评分。",
            "按证据可用性选择评分公式，映射四级风险，形成理由、风险点、建议和相似新闻。",
            "将检测记录、有效证据、分析载荷和审计信息写入 MySQL；按需生成受权限保护的 PDF 报告。",
        ],
    )
    add_figure(doc, "fig4-4-detection-overall-flow.png", "图 2-2 新闻可信度检测总体流程", max_width=14.8, max_height=15.5)
    add_heading(doc, "2.5 开发过程与主要成果", 2)
    add_table(
        doc,
        ["阶段", "主要活动", "形成成果"],
        [
            ["需求与论证", "问题分析、目标边界、可行性、角色与用例确认", "可行性研究报告、需求分析书"],
            ["总体设计", "架构、模块、接口、数据库、存储和安全设计", "概要设计书、数据库和 API 设计材料"],
            ["实现与集成", "前后端开发、RAG/LLM/规则融合、管理端和报告实现", "源代码、配置模板、迁移和种子数据"],
            ["详细设计", "程序级输入输出、处理、异常、依赖和测试设计", "详细设计书"],
            ["测试与验收", "知识图谱复核、自动化测试、生产构建和交付检查", "测试证据、项目验收总结报告"],
        ],
        [2.7, 7.0, 5.8],
        font_size=8.7,
    )
    add_heading(doc, "2.6 计划、进度、资源与费用评价", 2)
    add_table(
        doc,
        ["评价项", "实际情况", "验收评价"],
        [
            ["计划完成度", "需求、设计、实现、联调、测试、部署说明和验收文档阶段均形成对应成果。", "范围内工作已完成"],
            ["进度", "项目资料按七阶段组织；本报告不虚构未建立的计划/实际工期统计。", "以交付物完成状态确认，工期偏差不作量化"],
            ["人员", "由项目开发组承担需求、设计、前后端、数据库、测试和文档工作；未提供可核实的人数/工时台账。", "职责覆盖完整，工时生产率不作虚假估算"],
            ["软硬件资源", "使用现有 Windows 开发主机、Python/Node.js 开源工具链、MySQL、Chroma 和 Git。", "满足本期开发与验收需要"],
            ["费用", "未建立独立采购和人工成本核算；第三方模型、Embedding 与搜索按实际账号和调用量计费。", "不虚构金额；上线前应制定配额和费用预算"],
            ["成果规模", "知识图谱 3,645 节点/10,892 关系，8 个关系模型，37 个后端测试文件，406 项自动化测试，五类核心文档。", "成果可识别、可复核、可移交"],
        ],
        [3.0, 8.2, 4.3],
        font_size=8.2,
    )


def add_chapter_3(doc: Document) -> None:
    add_heading(doc, "3 验收组织、原则与准则", 1)
    add_heading(doc, "3.1 验收组织与职责", 2)
    add_table(
        doc,
        ["角色", "主要职责", "输出"],
        [
            ["验收负责人", "批准计划、主持评审、裁定问题级别和验收结论。", "验收意见、签署结论"],
            ["需求/业务代表", "确认范围、业务流程和输出可理解性。", "需求符合性意见"],
            ["技术评审人员", "复核架构、接口、数据库、安全、部署和可维护性。", "技术评审记录"],
            ["测试人员", "执行自动化与手工用例、记录环境和证据、复测问题。", "测试结果与缺陷清单"],
            ["开发人员", "说明实现、提供交付物、定位问题并完成整改。", "版本、说明和整改证据"],
            ["配置管理员", "固化代码、文档、迁移、依赖和发布包。", "配置基线与归档清单"],
        ],
        [3.0, 8.0, 4.5],
        font_size=8.8,
    )
    add_heading(doc, "3.2 验收原则", 2)
    add_bullets(
        doc,
        [
            "以批准的需求和设计基线为依据，以可复现证据为判断基础；",
            "对已执行项、未执行项和外部条件项分别记录，不以推测替代验证；",
            "阻断性安全、数据一致性或核心流程问题未关闭时不得判定通过；",
            "系统定位为可信度辅助评估工具，不把模型输出表述为权威真伪裁决；",
            "验收通过后应冻结配置项，后续变更履行版本和回归测试控制。",
        ],
    )
    add_heading(doc, "3.3 验收方法", 2)
    add_table(
        doc,
        ["方法", "实施内容", "主要证据"],
        [
            ["文档审查", "核对五类软件工程文档、README、数据库和接口资料。", "文档清单、章节与版本信息"],
            ["静态结构核验", "使用 codebase-memory-mcp 增量索引、搜索、调用链追踪和源片段读取。", "图谱节点/关系、模型和调用链结果"],
            ["自动化测试", "执行后端 unittest 与前端 Node Test。", "392/392、14/14 通过记录"],
            ["构建验证", "执行 Vite 生产构建和 Python 核心模块编译。", "1,650 模块构建成功、py_compile 成功"],
            ["配置与一致性检查", "检查迁移、环境变量模板、Git 空白错误和交付目录。", "迁移测试、diff check 和交付清单"],
            ["条件性验证", "第三方模型、Embedding、搜索和真实生产数据在部署环境复核。", "上线前联调记录"],
        ],
        [3.0, 7.7, 4.8],
        font_size=8.6,
    )
    add_heading(doc, "3.4 通过准则", 2)
    add_table(
        doc,
        ["编号", "准则", "判定规则"],
        [
            ["AC-01", "核心业务", "认证、检测、证据、评分、历史和报告主流程无阻断性缺陷。"],
            ["AC-02", "自动化测试", "既有后端与前端自动化测试全部通过，无失败或错误。"],
            ["AC-03", "可构建性", "前端生产构建和后端核心模块编译成功。"],
            ["AC-04", "数据设计", "8 个关系模型、外键/索引、迁移和 Chroma 派生索引职责明确。"],
            ["AC-05", "安全基线", "具备身份鉴别、角色授权、资源所有权、SSRF 防护、限流、秘密管理和审计。"],
            ["AC-06", "文档与交付", "交付物可识别、版本一致、内容可支持部署、使用、测试和维护。"],
            ["AC-07", "问题状态", "无未关闭的致命或严重问题；一般问题有责任、措施和复核节点。"],
        ],
        [1.8, 3.4, 10.3],
        font_size=8.7,
    )


def add_chapter_4(doc: Document) -> None:
    add_heading(doc, "4 验收环境与配置基线", 1)
    add_heading(doc, "4.1 实际验证环境", 2)
    add_table(
        doc,
        ["类别", "项目", "实际值/说明"],
        [
            ["操作系统", "验证主机", "Microsoft Windows 10 专业版，约 16 GB 内存"],
            ["后端运行时", "Python", "Python 3.14.2"],
            ["前端运行时", "Node.js / npm", "Node.js v22.22.3 / npm 10.9.8"],
            ["版本工具", "Git", "git 2.49.0.windows.1"],
            ["后端框架", "FastAPI/SQLAlchemy/Pydantic", "以 backend/requirements.txt 锁定范围为准"],
            ["前端框架", "Vue/Vite/Element Plus", "Vue 3.5、Vite 6、Element Plus 2.9 系列"],
            ["业务存储", "MySQL", "生产目标数据库；utf8mb4；结构以 Alembic head 为准"],
            ["向量存储", "Chroma", "knowledge_items 余弦集合；可从 MySQL 全量重建"],
            ["外部服务", "DeepSeek/DashScope/博查", "通过环境变量配置；本地自动化测试使用隔离或模拟"],
        ],
        [2.5, 4.1, 8.9],
        font_size=8.6,
    )
    add_heading(doc, "4.2 配置项与基线状态", 2)
    add_table(
        doc,
        ["配置项", "基线位置", "状态", "控制要求"],
        [
            ["后端源代码", "backend/app", "已验证", "归档前提交待确认变更并打版本标签"],
            ["前端源代码", "frontend/src", "已验证", "以成功构建快照为准"],
            ["数据库结构", "backend/alembic/versions", "已验证", "仅通过 Alembic 演进"],
            ["依赖", "requirements.txt / package-lock.json", "已识别", "生产安装使用受控锁文件/镜像"],
            ["秘密配置", "backend/.env（不入库）", "模板化", "不得提交真实口令、Token 或 API Key"],
            ["向量索引", "CHROMA_PATH", "可重建", "Provider 或维度变更后全量重建"],
            ["报告目录", "REPORT_DIR", "外置", "源代码目录外、最小权限、纳入备份"],
            ["验收文档", "开发管理目录", "齐备", "统一编号、版本和审批"],
        ],
        [3.1, 4.6, 2.4, 5.4],
        font_size=8.4,
    )
    add_heading(doc, "4.3 验收限制", 2)
    add_note(
        doc,
        "环境边界",
        "本轮证据充分覆盖代码结构、离线功能、接口行为、迁移、错误处理、前端逻辑和可构建性；未使用生产级 API Key 对第三方服务做配额、时延和真实内容质量验收，也未实施多机高并发压力测试。相关项目列入部署环境复核。",
        "warning",
    )


def add_chapter_5(doc: Document) -> None:
    add_heading(doc, "5 交付物审查", 1)
    add_heading(doc, "5.1 软件与工程资产", 2)
    add_table(
        doc,
        ["编号", "交付物", "位置/形式", "审查结论"],
        [
            ["D-01", "后端源代码", "backend/app", "分层清晰，包含 API、Service、CRUD、Model、Schema、Core、DB 与工具层"],
            ["D-02", "前端源代码", "frontend/src", "包含用户端、管理端、组件、路由、状态、API 与工具层"],
            ["D-03", "数据库迁移", "backend/alembic/versions", "具备可追踪版本，迁移测试通过"],
            ["D-04", "自动化测试", "backend/tests；frontend/src/**/*.test.js", "后端 392 项、前端 14 项全部通过"],
            ["D-05", "依赖与配置模板", "requirements.txt、package-lock.json、.env.example", "可支持环境重建，秘密与代码分离"],
            ["D-06", "部署与使用说明", "README.md、backend/README.md", "覆盖初始化、迁移、启动、构建和演示路径"],
            ["D-07", "报告模板", "backend/app/templates/reports", "支持按检测记录生成 PDF"],
            ["D-08", "知识图谱索引", "codebase-memory-mcp 项目索引", "验收前已增量更新至当前快照"],
        ],
        [1.5, 3.1, 4.1, 6.8],
        font_size=8.4,
    )
    add_heading(doc, "5.2 软件工程文档", 2)
    add_table(
        doc,
        ["序号", "文档", "编号", "状态"],
        [
            ["1", "可行性研究报告", "NCE-FS-001", "已交付"],
            ["2", "需求分析书", "NCE-SRS-001", "已交付"],
            ["3", "概要设计书（含数据库设计）", "NCE-HLD-001", "已交付"],
            ["4", "详细设计书", "NCE-LLD-001", "已交付"],
            ["5", "项目验收总结报告", DOC_NO, "本次交付"],
        ],
        [1.4, 7.4, 3.3, 3.4],
    )
    add_heading(doc, "5.3 交付物完整性结论", 2)
    add_body(
        doc,
        "交付物已覆盖软件源代码、数据库演进、依赖配置、测试、构建、部署、使用和软件工程文档。各项能够相互追踪：需求由概要/详细设计承接，程序和数据结构可在知识图谱中定位，测试覆盖主要业务与异常分支，验收报告汇总最终证据。",
    )


def add_chapter_6(doc: Document) -> None:
    add_heading(doc, "6 功能验收", 1)
    add_heading(doc, "6.1 用户端功能", 2)
    add_table(
        doc,
        ["编号", "验收项", "主要验证点", "结果"],
        [
            ["F-01", "注册、登录和会话", "唯一性、bcrypt、JWT、禁用用户、401/403、限流", "通过"],
            ["F-02", "新闻手工检测", "输入长度、游客/登录用户、限流、统一结果结构", "通过"],
            ["F-03", "URL 提取预览", "HTTP(S)、SSRF、重定向、体积/超时、正文和发布时间提取", "通过"],
            ["F-04", "本地 RAG 检索", "Embedding、Top-K、过滤、相似度格式化、异常重试", "通过"],
            ["F-05", "条件联网补证", "触发判定、检索规范化、去重合并、失败回退", "条件通过"],
            ["F-06", "LLM 分析与证据仲裁", "Prompt 边界、结构化契约、一次聚焦重试、失败降级", "条件通过"],
            ["F-07", "规则与综合评分", "规则命中、证据质量、分支公式、四级风险映射", "通过"],
            ["F-08", "结果展示", "分数、理由、风险点、证据质量、相似新闻、建议和免责声明", "通过"],
            ["F-09", "历史、详情与重评", "分页筛选、资源所有权、旧输入复用、新记录留痕", "通过"],
            ["F-10", "PDF 报告", "所有权校验、HTML 渲染、PDF 生成、路径安全、下载", "通过"],
            ["F-11", "公开高风险信息", "审核通过且允许公开、摘要字段、榜单/关键词/类别统计", "通过"],
        ],
        [1.4, 3.0, 8.7, 2.4],
        font_size=8.1,
    )
    add_body(
        doc,
        "“条件通过”表示代码路径、契约、降级和自动化测试已通过，但第三方服务真实响应质量、配额和网络时延需在正式部署环境以有效凭据复核；该条件不影响本地软件功能和容错逻辑验收。",
    )
    add_heading(doc, "6.2 管理端功能", 2)
    add_table(
        doc,
        ["编号", "验收项", "主要验证点", "结果"],
        [
            ["A-01", "用户管理", "列表、详情、启停、角色、最后管理员和自操作保护", "通过"],
            ["A-02", "检测管理", "全局列表、详情、删除和关联清理", "通过"],
            ["A-03", "知识库管理", "增删改查、单条向量化、状态可见、全量重建", "通过"],
            ["A-04", "Prompt 管理", "创建、修改、启停、默认模板、模板校验", "通过"],
            ["A-05", "高风险审核", "审核、公开、备注、状态组合和审计", "通过"],
            ["A-06", "报告管理", "列表、详情、受控下载和缺失文件处理", "通过"],
            ["A-07", "统计分析", "概览、趋势、风险/类别、关键词、活跃度、知识状态", "通过"],
            ["A-08", "系统日志", "分页、筛选、事件记录和管理员访问控制", "通过"],
        ],
        [1.5, 3.2, 8.4, 2.4],
        font_size=8.3,
    )
    add_heading(doc, "6.3 核心流程一致性核验", 2)
    add_body(
        doc,
        "codebase-memory-mcp 对 detect_news_credibility 的出向调用追踪确认：核心流程直接调用知识检索、条件联网搜索、DeepSeek 分析、证据仲裁、规则评分、风险等级映射、Prompt 读取和检测记录持久化；对 generate_detection_report 的追踪确认其执行所有权校验、上下文构造、HTML 渲染、PDF 转换、文件清理和报告记录保存。实现与概要/详细设计的职责划分一致。",
    )


def add_chapter_7(doc: Document) -> None:
    add_heading(doc, "7 数据库与数据验收", 1)
    add_heading(doc, "7.1 数据存储构成", 2)
    add_table(
        doc,
        ["数据对象", "职责", "关键关系/控制", "结论"],
        [
            ["users", "认证主体、角色和状态", "与检测、审核、日志关联；用户名/邮箱唯一", "通过"],
            ["detection_records", "新闻输入、评分、风险、审核和分析载荷", "用户可空；证据一对多；报告一对一", "通过"],
            ["evidence_matches", "有效证据快照和排序", "随检测记录级联删除；可回查知识条目", "通过"],
            ["knowledge_items", "知识事实源和向量同步状态", "分类/标签/风险/同步状态索引", "通过"],
            ["prompt_templates", "版本化 Prompt 内容和启停/默认状态", "模板类型和状态约束", "通过"],
            ["reports", "报告元数据与文件路径", "检测记录一对一；受所有权保护", "通过"],
            ["system_logs", "关键操作和审计事件", "按时间、用户、动作检索", "通过"],
            ["crawl_tasks", "定时采集执行结果", "记录状态、数量、错误和时间", "通过"],
            ["Chroma knowledge_items", "知识条目派生向量索引", "vector_id=knowledge:<id>；可重建", "通过"],
            ["REPORT_DIR", "生成的 HTML/PDF 文件", "源代码目录外；路径解析和授权下载", "通过"],
        ],
        [3.2, 4.4, 6.1, 1.8],
        font_size=8.1,
    )
    add_figure(doc, "fig4-7-database-er.png", "图 7-1 关系数据库实体联系", max_width=14.8, max_height=14.5)
    add_heading(doc, "7.2 完整性与一致性", 2)
    add_bullets(
        doc,
        [
            "关系主键统一采用大整数；外键明确 SET NULL、CASCADE 等删除语义；",
            "常用用户、风险、审核、公开、时间、类别、事实标签和同步状态字段建立索引；",
            "MySQL 是知识事实源，Chroma 是派生索引；双写失败通过 pending/failed/delete_failed 状态暴露并可补偿；",
            "检测保存与证据保存使用同一事务，失败时回滚，避免主记录和证据半成功；",
            "数据库结构以 Alembic head 为准，旧 SQL 仅保留历史追溯用途；",
            "切换 Embedding Provider 或维度后必须重建向量索引，避免维数不匹配和旧向量污染。",
        ],
    )
    add_heading(doc, "7.3 数据安全与备份要求", 2)
    add_table(
        doc,
        ["对象", "验收确认", "运行要求"],
        [
            ["密码与密钥", "密码散列、API Key 与 SECRET_KEY 不写入业务表或代码库", "使用秘密管理或受控环境变量，定期轮换"],
            ["业务数据库", "关系约束、迁移和访问层已实现", "定期全量+增量备份并演练恢复"],
            ["向量索引", "与知识 ID 关联且支持重建", "可按重建策略备份；恢复后校验集合数量和维度"],
            ["报告文件", "通过受控路径和授权接口访问", "纳入备份、保留周期和安全删除策略"],
            ["审计日志", "记录关键管理操作", "限制访问，保留周期与业务合规要求一致"],
        ],
        [3.0, 7.0, 5.5],
        font_size=8.5,
    )


def add_chapter_8(doc: Document) -> None:
    add_heading(doc, "8 非功能质量验收", 1)
    add_heading(doc, "8.1 质量属性评价", 2)
    add_table(
        doc,
        ["质量属性", "验收观察", "结论"],
        [
            ["功能适合性", "用户端、管理端和核心检测链路覆盖目标业务；异常与降级有明确结果。", "通过"],
            ["可靠性", "外部搜索失败回退、模型失败规则降级、Chroma 操作清缓存重试、数据库失败回滚。", "通过"],
            ["性能效率", "前端路由懒加载、API 分页、Top-K/正文/响应体/超时限制；未实施专项并发压测。", "基本通过"],
            ["兼容性", "标准浏览器 Web 架构、REST/JSON 接口、Windows 验证环境构建成功。", "通过"],
            ["易用性", "提供输入校验、加载/空状态、风险标签、证据说明、管理脚手架和统一错误提示。", "通过"],
            ["安全性", "JWT、bcrypt、角色与所有权、限流、SSRF、防路径穿越、秘密分离和审计。", "通过"],
            ["维护性", "前后端分层、Schema 契约、迁移、测试、日志、配置和五类文档完整。", "通过"],
            ["可移植性", "依赖清单和环境变量驱动；数据与文件路径可配置；未提供容器编排模板。", "基本通过"],
        ],
        [3.0, 10.3, 2.2],
        font_size=8.3,
    )
    add_heading(doc, "8.2 安全控制验收", 2)
    add_table(
        doc,
        ["控制域", "已实现措施", "证据/结论"],
        [
            ["身份鉴别", "bcrypt 密码散列、JWT 签发/校验、禁用状态实时检查", "认证测试通过"],
            ["访问控制", "user/admin 角色依赖、个人资源所有权、公开接口最小字段", "权限与越权测试通过"],
            ["输入安全", "Pydantic 类型/长度、文本清洗、ORM 参数化、extra 字段拒绝", "接口和 Schema 测试通过"],
            ["网络安全", "仅 HTTP(S)、DNS/IP 私网拦截、重定向重验、超时和响应体限制", "网页提取安全测试通过"],
            ["资源保护", "登录/检测限流、外部调用超时、Top-K 和分页上限", "限流与边界测试通过"],
            ["秘密管理", ".env 注入、生产 SECRET_KEY 启动校验、真实 Key 不入库", "启动校验测试通过"],
            ["文件安全", "报告根目录外置、路径解析后校验仍在允许目录、授权下载", "报告服务测试通过"],
            ["安全审计", "管理员关键操作写 system_logs，可分页筛选", "日志 API/Service 测试通过"],
        ],
        [2.8, 8.2, 4.5],
        font_size=8.2,
    )
    add_heading(doc, "8.3 性能与容量评价", 2)
    add_body(
        doc,
        "本轮未设定并验证生产并发量、P95 时延或吞吐量等定量 SLA，因此不对高并发能力作超出证据的结论。代码已具备分页、Top-K、文本长度、网页响应体、网络超时、外部调用条件触发和前端懒加载等基础控制。正式上线前应使用代表性数据和有效第三方账号执行容量、并发、长稳和故障注入测试。",
    )


def add_chapter_9(doc: Document) -> None:
    add_heading(doc, "9 验收测试执行与结果", 1)
    add_heading(doc, "9.1 执行汇总", 2)
    add_table(
        doc,
        ["验证项", "执行命令/方式", "结果", "判定"],
        [
            ["后端自动化测试", "python -m unittest discover -s tests -p test_*.py", "392 项；392 通过；0 失败/错误；32.584 s（精简复核）", "通过"],
            ["前端自动化测试", "npm test", "14 项；14 通过；0 失败；865.224 ms", "通过"],
            ["前端生产构建", "npm run build", "Vite 构建成功；1,650 模块转换；51.43 s", "通过"],
            ["后端核心模块编译", "python -m py_compile ...", "主应用、配置、迁移、检测、LLM、报告模块编译成功", "通过"],
            ["Git 空白错误检查", "git diff --check", "无空白错误；仅提示 LF/CRLF 转换", "通过"],
            ["知识图谱增量索引", "codebase-memory-mcp index_repository", "3,645 节点、10,892 关系；状态 indexed", "通过"],
        ],
        [3.0, 5.2, 5.6, 1.7],
        font_size=8.0,
    )
    add_note(
        doc,
        "总体测试结果",
        "后端与前端合计 406 项自动化测试全部通过，失败数为 0；生产构建和核心模块编译均成功。测试输出中的异常日志多为刻意模拟 Chroma、模型、网络和事务失败的负向用例，其断言结果均通过。",
        "success",
    )
    add_heading(doc, "9.2 测试覆盖域", 2)
    add_table(
        doc,
        ["覆盖域", "代表性验证内容"],
        [
            ["认证与权限", "注册登录、密码、JWT、禁用用户、管理员权限、资源所有权和限流"],
            ["检测与评分", "正常检测、无证据、模型失败、仲裁契约、重试、降级、分数和风险边界"],
            ["知识与向量", "Embedding、Chroma 查询/写入/删除重试、知识同步、失败状态和重建"],
            ["联网与网页", "搜索触发、查询构造、去重、URL 规范化、SSRF、重定向、正文/时间提取"],
            ["数据与迁移", "ORM 模型、时间戳、CRUD、历史筛选、Alembic 迁移和种子数据幂等"],
            ["报告与高风险", "报告生成/下载/所有权、公开条件、审核、榜单、关键词和分类"],
            ["管理与统计", "用户、报告、日志、Prompt、统计查询和知识状态"],
            ["前端工具", "游客结果缓存、证据质量状态、分数/百分比、发布时间精度与格式"],
        ],
        [4.0, 11.5],
        font_size=8.7,
    )
    add_heading(doc, "9.3 结果解释", 2)
    add_body(
        doc,
        "自动化测试使用隔离数据库、临时 Chroma 目录和模拟外部依赖验证确定性逻辑。测试日志中的 DeprecationWarning、ResourceWarning，以及刻意触发的“vector sync failed”“arbitration retry failed”等信息不代表测试失败；它们分别属于依赖升级提示、测试资源清理提示和负向场景日志。",
    )


def add_chapter_10(doc: Document) -> None:
    add_heading(doc, "10 问题、偏差与剩余风险", 1)
    add_heading(doc, "10.1 验收问题汇总", 2)
    add_table(
        doc,
        ["编号", "问题/观察项", "级别", "影响", "处置建议"],
        [
            ["OBS-01", "第三方模型、Embedding 和搜索未在生产凭据下进行实网质量/配额验收", "一般", "无法据此承诺真实网络时延和内容质量", "上线前执行真实凭据联调并留存样本"],
            ["OBS-02", "前端 ECharts 产物约 519 kB，Vite 给出大分块警告", "一般", "弱网首屏或统计页加载可能变慢", "按路由/图表动态导入并设置性能预算"],
            ["OBS-03", "测试运行出现 Python/依赖弃用警告", "提示", "未来升级 Python 可能需修改 API 用法", "纳入依赖升级计划并在升级后回归"],
            ["OBS-04", "部分测试出现临时 SQLite 连接 ResourceWarning", "一般", "不影响业务结论，但降低测试洁净度", "统一关闭 Session/Engine 和临时数据库"],
            ["OBS-05", "限流和调度状态默认驻留单进程", "一般", "多实例部署时配额不共享、任务可能重复", "扩容前改用 Redis/网关限流和分布式任务锁"],
            ["OBS-06", "本地 Chroma 与报告目录不适合无状态多实例共享", "一般", "横向扩展、备份和故障切换复杂", "生产扩容时迁移服务化向量库和对象存储"],
            ["OBS-07", "验收工作区含待提交变更", "一般", "若不冻结，归档版本与验收证据可能漂移", "签署前提交、标记版本并记录提交 SHA"],
        ],
        [1.5, 5.7, 1.7, 3.7, 3.6],
        font_size=7.8,
    )
    add_heading(doc, "10.2 问题分级与结论影响", 2)
    add_table(
        doc,
        ["级别", "定义", "本次数量", "对验收的影响"],
        [
            ["致命", "系统不可安装、核心数据破坏或存在不可接受安全风险", "0", "必须拒绝验收"],
            ["严重", "核心业务不可用、无替代路径或高概率造成错误结论", "0", "整改并复测后方可验收"],
            ["一般", "局部质量、工程或扩展性问题，有可行规避或后续方案", "6", "不阻断本期验收，纳入改进计划"],
            ["提示", "依赖升级、优化或洁净度建议", "1", "记录并跟踪"],
        ],
        [2.2, 8.0, 2.1, 3.2],
        font_size=8.5,
    )
    add_heading(doc, "10.3 风险接受原则", 2)
    add_body(
        doc,
        "上述观察项均不破坏当前单实例课程设计/小规模部署场景的核心功能和数据正确性。验收组接受这些剩余风险的前提是：明确系统辅助评估定位；正式部署前完成真实凭据联调；归档时冻结代码基线；生产扩容前完成共享限流、调度锁、向量服务和对象存储改造。",
    )


def add_chapter_11(doc: Document) -> None:
    add_heading(doc, "11 验收结论与意见", 1)
    add_heading(doc, "11.1 准则符合性", 2)
    add_table(
        doc,
        ["准则", "符合性说明", "结果"],
        [
            ["AC-01 核心业务", "认证、检测、证据、评分、历史、报告和管理主流程具备实现与测试证据。", "符合"],
            ["AC-02 自动化测试", "后端 392 项、前端 14 项全部通过。", "符合"],
            ["AC-03 可构建性", "Vite 生产构建、Python 核心模块编译成功。", "符合"],
            ["AC-04 数据设计", "8 个关系模型、迁移、索引、外键及 Chroma 协作机制明确。", "符合"],
            ["AC-05 安全基线", "认证、授权、SSRF、限流、秘密、路径和审计控制已实现并测试。", "符合"],
            ["AC-06 文档交付", "五类核心软件工程文档及部署/使用资料齐备。", "符合"],
            ["AC-07 问题状态", "无致命和严重未关闭问题；一般问题均有建议和复核节点。", "符合"],
        ],
        [3.0, 10.1, 2.4],
        font_size=8.5,
    )
    add_heading(doc, "11.2 综合结论", 2)
    add_note(
        doc,
        "建议验收通过",
        "项目已完成约定范围内的分析、设计、实现、测试和文档交付；软件结构、数据库设计、核心功能、安全基线、自动化测试和生产构建满足本期验收准则。未发现致命或严重缺陷，建议通过项目验收并进入归档、移交和试运行阶段。",
        "success",
    )
    add_heading(doc, "11.3 附带要求", 2)
    add_numbered(
        doc,
        [
            "验收签署前将当前认可的工作区变更提交到版本库，记录最终提交 SHA，并创建与 V1.0 对应的发布标签。",
            "在正式部署环境使用有效 DeepSeek、DashScope 和博查凭据完成实网联调，记录成功率、时延、配额和代表性样本。",
            "生产上线前设置强 SECRET_KEY、最小权限数据库账户、明确 CORS 域名、TLS、备份、日志轮转和目录权限。",
            "将前端大分块、测试资源清理和依赖弃用警告纳入维护计划；不迟于下一次版本升级完成复核。",
            "对外展示必须保留辅助评估免责声明，高风险结论应由人工结合权威来源复核。",
        ],
    )
    add_heading(doc, "11.4 验收签署", 2)
    add_table(
        doc,
        ["验收角色", "验收意见", "姓名", "签字", "日期"],
        [
            ["业务/需求代表", "□同意通过  □整改后通过  □不通过", "", "", ""],
            ["技术评审代表", "□同意通过  □整改后通过  □不通过", "", "", ""],
            ["测试代表", "□同意通过  □整改后通过  □不通过", "", "", ""],
            ["验收组长", "□同意通过  □整改后通过  □不通过", "", "", ""],
            ["批准人", "□批准结项  □暂缓结项", "", "", ""],
        ],
        [3.0, 6.1, 2.2, 2.2, 2.0],
        font_size=8.4,
    )


def add_chapter_12(doc: Document) -> None:
    add_heading(doc, "12 移交、运行与后续工作", 1)
    add_heading(doc, "12.1 移交要求", 2)
    add_table(
        doc,
        ["移交项", "接收检查", "责任方"],
        [
            ["源代码与版本", "提交 SHA、标签、分支、变更说明可追踪", "开发/配置管理"],
            ["数据库", "迁移版本、初始化方法、备份和恢复步骤明确", "开发/运维"],
            ["配置与秘密", "模板齐全；真实秘密通过安全渠道交接", "运维/安全"],
            ["外部服务", "账号、配额、告警、续费和故障联系人明确", "项目负责人/运维"],
            ["数据目录", "Chroma、报告和日志目录权限、容量和备份策略明确", "运维"],
            ["文档", "五类文档、README、使用和部署资料统一归档", "配置管理"],
            ["问题清单", "OBS-01～OBS-07 纳入跟踪，明确责任和计划", "项目负责人"],
        ],
        [3.2, 9.2, 3.1],
        font_size=8.6,
    )
    add_heading(doc, "12.2 试运行观察指标", 2)
    add_table(
        doc,
        ["类别", "建议指标", "异常信号"],
        [
            ["API", "请求量、P95 时延、4xx/5xx/429、健康检查", "持续 5xx、时延突增、限流异常"],
            ["检测业务", "检测成功率、模型降级率、仲裁重试率、无证据率", "降级/重试显著上升、结果字段缺失"],
            ["外部服务", "DeepSeek/DashScope/博查时延、错误率、401/429 和配额", "持续超时、鉴权失败、配额耗尽"],
            ["MySQL", "连接、慢查询、容量、备份成功率", "连接耗尽、备份失败、增长异常"],
            ["Chroma", "集合数量、查询/写入失败、failed 状态、磁盘", "维度不匹配、同步失败积压"],
            ["报告", "生成/下载成功率、目录容量、缺失文件", "生成失败、路径异常、容量逼近阈值"],
            ["安全", "登录失败、SSRF 拒绝、管理员操作和异常来源", "高频失败、越权尝试、异常公开操作"],
        ],
        [2.8, 8.2, 4.5],
        font_size=8.3,
    )
    add_heading(doc, "12.3 后续改进路线", 2)
    add_table(
        doc,
        ["优先级", "改进项", "完成标志"],
        [
            ["P0/上线前", "冻结 V1.0、真实服务联调、强秘密、CORS/TLS、备份恢复演练", "形成发布记录和上线检查表"],
            ["P1/近期", "消除测试资源警告、处理依赖弃用、优化前端图表分块", "测试无资源泄漏警告，构建性能预算达标"],
            ["P1/近期", "补充代表性性能和长稳测试", "形成 P95/吞吐/容量基线及瓶颈结论"],
            ["P2/扩容前", "共享限流、分布式调度锁、服务化向量库、对象存储", "多实例无重复任务和状态漂移"],
            ["P2/质量", "建立人工复核样本集和模型效果评估", "形成准确性、召回率、证据相关性和偏差报告"],
        ],
        [2.5, 9.1, 3.9],
        font_size=8.4,
    )
    add_heading(doc, "12.4 项目经验总结", 2)
    add_bullets(
        doc,
        [
            "RAG 的价值不只是检索候选，而是建立证据标识、仲裁契约和可追踪结果；",
            "对外部模型输出实施服务端校验和确定性降级，可以显著降低不可控响应对业务的影响；",
            "MySQL 作为事实源、Chroma 作为可重建派生索引，使双存储职责和恢复路径更清晰；",
            "将权限、SSRF、限流、路径和秘密校验纳入自动化测试，比只在部署说明中声明更可靠；",
            "以知识图谱、测试结果和构建结果共同支撑验收，能够把文档结论落到可复现证据。",
        ],
    )
    add_heading(doc, "12.5 开发工作综合评价", 2)
    add_table(
        doc,
        ["评价维度", "评价"],
        [
            ["技术方法", "前后端分离、分层后端、事实库与派生向量分离、受约束 LLM 契约和确定性降级方法适合本项目。"],
            ["工程工具", "Git、Alembic、unittest、Node Test、Vite 和 codebase-memory-mcp 共同支持版本、迁移、测试、构建和追踪。"],
            ["产品质量", "406 项自动化测试全部通过，核心安全与异常分支具备验证；真实外部服务和高并发指标仍需生产环境补充。"],
            ["文档质量", "五类核心软件工程文档形成统一编号、版本、基线和版式，并建立需求—设计—测试—验收追踪。"],
            ["管理改进", "后续应强化计划/工时/费用台账、发布标签、性能预算和线上服务等级记录。"],
        ],
        [3.3, 12.2],
        font_size=8.6,
    )


def add_appendices(doc: Document) -> None:
    add_heading(doc, "附录 A 验收检查表", 1)
    add_table(
        doc,
        ["序号", "检查项", "结果", "验收人/日期"],
        [
            ["1", "验收对象、范围和版本已确认", "□通过 □不通过", ""],
            ["2", "五类软件工程文档齐备且版本一致", "□通过 □不通过", ""],
            ["3", "后端 392 项自动化测试全部通过", "□通过 □不通过", ""],
            ["4", "前端 14 项自动化测试全部通过", "□通过 □不通过", ""],
            ["5", "前端生产构建和后端核心模块编译成功", "□通过 □不通过", ""],
            ["6", "数据库模型、迁移、索引和备份要求明确", "□通过 □不通过", ""],
            ["7", "认证、授权、SSRF、限流、秘密和审计控制满足要求", "□通过 □不通过", ""],
            ["8", "外部服务部署联调责任和条件已确认", "□通过 □不通过", ""],
            ["9", "遗留问题已分级并纳入跟踪", "□通过 □不通过", ""],
            ["10", "最终提交 SHA、V1.0 标签和归档介质已确认", "□通过 □不通过", ""],
        ],
        [1.3, 8.8, 3.0, 2.4],
        font_size=8.5,
    )

    add_heading(doc, "附录 B 需求—设计—测试—验收追踪矩阵", 1)
    add_table(
        doc,
        ["需求域", "设计/实现模块", "验证证据", "验收结果"],
        [
            ["认证与权限", "auth、security、deps、user store/router", "auth/admin user/CORS/secret tests", "通过"],
            ["新闻检测", "detect API、detection_service、schemas", "detect/history/CRUD tests", "通过"],
            ["RAG 与向量", "knowledge_service、embedding、chroma", "embedding/chroma/RAG/knowledge sync tests", "通过"],
            ["LLM 与 Prompt", "llm_service、prompt_service、validator", "LLM/prompt/arbitration tests", "条件通过"],
            ["联网与 URL", "bocha_client、web_search、fetcher", "web search/fetcher/SSRF tests", "条件通过"],
            ["评分与风险", "rule_score、risk_level、high_risk", "rule/risk/high-risk tests", "通过"],
            ["历史与报告", "detection CRUD、report_service", "history/report API/service tests", "通过"],
            ["管理与统计", "admin APIs、statistics、system logs", "admin/statistics/log tests", "通过"],
            ["数据库演进", "models、Alembic、migration guard", "migration/model timestamp tests", "通过"],
            ["前端展示", "views、components、utils", "14 项 Node Test + Vite build", "通过"],
        ],
        [3.0, 5.0, 5.5, 2.0],
        font_size=8.0,
    )

    add_heading(doc, "附录 C 验收命令与证据摘要", 1)
    add_table(
        doc,
        ["编号", "命令/工具", "证据摘要"],
        [
            ["E-01", "codebase-memory-mcp index_repository", "增量索引成功；3,645 节点、10,892 关系"],
            ["E-02", "get_architecture/search_graph/trace_path/get_code_snippet", "确认 8 个关系模型、37 个后端测试文件及核心调用链"],
            ["E-03", "python -m unittest discover -s tests -p test_*.py", "Ran 392 tests；OK"],
            ["E-04", "npm test", "14 tests；pass 14；fail 0"],
            ["E-05", "npm run build", "1,650 modules transformed；built successfully"],
            ["E-06", "python -m py_compile ...", "核心模块编译成功"],
            ["E-07", "git diff --check", "无空白错误；存在行尾格式提示"],
        ],
        [1.5, 7.6, 6.4],
        font_size=8.5,
    )

    add_heading(doc, "附录 D 验收结论单", 1)
    add_table(
        doc,
        ["项目", "填写内容"],
        [
            ["项目名称", PROJECT_TITLE],
            ["验收版本", VERSION],
            ["验收日期", "____________________"],
            ["验收地点", "____________________"],
            ["验收结论", "□通过  □整改后通过  □不通过"],
            ["主要意见", "________________________________________________________________"],
            ["整改要求", "________________________________________________________________"],
            ["验收组长签字", "____________________"],
            ["批准人签字", "____________________"],
        ],
        [4.2, 11.3],
        first_col_bold=True,
    )


def add_main_content(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, header_footer=True, page_format="decimal", page_start=1)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_chapter_6(doc)
    add_chapter_7(doc)
    add_chapter_8(doc)
    add_chapter_9(doc)
    add_chapter_10(doc)
    add_chapter_11(doc)
    add_chapter_12(doc)
    add_appendices(doc)


def set_document_metadata(doc: Document) -> None:
    props = doc.core_properties
    props.title = f"《{PROJECT_TITLE}》{DOC_TITLE}"
    props.subject = "软件项目验收、测试结果、交付物、问题与移交总结"
    props.author = "项目开发组"
    props.keywords = "项目验收, 软件工程, RAG, 大语言模型, 新闻真伪鉴别, 测试总结"
    props.comments = f"文档编号 {DOC_NO}；版本 {VERSION}；基线 {BASELINE}"

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_main_content(doc)
    set_document_metadata(doc)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
