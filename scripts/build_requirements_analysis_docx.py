from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import build_feasibility_report_docx as base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = Path(__file__).resolve().parents[1] / ".artifacts" / "docs"
OUTPUT_PATH = OUTPUT_DIR / "02_《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》需求分析书.docx"

PROJECT_TITLE = "基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计"
DOC_TITLE = "需求分析书"
DOC_NO = "NCE-SRS-001"
VERSION = "V1.0"
DATE_TEXT = "2026年6月20日"

MODULE_FIGURE = PROJECT_ROOT / "docs/thesis/figures/chapter4/fig4-3-functional-modules.png"
FLOW_FIGURE = PROJECT_ROOT / "docs/thesis/figures/chapter4/fig4-4-detection-overall-flow.png"
SCORE_FIGURE = PROJECT_ROOT / "docs/thesis/figures/chapter4/fig4-6-scoring-risk-classification.png"
ER_FIGURE = PROJECT_ROOT / "docs/thesis/figures/chapter4/fig4-7-database-er.png"
USER_FIGURE = PROJECT_ROOT / "docs/thesis/figures/chapter4/fig4-9-user-page-structure.png"
ADMIN_FIGURE = PROJECT_ROOT / "docs/thesis/figures/chapter4/fig4-10-admin-page-structure.png"


# Priority: M=Must, S=Should, C=Could.
# Status: I=implemented baseline, A=acceptance target, T=to be determined.
# Verification: T=test, D=demonstration, I=inspection, A=analysis.
REQ_GROUPS: list[tuple[str, str, list[tuple[str, str, str, str, str]]]] = [
    (
        "4.1",
        "用户注册、认证与授权",
        [
            ("FR-AUTH-001", "系统应允许游客注册普通用户账号；用户名长度3～50个字符，密码长度6～128个字符，邮箱可选且格式有效。", "M", "I", "T"),
            ("FR-AUTH-002", "系统应保证用户名唯一；填写邮箱时还应保证邮箱唯一，并对并发注册冲突返回明确错误。", "M", "I", "T"),
            ("FR-AUTH-003", "系统不得以明文保存密码，应使用带盐的密码哈希进行校验。", "M", "I", "I/T"),
            ("FR-AUTH-004", "系统应按客户端IP限制注册请求，基线为每60秒最多2次，超限返回HTTP 429。", "M", "I", "T"),
            ("FR-AUTH-005", "系统应允许活动账号使用用户名和密码登录，并返回Bearer类型JWT及最小用户信息。", "M", "I", "T"),
            ("FR-AUTH-006", "系统应按客户端IP限制登录请求，基线为每60秒最多5次；失败登录应记录审计事件。", "M", "I", "T"),
            ("FR-AUTH-007", "系统应拒绝密码错误的登录并返回401，拒绝disabled账号登录并返回403。", "M", "I", "T"),
            ("FR-AUTH-008", "登录用户应能通过/api/auth/me获取当前账号信息；过期、伪造或已禁用账号的令牌不得通过认证。", "M", "I", "T"),
            ("FR-AUTH-009", "用户端受保护页面应在未登录时跳转登录页并保留redirect；管理员页面还应检查admin角色。", "M", "I", "D/T"),
            ("FR-AUTH-010", "后端每个管理接口均应执行管理员依赖校验，不能只依赖前端路由守卫。", "M", "I", "I/T"),
            ("FR-AUTH-011", "用户退出登录后，前端应清除本地令牌和用户缓存，并阻止继续访问受保护页面。", "M", "I", "D/T"),
        ],
    ),
    (
        "4.2",
        "新闻输入、链接预览与检测提交",
        [
            ("FR-DET-001", "系统应支持手工输入新闻标题、正文、类别、来源名称、来源链接、发布时间及联网搜索开关。", "M", "I", "D/T"),
            ("FR-DET-002", "新闻标题去除首尾空白后不得少于4个字符且不得超过255个字符。", "M", "I", "T"),
            ("FR-DET-003", "新闻正文去除首尾空白后不得少于20个字符；无效输入应返回统一的422校验响应。", "M", "I", "T"),
            ("FR-DET-004", "检测请求应拒绝未声明的额外字段，防止客户端静默提交未受控数据。", "M", "I", "T"),
            ("FR-DET-005", "系统应允许游客提交检测；游客结果只通过本次响应/短期前端缓存展示，不形成可访问的个人历史。", "M", "I", "D/T"),
            ("FR-DET-006", "登录用户提交检测时，系统应把检测记录关联到当前用户。", "M", "I", "T"),
            ("FR-DET-007", "系统应按客户端IP限制检测与链接提取请求，基线为每60秒最多3次，参数可配置。", "M", "I", "T"),
            ("FR-DET-008", "系统应提供URL提取预览接口，仅接受http://或https://链接，最大长度2048字符。", "S", "I", "T"),
            ("FR-DET-009", "URL提取应默认拒绝私有、回环、链路本地和保留网络地址，并在每次重定向后重新校验目标。", "M", "I", "T"),
            ("FR-DET-010", "URL提取应限制重定向次数、下载字节数和超时时间；失败或无法提取标题/正文时返回422和可操作提示。", "M", "I", "T"),
            ("FR-DET-011", "链接提取结果应作为预览回填标题、正文、来源与发布时间；用户确认或编辑后方可提交检测。", "S", "I", "D/T"),
            ("FR-DET-012", "发布时间应携带date或datetime精度，字段值与精度必须一致；系统不得为只有日期的值伪造午夜时间。", "S", "I", "T"),
            ("FR-DET-013", "系统应允许用户关闭本次联网搜索；全局开关关闭时不得发起联网搜索。", "M", "I", "T"),
            ("FR-DET-014", "知识检索失败应返回503并说明证据检索失败；一般检测业务错误应返回400，不得返回伪造成功结果。", "M", "I", "T"),
            ("FR-DET-015", "检测成功后应记录用户、检测ID、风险等级、最终分和客户端IP等审计信息，日志失败不得阻断主业务。", "S", "I", "T"),
        ],
    ),
    (
        "4.3",
        "RAG检索、联网证据与证据仲裁",
        [
            ("FR-RAG-001", "系统应将新闻标题和正文清洗并组合为RAG查询文本。", "M", "I", "T"),
            ("FR-RAG-002", "正式环境应使用语义Embedding；hash向量仅可用于无密钥的本地流程演示，不得宣称具有语义检索能力。", "M", "I", "I/T"),
            ("FR-RAG-003", "系统应从Chroma召回最多10条本地候选证据，并用knowledge_id回查MySQL业务数据。", "M", "I", "T"),
            ("FR-RAG-004", "系统应过滤已失效或无法回查的向量条目，并规范化标题、摘要、来源、标签和相似度。", "M", "I", "T"),
            ("FR-RAG-005", "只有本地证据数量或相似度不满足配置条件且用户允许联网时，系统才应触发联网补充。", "M", "I", "T"),
            ("FR-RAG-006", "联网搜索失败时系统应记录警告并继续使用本地候选，不得因此使整个检测崩溃。", "M", "I", "T"),
            ("FR-RAG-007", "系统应合并本地证据和网络证据，并按标题等规则去重；每条候选必须有稳定candidate_id和source_type。", "M", "I", "T"),
            ("FR-RAG-008", "提交给模型的候选顺序应来源中立且可复现，不得以来源类型、相似度或原始召回顺序暗示模型排序。", "M", "I", "A/I"),
            ("FR-RAG-009", "模型证据仲裁应输出ranked_evidence、rejected_evidence、candidate_id、相关性、质量、立场和理由。", "M", "I", "T"),
            ("FR-RAG-010", "系统应校验证据仲裁中的candidate_id唯一性与存在性，分数范围为0～100，立场仅允许support/contradict/neutral，理由不得为空。", "M", "I", "T"),
            ("FR-RAG-011", "模型应输出coverage、consistency、score和assessment；服务端应按0.6×coverage＋0.4×consistency重算质量分。", "M", "I", "T"),
            ("FR-RAG-012", "首次证据契约无效时，系统应执行一次聚焦仲裁重试；重试后仍无效应标记retry_exhausted并保存错误摘要。", "M", "I", "T"),
            ("FR-RAG-013", "未经成功仲裁的候选不得参与证据评分；被拒绝证据不得进入最终有效证据表。", "M", "I", "T"),
            ("FR-RAG-014", "系统应分别记录本地知识库和联网证据是否存在有效匹配，供结果解释与质量分析使用。", "S", "I", "T"),
        ],
    ),
    (
        "4.4",
        "评分、风险分级与结果解释",
        [
            ("FR-SCORE-001", "系统应输出0～100范围的llm_score、evidence_score、rule_score和final_score。", "M", "I", "T"),
            ("FR-SCORE-002", "规则评分应检查来源缺失、夸张表达、情绪化词、绝对化表述及与有效证据冲突，并输出命中规则。", "M", "I", "T"),
            ("FR-SCORE-003", "有效证据存在且质量有效时，最终分应为0.5×LLM分＋0.3×证据质量分＋0.2×规则分。", "M", "I", "T"),
            ("FR-SCORE-004", "无有效证据但LLM可用时，最终分应为0.6×LLM分＋0.4×规则分。", "M", "I", "T"),
            ("FR-SCORE-005", "LLM服务不可用时，未仲裁候选不得参与评分，最终分应退化为规则分并明确标注降级。", "M", "I", "T"),
            ("FR-SCORE-006", "最终分≥80为可信新闻，60～79.99为存疑信息，40～59.99为疑似谣言，<40为高风险谣言。", "M", "I", "T"),
            ("FR-SCORE-007", "final_score<40或风险等级为高风险谣言时，系统应自动标记is_high_risk=true。", "M", "I", "T"),
            ("FR-SCORE-008", "结果应包含判断结论、原因、风险点、关键词、建议、代理步骤、有效证据、排除证据和相似新闻。", "M", "I", "D/T"),
            ("FR-SCORE-009", "结果应包含免责声明，说明自动评分仅用于辅助核验，不能替代权威事实认定或专业判断。", "M", "I", "I/D"),
            ("FR-SCORE-010", "系统应保存analysis_contract_version、仲裁状态、质量状态、尝试次数和错误摘要，以支持旧记录兼容与追溯。", "S", "I", "T"),
        ],
    ),
    (
        "4.5",
        "检测历史、详情与重新评估",
        [
            ("FR-HIS-001", "登录用户应能分页查询自己的检测历史，默认页码1、每页20、每页最多100条。", "M", "I", "T"),
            ("FR-HIS-002", "历史查询应支持风险等级和新闻标题关键词筛选，并返回总数、页码、每页数量及列表。", "M", "I", "T"),
            ("FR-HIS-003", "普通用户只能查看自己的检测详情；记录不存在或不属于当前用户时返回404，避免泄露记录存在性。", "M", "I", "T"),
            ("FR-HIS-004", "管理员应能通过管理接口查看任意检测详情。", "M", "I", "T"),
            ("FR-HIS-005", "详情应返回持久化的有效证据、候选证据、排除证据、证据质量、来源信息和仲裁元数据。", "M", "I", "D/T"),
            ("FR-HIS-006", "登录用户应能对自己的历史记录发起重新评估；原记录保持不可变并生成新的检测记录。", "S", "I", "T"),
            ("FR-HIS-007", "重新评估应使用原输入和当前检索/模型契约，重新执行证据检索和分析，而非复制旧分数。", "M", "I", "T"),
            ("FR-HIS-008", "不符合当前输入约束的旧记录不得重新评估，应返回422和明确原因。", "S", "I", "T"),
            ("FR-HIS-009", "管理员应能分页筛选全站检测记录、查看详情并删除指定记录；删除应处理关联证据和报告。", "M", "I", "T"),
        ],
    ),
    (
        "4.6",
        "报告生成与下载",
        [
            ("FR-REP-001", "登录用户应能为自己的检测生成报告；管理员可为任意检测生成报告。", "M", "I", "T"),
            ("FR-REP-002", "报告内容应读取已保存的检测和证据，不应重新调用LLM而改变结论。", "M", "I", "I/T"),
            ("FR-REP-003", "同一检测只保留一条报告记录；重新生成应替换文件并保持数据库唯一约束。", "M", "I", "T"),
            ("FR-REP-004", "报告文件必须位于REPORT_DIR配置目录内，数据库保存受控路径；越界路径应拒绝下载。", "M", "I", "T"),
            ("FR-REP-005", "报告记录或PDF文件不存在时返回404；渲染或转换失败时返回明确500错误。", "M", "I", "T"),
            ("FR-REP-006", "管理员应能分页筛选、查看详情并下载全站报告。", "S", "I", "D/T"),
        ],
    ),
    (
        "4.7",
        "公开高风险信息",
        [
            ("FR-PUB-001", "公开接口仅应返回is_high_risk=true、review_status=approved且is_public=true的记录。", "M", "I", "T"),
            ("FR-PUB-002", "公开列表应支持分页、关键词、类别、风险等级和时间范围筛选。", "S", "I", "T"),
            ("FR-PUB-003", "公开响应不得包含完整正文、管理员备注、审核人、内部状态和分析错误等内部字段。", "M", "I", "T"),
            ("FR-PUB-004", "公开排名应在已审核公开集合内按最终分升序、审核时间倒序排列。", "S", "I", "T"),
            ("FR-PUB-005", "公开关键词和类别分布只能基于已审核公开集合计算。", "M", "I", "T"),
            ("FR-PUB-006", "用户端应明确说明高风险列表是辅助评估案例，不得把系统结论表述为权威通报。", "M", "A", "I/D"),
        ],
    ),
    (
        "4.8",
        "管理员用户与检测管理",
        [
            ("FR-ADM-001", "管理员应能按关键词、角色和状态分页查询用户，返回检测数量但不得返回密码哈希或令牌。", "M", "I", "T"),
            ("FR-ADM-002", "管理员应能查看用户详情和该用户的检测记录。", "M", "I", "D/T"),
            ("FR-ADM-003", "启用操作应幂等；禁用用户后，其登录、令牌恢复和受保护接口访问均应被拒绝。", "M", "I", "T"),
            ("FR-ADM-004", "系统不得允许管理员禁用或降级当前登录账号。", "M", "I", "T"),
            ("FR-ADM-005", "系统不得禁用或降级最后一个可用管理员；冲突应返回409。", "M", "I", "T"),
            ("FR-ADM-006", "角色只允许user和admin，状态只允许active和disabled。", "M", "I", "T"),
            ("FR-ADM-007", "为保留历史检测记录，系统不得开放用户物理删除接口。", "M", "I", "I/T"),
            ("FR-ADM-008", "关键用户状态和角色变更应记录管理员、对象、动作、结果和IP。", "M", "I", "T"),
        ],
    ),
    (
        "4.9",
        "知识库与向量索引管理",
        [
            ("FR-KB-001", "管理员应能按关键词、类别、真实性标签、风险等级和向量状态分页查询知识项。", "M", "I", "T"),
            ("FR-KB-002", "知识项必须包含非空标题、正文和真实性标签；来源、摘要、关键词、辟谣说明等字段可选。", "M", "I", "T"),
            ("FR-KB-003", "新增知识项应先保存MySQL业务数据，再执行向量同步并回填vector_id和同步状态。", "M", "I", "T"),
            ("FR-KB-004", "更新标题、正文、摘要、关键词等向量输入字段时应重新向量化；同步失败时业务变更应按事务策略回滚或标记失败。", "M", "I", "T"),
            ("FR-KB-005", "向量状态仅允许pending、synced、failed、delete_failed，并保存最近错误摘要。", "M", "I", "T"),
            ("FR-KB-006", "管理员应能对单条知识项重新向量化并看到同步结果。", "M", "I", "D/T"),
            ("FR-KB-007", "管理员应能全量重建索引，结果返回total、success、failed和failed_ids。", "M", "I", "T"),
            ("FR-KB-008", "切换Embedding供应商、模型或向量维度后必须重建全部索引。", "M", "I", "I/D"),
            ("FR-KB-009", "删除知识项应协调Chroma向量与MySQL记录；任一步失败时应执行补偿并保留可诊断状态。", "M", "I", "T"),
            ("FR-KB-010", "Chroma操作失败时应清除缓存句柄并最多重试一次，仍失败则返回明确错误。", "S", "I", "T"),
            ("FR-KB-011", "系统应支持由定时采集创建知识项，但自动入库必须保留来源URL、任务名和同步状态。", "S", "I", "T"),
        ],
    ),
    (
        "4.10",
        "Prompt模板管理",
        [
            ("FR-PRM-001", "管理员应能分页查询、查看、新增、编辑和删除Prompt模板。", "M", "I", "D/T"),
            ("FR-PRM-002", "模板名称、类型和内容必须非空；状态仅允许enabled和disabled。", "M", "I", "T"),
            ("FR-PRM-003", "news_credibility模板必须包含{title}、{content}和{evidence_list}，兼容{evidence_json}。", "M", "I", "T"),
            ("FR-PRM-004", "创建、编辑、启用和设为默认时均应执行模板契约校验。", "M", "I", "T"),
            ("FR-PRM-005", "同一type下最多一个默认模板；设为默认时应在事务中取消同类型其他默认标记。", "M", "I", "T"),
            ("FR-PRM-006", "disabled或契约无效的模板不得设为默认。", "M", "I", "T"),
            ("FR-PRM-007", "默认模板不得直接删除；删除冲突应返回409。", "M", "I", "T"),
            ("FR-PRM-008", "没有可用数据库模板或模板无效时，检测应使用受控的代码内置模板并记录告警。", "M", "I", "T"),
            ("FR-PRM-009", "模板变更应记录创建者/操作者、动作和时间，生产变更前应通过代表性样例回归。", "S", "I/A", "I/T"),
        ],
    ),
    (
        "4.11",
        "高风险审核管理",
        [
            ("FR-HR-001", "高风险记录创建后默认review_status=pending且is_public=false。", "M", "I", "T"),
            ("FR-HR-002", "管理员应能按风险、类别、审核、公开、关键词和时间范围筛选高风险记录。", "M", "I", "T"),
            ("FR-HR-003", "管理员详情应包含完整正文、分项评分、证据、风险点和审核信息。", "M", "I", "D/T"),
            ("FR-HR-004", "审核状态只允许pending、approved和rejected；改为pending或rejected时应强制取消公开。", "M", "I", "T"),
            ("FR-HR-005", "只有is_high_risk=true且review_status=approved的记录才能设置公开，否则返回409。", "M", "I", "T"),
            ("FR-HR-006", "管理员应能编辑内部备注；内部备注不得通过公开接口返回。", "M", "I", "T"),
            ("FR-HR-007", "审核、公开和备注变更应保存reviewed_by、reviewed_at并记录审计事件。", "M", "I", "T"),
        ],
    ),
    (
        "4.12",
        "统计分析与系统日志",
        [
            ("FR-STAT-001", "管理员总览应返回用户、检测、知识、高风险、今日检测和报告总数。", "M", "I", "T"),
            ("FR-STAT-002", "系统应提供检测趋势、风险分布、类别分布、高频关键词和用户活跃度。", "M", "I", "D/T"),
            ("FR-STAT-003", "趋势和活跃度默认按days统计，范围1～366；也应支持明确起止日期。", "S", "I", "T"),
            ("FR-STAT-004", "开始日期晚于结束日期或范围超过366天时应返回422。", "M", "I", "T"),
            ("FR-STAT-005", "没有数据的日期应补零，保证时间序列连续。", "S", "I", "T"),
            ("FR-STAT-006", "知识概览应包含类别、真实性标签和向量状态分布。", "M", "I", "T"),
            ("FR-LOG-001", "系统应记录注册、登录成功/失败、检测、用户管理、知识、Prompt、高风险和报告等关键事件。", "M", "I", "I/T"),
            ("FR-LOG-002", "管理员应能按关键词、模块、动作、用户和时间范围分页查询系统日志。", "M", "I", "T"),
            ("FR-LOG-003", "审计日志写入失败不得阻断主要业务，但应输出应用错误日志。", "M", "I", "T"),
            ("FR-LOG-004", "公开接口和普通用户接口不得返回管理员内部日志。", "M", "I", "T"),
        ],
    ),
    (
        "4.13",
        "定时新闻采集",
        [
            ("FR-CRAWL-001", "系统应支持通过配置开关启停定时采集，不启用时不得创建调度任务。", "S", "I", "I/T"),
            ("FR-CRAWL-002", "采集周期应由Cron表达式配置，基线为每6小时。", "S", "I", "I/T"),
            ("FR-CRAWL-003", "采集应限制并发、单页超时和最大下载字节数，并默认禁止私有网络地址。", "M", "I", "T"),
            ("FR-CRAWL-004", "采集结果应去重，并记录发现数量、新增数量、跳过数量、错误和执行时间。", "S", "I", "T"),
            ("FR-CRAWL-005", "CRAWL_AUTO_SYNC_VECTOR=true时新增知识应同步向量；关闭时仅写MySQL并标记待同步。", "S", "I", "T"),
            ("FR-CRAWL-006", "单个采集任务失败不得终止其他任务或核心检测服务。", "M", "I", "T"),
        ],
    ),
]


NFR_GROUPS: list[tuple[str, str, list[tuple[str, str, str, str, str]]]] = [
    ("7.1", "性能与容量", [
        ("NFR-PERF-001", "在目标试运行环境中，不依赖外部智能服务的普通查询P95响应时间应不超过2秒。", "M", "A", "T"),
        ("NFR-PERF-002", "本地向量检索P95应不超过3秒；完整智能检测在外部服务正常时P95应不超过120秒。", "M", "A", "T"),
        ("NFR-PERF-003", "系统应支持至少20个并发交互用户，错误率不高于1%，且不得发生数据串扰。", "M", "A", "T"),
        ("NFR-PERF-004", "列表接口应分页，每页最多100条；统计日期范围最多366天。", "M", "I", "I/T"),
        ("NFR-PERF-005", "前端首屏压缩资源建议不超过1.5MB；ECharts等大组件应按页面懒加载。", "S", "A", "A/T"),
        ("NFR-PERF-006", "外部API调用应有可配置超时和调用数量上限，防止单请求无限等待或成本失控。", "M", "I", "I/T"),
    ]),
    ("7.2", "可靠性、可用性与恢复", [
        ("NFR-REL-001", "业务数据库写入失败时应回滚当前事务，不得返回部分成功。", "M", "I", "T"),
        ("NFR-REL-002", "搜索失败应保留本地证据，LLM失败应显式降级，Embedding失败应明确报错，不得伪造模型结果。", "M", "I", "T"),
        ("NFR-REL-003", "MySQL、Chroma和报告文件应纳入一致的备份集；建议RPO≤24小时、RTO≤4小时。", "M", "A", "D/T"),
        ("NFR-REL-004", "至少每季度执行一次恢复演练并保存结果。", "S", "A", "I/D"),
        ("NFR-REL-005", "系统健康接口应在后端运行时返回服务名和ok状态；生产监控不得只依赖该浅层检查。", "S", "I", "T/A"),
        ("NFR-REL-006", "报告、向量和外部API错误应包含可定位的日志上下文，同时避免泄露密钥和敏感数据。", "M", "I/A", "I/T"),
    ]),
    ("7.3", "安全性", [
        ("NFR-SEC-001", "生产环境必须使用HTTPS；JWT密钥至少32个随机字符且不得使用模板值。", "M", "A", "I/T"),
        ("NFR-SEC-002", "生产环境CORS必须使用明确来源白名单，不得使用通配符与凭据组合。", "M", "A", "I/T"),
        ("NFR-SEC-003", "认证、检测、URL抓取和高成本接口应实施限流；多实例部署时应使用共享限流或网关。", "M", "I/A", "I/T"),
        ("NFR-SEC-004", "所有用户输入应进行类型、长度、枚举和空值校验；SQL操作必须通过ORM参数绑定。", "M", "I", "I/T"),
        ("NFR-SEC-005", "URL抓取必须防止SSRF、DNS/重定向绕过、超大响应和非HTTP协议。", "M", "I", "T"),
        ("NFR-SEC-006", "报告下载必须验证所有权/管理员权限并验证规范化路径位于REPORT_DIR内。", "M", "I", "T"),
        ("NFR-SEC-007", "任何接口不得返回密码哈希、SECRET_KEY或第三方API Key。", "M", "I", "I/T"),
        ("NFR-SEC-008", "生产数据库账号应遵循最小权限；备份和日志访问应受控并留痕。", "M", "A", "I"),
        ("NFR-SEC-009", "依赖和容器/主机镜像应在发布前完成已知漏洞与许可证检查。", "S", "A", "I"),
    ]),
    ("7.4", "个人信息、版权与内容治理", [
        ("NFR-PRI-001", "系统应在提交前告知新闻文本可能被发送给模型/搜索供应商，并说明处理目的和范围。", "M", "A", "I/D"),
        ("NFR-PRI-002", "系统应最小化收集账号和新闻内容，不得要求与鉴别无关的敏感个人信息。", "M", "A", "I"),
        ("NFR-PRI-003", "应定义用户、检测、日志、报告和抓取内容的保留期限及删除/更正流程。", "M", "T", "I"),
        ("NFR-PRI-004", "网络证据应优先保存必要摘要、元数据和来源链接，避免无授权长期复制完整受保护内容。", "M", "A", "I"),
        ("NFR-PRI-005", "公开高风险内容必须经人工复核，并提供撤回、更正或申诉处理机制。", "M", "A", "D/I"),
        ("NFR-PRI-006", "所有自动结论应明确为辅助评估；不得用于对个人或机构作出自动化不利决定的唯一依据。", "M", "A", "I/D"),
    ]),
    ("7.5", "易用性与可访问性", [
        ("NFR-USA-001", "页面应使用一致的导航、术语、风险颜色和错误提示；关键操作应有加载、成功和失败状态。", "M", "I", "D"),
        ("NFR-USA-002", "用户应能在不阅读技术文档的情况下完成输入、检测、查看结果和下载报告。", "M", "A", "D"),
        ("NFR-USA-003", "高风险公开、删除、重建索引、禁用账号等不可逆或高影响操作应二次确认。", "M", "I/A", "D/T"),
        ("NFR-USA-004", "评分页应同时展示总分、分项分数、有效证据、风险点、建议和不确定性状态。", "M", "I", "D"),
        ("NFR-USA-005", "界面应支持现代Chrome和Edge；建议最低分辨率1280×720并在较窄窗口保持可用。", "S", "A", "D/T"),
        ("NFR-USA-006", "文字与控件应满足基本键盘操作、焦点可见、标签关联和对比度要求。", "S", "A", "I/D"),
    ]),
    ("7.6", "可维护性、可测试性与可移植性", [
        ("NFR-MNT-001", "后端应保持API、Service、CRUD、Model/Schema分层，不在路由中堆叠业务规则。", "M", "I", "I"),
        ("NFR-MNT-002", "数据库结构变更必须通过Alembic版本管理；启动时应检查数据库位于迁移head。", "M", "I", "I/T"),
        ("NFR-MNT-003", "外部模型、Embedding、搜索和存储路径应通过配置注入，不得硬编码真实密钥。", "M", "I", "I/T"),
        ("NFR-MNT-004", "关键业务、权限、异常降级和数据一致性需求应有自动化测试；发布前后端测试必须全部通过。", "M", "I/A", "T"),
        ("NFR-MNT-005", "生产构建应成功，不得存在阻断告警；大分块告警应登记或优化。", "M", "I/A", "T/I"),
        ("NFR-MNT-006", "系统应能在Windows开发环境运行，并可迁移到常见Linux服务器；路径处理不得依赖固定工作目录。", "S", "I/A", "D/T"),
        ("NFR-MNT-007", "更换Embedding模型或维度必须触发索引重建和回归；更换LLM应通过同一结构化契约。", "M", "I/A", "I/T"),
        ("NFR-MNT-008", "需求、接口、数据库、测试和验收文档应使用稳定编号并保持追踪关系。", "M", "A", "I"),
    ]),
    ("7.7", "可观测性与审计", [
        ("NFR-OBS-001", "应用日志应至少包含时间、级别、模块、事件和关联ID/业务ID，不得记录密钥和完整密码。", "M", "A", "I/T"),
        ("NFR-OBS-002", "应监控请求错误率和延迟、外部API失败/耗时、向量同步失败、调度失败、数据库连接和磁盘容量。", "M", "A", "D/I"),
        ("NFR-OBS-003", "重大告警应包含阈值、责任人和处置说明；恢复后应关闭并记录复盘。", "S", "A", "I"),
        ("NFR-OBS-004", "安全与管理审计日志的保留期应由部署方确定，且普通用户不可访问。", "M", "T", "I/T"),
    ]),
]


def configure_base() -> None:
    base.DOC_NO = DOC_NO
    base.DOC_TITLE = DOC_TITLE
    base.PROJECT_TITLE = PROJECT_TITLE
    base.VERSION = VERSION
    base.DATE_TEXT = DATE_TEXT


def add_req_table(doc: Document, requirements: list[tuple[str, str, str, str, str]]) -> None:
    base.add_table(
        doc,
        ["需求编号", "需求描述", "优先级", "状态", "验证"],
        [list(item) for item in requirements],
        widths=[2.6, 10.5, 1.1, 1.0, 1.3],
        font_size=7.8,
    )


def build_cover(doc: Document) -> None:
    accent = doc.add_table(rows=1, cols=2)
    accent.alignment = WD_TABLE_ALIGNMENT.CENTER
    accent.autofit = False
    accent.cell(0, 0).width = Cm(10.5)
    accent.cell(0, 1).width = Cm(5.0)
    base.set_cell_shading(accent.cell(0, 0), base.NAVY)
    base.set_cell_shading(accent.cell(0, 1), base.TEAL)
    for cell in accent.rows[0].cells:
        cell.height = Cm(0.22)
        cell.text = ""
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("SOFTWARE REQUIREMENTS SPECIFICATION")
    base.set_run_font(r, east_asia="Arial", latin="Arial", size=10, bold=True, color=base.TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(f"《{PROJECT_TITLE}》")
    base.set_run_font(r, east_asia="黑体", size=24, bold=True, color=base.NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("需 求 分 析 书")
    base.set_run_font(r, east_asia="黑体", size=30, bold=True, color=base.BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("（软件需求规格说明）")
    base.set_run_font(r, east_asia="黑体", size=15, bold=True, color=base.GRAY)

    for _ in range(2):
        doc.add_paragraph()

    info = [
        ("文档编号", DOC_NO),
        ("版本号", VERSION),
        ("需求基线", "SRS-BL-2026-06-20"),
        ("文档状态", "正式版"),
        ("密级", "内部资料"),
        ("编制日期", DATE_TEXT),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (key, value) in zip(table.rows, info):
        row.cells[0].width = Cm(4.2)
        row.cells[1].width = Cm(8.1)
        base.set_cell_shading(row.cells[0], base.PALE_BLUE)
        for cell in row.cells:
            base.set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
            cell.paragraphs[0].paragraph_format.first_line_indent = Cm(0)
        r1 = row.cells[0].paragraphs[0].add_run(key)
        base.set_run_font(r1, east_asia="黑体", size=10.5, bold=True, color=base.NAVY)
        r2 = row.cells[1].paragraphs[0].add_run(value)
        base.set_run_font(r2, size=10.5)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("项目组")
    base.set_run_font(r, east_asia="黑体", size=12, bold=True, color=base.NAVY)


def build_control(doc: Document) -> None:
    base.add_center_title(doc, "文档控制信息", 20)
    base.add_table(doc, ["项目", "内容"], [
        ["文档名称", f"《{PROJECT_TITLE}》{DOC_TITLE}"],
        ["主要编制依据", "GB/T 9385—2008《计算机软件需求规格说明规范》、GB/T 8567—2006《计算机软件文档编制规范》"],
        ["需求基线", "SRS-BL-2026-06-20"],
        ["适用范围", "课程设计/毕业设计原型、验收准备和小规模受控部署"],
        ["事实来源", "项目设计资料、当前源码、codebase知识图谱、数据库模型、接口与测试基线"],
        ["读者", "项目负责人、需求/设计/开发/测试人员、指导教师、验收人员和维护人员"],
    ], widths=[3.5, 13.0], font_size=9.2)

    doc.add_heading("修订记录", level=2)
    base.add_table(doc, ["版本", "日期", "说明", "修订者"], [
        [VERSION, DATE_TEXT, "首次正式编制；建立功能、接口、数据、非功能及追踪需求基线", "项目组"],
    ], widths=[2.2, 3.3, 8.4, 2.6])

    doc.add_heading("审核与批准", level=2)
    base.add_table(doc, ["角色", "姓名", "意见", "签字", "日期"], [
        ["编制", "", "", "", ""], ["审核", "", "", "", ""], ["批准", "", "", "", ""],
    ], widths=[2.4, 3.0, 5.5, 2.8, 2.8])

    doc.add_heading("需求变更控制", level=2)
    base.add_body(doc, "需求基线批准后，任何新增、删除或变更均应提交变更记录，说明业务原因、受影响的需求编号、接口、数据、测试和文档；经审核后更新版本号和追踪矩阵。不得仅修改代码而不更新需求状态。")


def build_abstract_toc(doc: Document) -> None:
    base.page_break(doc)
    base.add_center_title(doc, "摘要", 20)
    total_functional = sum(len(group[2]) for group in REQ_GROUPS)
    total_nfr = sum(len(group[2]) for group in NFR_GROUPS)
    base.add_body(doc, f"本需求分析书依据GB/T 9385—2008和GB/T 8567—2006编制，对《{PROJECT_TITLE}》的业务目标、用户角色、功能能力、外部接口、数据、质量属性、资源环境、约束、合格性规定和需求追踪关系作出可验证的规格说明。文档共定义{total_functional}条功能需求和{total_nfr}条非功能需求，每条关键需求均具有稳定编号、优先级、当前状态和验证方法。")
    base.add_body(doc, "需求事实基线来自当前可运行原型。codebase知识图谱显示项目包含3,073个节点、8,131条关系、8个核心数据实体和35项环境配置，主检测链路覆盖文本清洗、本地RAG、条件联网、证据仲裁、规则与模型评分、风险分级、持久化和报告。文档同时把当前已实现能力与验收目标、待定事项分开标识，避免把尚未验证的性能、安全或合规指标写成既成事实。")
    base.add_note(doc, "需求口径", "“系统应”表示可验证的强制或建议需求；I表示当前实现基线，A表示验收前应满足的目标，T表示必须由部署方决策后基线化的事项。", kind="info")

    base.page_break(doc)
    base.add_center_title(doc, "目录", 20)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    base.add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "目录将在Word中自动更新")


def build_intro(doc: Document) -> None:
    doc.add_heading("1 引言", level=1)
    doc.add_heading("1.1 编写目的", level=2)
    base.add_body(doc, "本文件建立系统需求基线，作为概要设计、详细设计、编码、测试、验收、用户培训和维护的共同依据。它描述系统必须提供什么能力以及如何验证，不规定不必要的内部实现细节。")

    doc.add_heading("1.2 系统范围", level=2)
    base.add_body(doc, "系统面向网络新闻辅助核验场景。用户提交新闻文本或链接预览，系统利用本地知识库、可选联网证据、规则分析和大语言模型形成可解释的可信度评估。管理员负责知识、Prompt、用户、检测、高风险审核、报告、统计和日志。系统不承担权威事实裁决、自动执法或专业领域最终决策。")
    base.add_table(doc, ["范围内", "范围外"], [[
        "Web端注册登录、新闻检测、RAG与联网证据、LLM仲裁、规则评分、历史、报告、高风险公开与审核、知识/Prompt/用户/统计/日志/采集管理",
        "训练基础大模型、建设通用搜索引擎、全网舆情监控、自动法律/医疗/金融结论、移动原生应用、多租户计费、权威辟谣发布平台",
    ]], widths=[8.25, 8.25], font_size=9.0)

    doc.add_heading("1.3 读者对象", level=2)
    base.add_table(doc, ["读者", "使用目的"], [
        ["项目负责人/需求人员", "确认范围、优先级、待定事项与变更"],
        ["设计与开发人员", "建立模块、接口、数据和错误处理设计"],
        ["测试与验收人员", "由需求编号导出测试、检查和演示用例"],
        ["管理员与维护人员", "理解运行、权限、配置、备份和审计要求"],
        ["指导教师/评审人员", "检查需求完整性、实现一致性与国标符合性"],
    ], widths=[4.3, 12.2])

    doc.add_heading("1.4 术语和缩略语", level=2)
    base.add_table(doc, ["术语", "定义"], [
        ["RAG", "检索增强生成：先检索证据，再基于证据执行模型分析。"],
        ["LLM", "大语言模型；本系统通过结构化契约要求其输出分析和证据仲裁。"],
        ["候选证据", "本地知识库或联网搜索召回、尚未通过模型仲裁的材料。"],
        ["有效证据", "通过仲裁进入ranked_evidence并可用于证据质量和结果解释的材料。"],
        ["排除证据", "被仲裁拒绝、不参与评分但可保留拒绝原因的候选。"],
        ["证据质量", "coverage与consistency的组合分，服务端按0.6/0.4重算。"],
        ["高风险记录", "final_score<40或风险等级为高风险谣言的检测记录。"],
        ["I/A/T", "Implemented当前已实现；Acceptance验收目标；TBD待定。"],
        ["T/D/I/A", "Test测试；Demonstration演示；Inspection检查；Analysis分析。"],
    ], widths=[3.5, 13.0], font_size=9.1)

    doc.add_heading("1.5 引用文件", level=2)
    base.add_table(doc, ["序号", "文件", "用途", "来源"], [
        ["[1]", "GB/T 9385—2008《计算机软件需求规格说明规范》", "主要需求规格依据，现行", "https://openstd.samr.gov.cn/bzgk/gb/newGbInfo?hcno=2790825C43AD0B69E3C38C140BFFCFE6"],
        ["[2]", "GB/T 8567—2006《计算机软件文档编制规范》", "文档组织与内容依据，现行", "https://openstd.samr.gov.cn/bzgk/gb/newGbInfo?hcno=84C42B6277D2714B7176B10C6E6B1A44"],
        ["[3]", "GB/T 11457—2006《信息技术 软件工程术语》", "术语依据，现行", "https://openstd.samr.gov.cn/bzgk/gb/newGbInfo?hcno=07E3E9867D23EA5A74EB525A44622E86"],
        ["[4]", "GB/T 25000.10—2016《系统与软件质量模型》", "质量属性依据，现行", "https://openstd.samr.gov.cn/bzgk/gb/newGbInfo?hcno=DB14B9415F3D51A9EF321EA62EDDE9A6"],
        ["[5]", "NCE-FS-001 可行性研究报告", "项目决策、边界和条件", "同目录01号文档"],
        ["[6]", "项目源码、docs设计资料与codebase知识图谱", "实际需求事实基线", "NewsCredibilityEvaluator/"],
    ], widths=[1.2, 5.4, 4.2, 5.7], font_size=8.0)

    doc.add_heading("1.6 需求表述约定", level=2)
    base.add_bullets(doc, [
        "M（Must）：缺失将导致核心目标、法律安全或验收失败；S（Should）：重要但可在受控情况下延期；C（Could）：增强能力。",
        "状态I说明当前源码已有实现，并不自动代表生产验收；状态A需要在验收环境验证；状态T必须先决策再批准。",
        "需求编号批准后不得复用。废弃需求保留编号并标记状态，避免追踪断裂。",
        "除特别说明外，‘用户’指普通注册用户，‘管理员’指role=admin且status=active的账号。",
    ])


def build_overview(doc: Document) -> None:
    doc.add_heading("2 总体描述", level=1)
    doc.add_heading("2.1 产品前景和系统边界", level=2)
    base.add_body(doc, "系统是一个前后端分离的新闻可信度辅助评估应用。浏览器用户通过Vue界面访问FastAPI接口；后端协调MySQL、Chroma、报告目录和外部智能服务。系统对外输出的是可解释的辅助评估，不输出法律意义或权威媒体意义上的真伪裁决。")
    base.add_figure(doc, MODULE_FIGURE, "图2-1 系统功能领域与边界", width_inches=6.15)

    doc.add_heading("2.2 业务目标", level=2)
    base.add_table(doc, ["目标编号", "目标", "衡量方式"], [
        ["BG-01", "降低新闻初筛与证据整理时间", "基准场景从约15分钟/条降低至约5分钟人工复核"],
        ["BG-02", "提高结论可解释与可追溯性", "结果同时包含分项分、有效证据、风险点、建议和契约状态"],
        ["BG-03", "统一高风险信息公开流程", "未经approved＋public不得进入公开列表"],
        ["BG-04", "沉淀可复用知识和报告", "知识可维护/向量化，检测可查询并生成报告"],
        ["BG-05", "形成可验收的软件工程资产", "需求—设计—实现—测试—验收文档可追踪"],
    ], widths=[2.2, 7.3, 7.0], font_size=8.9)

    doc.add_heading("2.3 用户类别与特征", level=2)
    base.add_table(doc, ["角色", "技能假设", "主要任务", "权限边界"], [
        ["游客", "能使用浏览器", "浏览公开内容、提交临时检测、查看本次结果", "无历史、报告和后台权限"],
        ["普通用户", "一般互联网用户", "登录、检测、历史、详情、重新评估、报告、个人中心", "仅访问自己的非公开数据"],
        ["管理员", "接受后台培训", "用户、检测、知识、Prompt、高风险、统计、报告、日志", "不得绕过最后管理员和公开审核规则"],
        ["维护人员", "具备部署/数据库能力", "配置、迁移、备份、恢复、监控和升级", "通过系统外受控运维权限操作"],
        ["领域复核人员", "具备新闻/事实核查知识", "抽检知识标签、审核高风险公开、处理更正", "可由管理员角色承载，但职责应独立"],
    ], widths=[2.4, 3.5, 6.2, 4.4], font_size=8.5)

    doc.add_heading("2.4 运行方式与降级模式", level=2)
    base.add_table(doc, ["模式", "触发条件", "系统行为", "用户提示"], [
        ["正常模式", "MySQL、Chroma、Embedding、LLM可用", "完整检索、仲裁、评分和持久化", "展示完整结果"],
        ["搜索降级", "联网搜索关闭/失败", "保留本地候选并继续", "说明是否触发及来源数量"],
        ["无有效证据", "无候选或仲裁后无有效证据", "0.6×LLM＋0.4×规则", "证据质量不可用/无证据"],
        ["LLM降级", "LLM调用失败", "未仲裁证据不计分，最终分使用规则分", "明确模型服务不可用，建议人工核验"],
        ["仲裁失败", "两次契约校验均失败", "候选不计分，保存retry_exhausted和错误摘要", "提供重新评估入口"],
        ["检索不可用", "Embedding或Chroma检索失败", "终止本次检测并返回503", "提示稍后重试，不伪造结果"],
        ["数据库不可用", "事务/连接失败", "回滚并返回统一错误", "提示服务暂不可用"],
    ], widths=[2.6, 4.0, 6.0, 3.9], font_size=8.2)

    doc.add_heading("2.5 假定、依赖和约束", level=2)
    base.add_bullets(doc, [
        "部署方能够提供MySQL、可写Chroma目录和报告目录，并安全配置模型、向量和搜索服务凭据。",
        "正式RAG使用的Embedding模型输出维度与EMBEDDING_DIMENSION一致；切换后重建索引。",
        "知识库材料来源合法、标签可解释且覆盖目标新闻类别；领域人员负责抽检。",
        "系统优先服务课程/小规模试点，不承诺大规模实时舆情平台的容量和SLA。",
        "浏览器端依赖JavaScript与现代Web API；后端接口统一以/api为前缀。",
        "生产环境需补充HTTPS、共享限流、集中监控、备份恢复和合规告知。",
    ])


def add_use_case(doc: Document, title: str, rows: list[list[str]]) -> None:
    doc.add_heading(title, level=3)
    base.add_table(doc, ["项目", "内容"], rows, widths=[3.2, 13.3], font_size=8.8)


def build_use_cases(doc: Document) -> None:
    doc.add_heading("3 业务需求与用例分析", level=1)
    doc.add_heading("3.1 总体业务流程", level=2)
    base.add_figure(doc, FLOW_FIGURE, "图3-1 新闻可信度评估总体业务流程", width_inches=4.8)

    doc.add_heading("3.2 核心用例", level=2)
    add_use_case(doc, "3.2.1 UC-01 注册与登录", [
        ["参与者", "游客、普通用户、管理员"],
        ["前置条件", "前端与后端可访问；数据库迁移处于head"],
        ["触发", "用户提交注册或登录表单"],
        ["主流程", "校验输入→检查限流→注册/认证→生成JWT→记录审计→进入对应页面"],
        ["异常", "重复账号400；错误密码401；禁用账号403；限流429；校验失败422"],
        ["后置条件", "注册产生user记录；登录产生有效会话；密码不以明文存储"],
        ["关联需求", "FR-AUTH-001～011"],
    ])
    add_use_case(doc, "3.2.2 UC-02 提交新闻检测", [
        ["参与者", "游客、普通用户、管理员"],
        ["前置条件", "标题≥4字符、正文≥20字符；服务未触发限流"],
        ["触发", "用户提交手工文本，或先提取URL并确认"],
        ["主流程", "输入校验→清洗→本地RAG→条件联网→候选池→LLM仲裁→规则/模型评分→风险分级→保存→返回结果"],
        ["异常", "URL/输入422；检索失败503；检测业务错误400；LLM失败进入降级"],
        ["后置条件", "产生检测记录；登录用户关联user_id；有效证据持久化；高风险默认待审不公开"],
        ["关联需求", "FR-DET、FR-RAG、FR-SCORE"],
    ])
    add_use_case(doc, "3.2.3 UC-03 查看、重新评估与生成报告", [
        ["参与者", "普通用户、管理员"],
        ["前置条件", "用户已登录且拥有记录，或为管理员"],
        ["主流程", "筛选历史→查看详情→必要时重新评估生成新记录→生成/下载PDF"],
        ["异常", "越权/不存在404或403；旧输入不兼容422；文件不存在404；生成失败500"],
        ["后置条件", "原检测不可变；新评估有新ID；同一检测至多一条报告记录"],
        ["关联需求", "FR-HIS、FR-REP"],
    ])
    add_use_case(doc, "3.2.4 UC-04 管理知识库", [
        ["参与者", "管理员"],
        ["前置条件", "管理员已认证；MySQL和Chroma可访问"],
        ["主流程", "查询→新增/编辑→向量化→检查同步状态；必要时单条重试或全量重建"],
        ["异常", "校验422；不存在404；向量失败记录failed；删除补偿失败记录delete_failed"],
        ["后置条件", "MySQL业务记录与Chroma向量可通过knowledge_id/vector_id关联"],
        ["关联需求", "FR-KB-001～011"],
    ])
    add_use_case(doc, "3.2.5 UC-05 管理Prompt模板", [
        ["参与者", "管理员"],
        ["主流程", "查询→新增/编辑→契约校验→启用→设为同类型默认→样例回归"],
        ["异常", "缺占位符422；禁用/无效模板不得默认；默认模板删除409"],
        ["后置条件", "每个模板类型最多一个有效默认模板；变更可审计"],
        ["关联需求", "FR-PRM-001～009"],
    ])
    add_use_case(doc, "3.2.6 UC-06 高风险审核与公开", [
        ["参与者", "管理员、公开访问者"],
        ["前置条件", "记录已由检测标记为高风险"],
        ["主流程", "管理员查看详情→审核通过/驳回→填写备注→通过后按需公开→公众浏览脱敏摘要"],
        ["异常", "未通过审核却请求公开时409；驳回/待审自动取消公开"],
        ["后置条件", "只有高风险＋已通过＋公开三条件同时满足时对外展示"],
        ["关联需求", "FR-HR、FR-PUB"],
    ])
    add_use_case(doc, "3.2.7 UC-07 用户、检测、统计与日志管理", [
        ["参与者", "管理员"],
        ["主流程", "筛选用户/检测→启禁用或查看/删除→查看统计→查询审计日志"],
        ["异常", "非管理员401/403；最后管理员保护409；统计日期非法422"],
        ["后置条件", "管理动作受最小权限和审计约束"],
        ["关联需求", "FR-ADM、FR-STAT、FR-LOG"],
    ])

    doc.add_heading("3.3 业务规则", level=2)
    base.add_table(doc, ["规则编号", "规则"], [
        ["BR-01", "风险等级只能由final_score统一映射，不允许前端另行推导。"],
        ["BR-02", "高风险标记由final_score<40或风险等级=高风险谣言触发。"],
        ["BR-03", "公开高风险记录必须同时满足高风险、审核通过、明确公开。"],
        ["BR-04", "未成功仲裁的候选证据不得参与评分和最终证据持久化。"],
        ["BR-05", "LLM失败时最终分只使用确定性规则分，并明确能力降级。"],
        ["BR-06", "同类型Prompt最多一个默认模板，默认模板必须enabled且契约有效。"],
        ["BR-07", "不得禁用或降级当前管理员及最后一个可用管理员。"],
        ["BR-08", "同一检测记录至多一条报告记录；报告不能越权下载。"],
        ["BR-09", "正式语义向量模型或维度改变后必须全量重建Chroma索引。"],
        ["BR-10", "游客检测不形成可访问个人历史；登录用户的数据按user_id隔离。"],
    ], widths=[2.5, 14.0], font_size=8.8)


def build_functional(doc: Document) -> None:
    doc.add_heading("4 系统功能需求", level=1)
    base.add_note(doc, "表格说明", "优先级M/S/C；状态I/A/T；验证方法T/D/I/A。组合值如I/A表示已有基础实现但仍有验收增强项。", kind="info")
    for section_no, title, requirements in REQ_GROUPS:
        doc.add_heading(f"{section_no} {title}", level=2)
        add_req_table(doc, requirements)
        if section_no == "4.4":
            base.add_figure(doc, SCORE_FIGURE, "图4-1 评分与风险等级需求规则", width_inches=4.8)


def build_interfaces(doc: Document) -> None:
    doc.add_heading("5 外部接口需求", level=1)
    doc.add_heading("5.1 用户界面", level=2)
    base.add_figure(doc, USER_FIGURE, "图5-1 用户端页面与访问要求", width_inches=4.9)
    base.add_figure(doc, ADMIN_FIGURE, "图5-2 管理员端页面与权限要求", width_inches=4.85)
    base.add_table(doc, ["界面编号", "页面/路径", "角色", "主要内容"], [
        ["UI-01", "首页 /", "公开", "系统说明、检测入口、风险等级说明"],
        ["UI-02", "登录 /login", "游客", "用户名、密码、错误提示、注册链接"],
        ["UI-03", "注册 /register", "游客", "用户名、邮箱、密码、确认与校验"],
        ["UI-04", "新闻检测 /detect", "公开", "链接预览、文本输入、联网开关、进度和提交"],
        ["UI-05", "检测结果 /result/:id", "本次游客/登录用户", "总分、分项、风险、理由、证据、排除证据、建议、报告"],
        ["UI-06", "历史 /history", "登录用户", "筛选、分页、详情、重新评估和报告"],
        ["UI-07", "高风险 /high-risk", "公开", "审核后公开列表、排名、关键词和类别"],
        ["UI-08", "个人中心 /profile", "登录用户", "账号信息和快捷入口"],
        ["UI-09", "后台 /admin/*", "管理员", "首页、用户、检测、知识、Prompt、高风险、统计、报告和日志"],
    ], widths=[2.0, 4.5, 3.0, 7.0], font_size=8.4)

    doc.add_heading("5.2 REST API通用约定", level=2)
    base.add_bullets(doc, [
        "所有业务接口使用/api前缀，数据交换编码为UTF-8 JSON；PDF下载除外。",
        "成功/失败JSON统一包含code、message和data；HTTP状态码与业务code应一致。",
        "认证头格式为Authorization: Bearer <token>；公开接口不得要求令牌。",
        "参数错误422、未认证401、无权限403、不存在404、状态冲突409、限流429、服务不可用503。",
        "分页参数page从1开始，page_size通常默认20且最大100；日期使用ISO 8601。",
        "接口Schema应拒绝未声明的额外字段，除明确的兼容场景外不得静默忽略。",
    ])

    doc.add_heading("5.3 API接口清单", level=2)
    public_apis = [
        ["POST", "/api/auth/register", "公开", "注册普通用户"], ["POST", "/api/auth/login", "公开", "登录并获取JWT"], ["GET", "/api/auth/me", "登录", "当前用户"],
        ["POST", "/api/detect/extract-preview", "公开/限流", "安全提取新闻链接预览"], ["POST", "/api/detect/news", "公开/限流", "执行新闻检测"],
        ["GET", "/api/detect/history", "登录", "个人历史"], ["GET", "/api/detect/{id}", "所有者/管理员", "检测详情"], ["POST", "/api/detect/{id}/re-evaluate", "所有者/管理员", "重新评估并生成新记录"],
        ["POST", "/api/rag/search", "登录", "相似知识检索"], ["POST", "/api/report/generate/{detection_id}", "所有者/管理员", "生成报告"], ["GET", "/api/report/download/{report_id}", "所有者/管理员", "下载PDF"],
        ["GET", "/api/high-risk/public", "公开", "公开高风险列表"], ["GET", "/api/high-risk/ranking", "公开", "公开排名"], ["GET", "/api/high-risk/keywords", "公开", "公开关键词"], ["GET", "/api/high-risk/category-distribution", "公开", "公开类别分布"], ["GET", "/api/health", "公开", "服务健康"],
    ]
    base.add_table(doc, ["方法", "路径", "权限", "用途"], public_apis, widths=[1.6, 7.2, 3.1, 4.6], font_size=8.0)

    admin_apis = [
        ["检测", "GET /api/admin/detections；GET/DELETE /api/admin/detections/{id}", "列表、详情、删除"],
        ["用户", "GET /api/admin/users；GET /{user_id}；GET /{user_id}/detections；POST /enable、/disable、/role", "查询、启禁用、角色"],
        ["知识", "GET/POST /api/admin/knowledge；GET/PUT/DELETE /{id}；POST /{id}/vectorize；POST /rebuild-index", "CRUD与向量维护"],
        ["Prompt", "GET/POST /api/admin/prompts；GET/PUT/DELETE /{id}；POST /{id}/enable、/disable、/set-default", "模板生命周期"],
        ["高风险", "GET /api/admin/high-risk；GET /{record_id}；PUT /review、/public、/remark", "审核与公开"],
        ["统计", "GET /api/admin/statistics/overview、trend、risk-distribution、category-distribution、keywords、user-activity、knowledge-overview", "管理统计"],
        ["报告", "GET /api/admin/reports；GET /{report_id}；GET /{report_id}/download", "报告查询与下载"],
        ["日志", "GET /api/admin/logs", "审计查询"],
        ["诊断", "GET /api/admin/ping", "管理员连通性检查"],
    ]
    base.add_table(doc, ["模块", "路径概要", "用途"], admin_apis, widths=[2.2, 10.2, 4.1], font_size=7.9)

    doc.add_heading("5.4 外部智能与搜索接口", level=2)
    base.add_table(doc, ["接口", "输入", "输出", "错误/约束"], [
        ["DeepSeek Chat", "新闻标题、正文、限量候选证据、Prompt", "LLM分、理由、风险点、建议、证据仲裁与质量", "30秒基线超时；JSON契约；失败降级"],
        ["DashScope Embedding", "清洗后的文本批次", "与配置维度一致的向量", "正式默认text-embedding-v4/1024维；切换后重建"],
        ["兼容Embedding", "文本", "向量", "DeepSeek兼容或hash演示；hash非语义"],
        ["Bocha搜索", "标题、关键词、条数、freshness", "标题、摘要、来源、URL等网络候选", "可选；8秒基线超时；失败保留本地证据"],
        ["网页提取", "HTTP(S) URL", "标题、正文、来源、发布时间与精度", "SSRF默认拒绝；限大小、重定向和超时"],
    ], widths=[3.0, 4.4, 5.0, 4.1], font_size=8.2)

    doc.add_heading("5.5 软件、通信与文件接口", level=2)
    base.add_table(doc, ["接口对象", "协议/格式", "需求"], [
        ["浏览器—前端", "HTTPS/HTML/CSS/JavaScript", "现代Chrome/Edge；响应式布局；静态资源缓存"],
        ["前端—后端", "HTTPS REST/JSON/JWT", "同源代理或白名单CORS；15～120秒分接口超时"],
        ["后端—MySQL", "SQLAlchemy/PyMySQL", "utf8mb4；事务；最小权限；连接失败可诊断"],
        ["后端—Chroma", "本地客户端/持久化目录", "collection名称稳定；cosine相似；元数据含knowledge_id"],
        ["后端—报告目录", "文件系统HTML/PDF", "目录位于源码外；规范化路径；可写、可备份"],
        ["运维—配置", ".env/环境变量", "密钥不入库；开发/测试/生产隔离"],
    ], widths=[3.5, 5.0, 8.0], font_size=8.5)


def build_data(doc: Document) -> None:
    doc.add_heading("6 数据需求", level=1)
    doc.add_heading("6.1 数据实体", level=2)
    base.add_figure(doc, ER_FIGURE, "图6-1 核心数据实体与关系（逻辑需求视图）", width_inches=6.35)
    base.add_table(doc, ["实体", "用途", "关键标识/约束", "敏感性"], [
        ["User", "账号、角色和状态", "username唯一；email可空唯一；role/status枚举", "密码哈希、邮箱"],
        ["KnowledgeItem", "新闻、核查与辟谣知识", "title/content/truth_label必填；vector_id与同步状态", "可能含版权/个人信息"],
        ["DetectionRecord", "输入、评分、风险、解释和审核", "user_id可空；分数0～100；analysis_payload", "新闻文本、用户行为"],
        ["EvidenceMatch", "最终有效证据快照", "detection_id必填；rank_order≥1；检测删除级联", "来源摘要"],
        ["PromptTemplate", "模型提示模板", "type、status、is_default；同类型唯一默认由业务保证", "内部配置"],
        ["Report", "报告元数据与文件路径", "detection_id唯一；用户所有权；受控路径", "汇总敏感内容"],
        ["SystemLog", "关键业务与管理审计", "user_id可空；action/module/time", "IP、用户活动"],
        ["CrawlTask", "定时采集执行日志", "job_name、status、数量、错误和时间", "查询词/来源"],
    ], widths=[3.0, 4.7, 5.6, 3.2], font_size=8.2)

    doc.add_heading("6.2 数据关系与完整性", level=2)
    base.add_table(doc, ["关系", "基数", "完整性规则"], [
        ["User—DetectionRecord", "1 : 0..N", "游客记录user_id为空；用户删除不开放；外键SET NULL"],
        ["User—PromptTemplate", "1 : 0..N", "created_by可空；用户保留策略不破坏模板"],
        ["User—SystemLog", "1 : 0..N", "日志允许匿名事件；外键SET NULL"],
        ["DetectionRecord—EvidenceMatch", "1 : 0..N", "只保存最终有效证据；检测删除级联"],
        ["KnowledgeItem—EvidenceMatch", "1 : 0..N", "知识删除后历史证据快照保留，knowledge_id可SET NULL"],
        ["DetectionRecord—Report", "1 : 0..1", "detection_id唯一；检测删除级联"],
        ["MySQL Knowledge—Chroma", "1 : 0..1向量", "通过knowledge_id/vector_id和vector_sync_status对账；非分布式事务"],
    ], widths=[5.0, 2.4, 9.1], font_size=8.5)

    doc.add_heading("6.3 状态与生命周期", level=2)
    base.add_table(doc, ["对象", "状态", "允许转换/规则"], [
        ["User", "active / disabled", "管理员启用或禁用；不得禁用当前或最后可用管理员"],
        ["Knowledge vector", "pending / synced / failed / delete_failed", "向量化成功到synced；失败记录错误；可重试/重建"],
        ["Prompt", "enabled / disabled＋is_default", "默认必须enabled且契约有效；同类型最多一个默认"],
        ["High-risk review", "pending / approved / rejected", "pending/rejected强制不公开；approved后才可公开"],
        ["Report", "无记录→generated→regenerated", "同一检测唯一记录；重新生成替换受控文件"],
        ["Arbitration", "unavailable / no_evidence / provider_error / ok / retry_exhausted", "最多2次尝试；只有ok结果可产生有效证据"],
    ], widths=[3.5, 5.3, 7.7], font_size=8.4)

    doc.add_heading("6.4 数据校验与格式", level=2)
    base.add_bullets(doc, [
        "所有文本输入应去除首尾空白；必填文本清洗后不得为空。",
        "分数统一为0～100；相似度持久化为非负数；rank_order从1开始。",
        "日期时间API采用ISO 8601；只有日期时必须保留date精度，不伪造时分秒。",
        "risk_points等列表持久化为JSON字符串时，读取层应兼容历史字符串并归一为字符串数组。",
        "analysis_payload用于保存候选/排除证据和契约元数据，必须保持可解析JSON对象并带版本号。",
        "公开响应采用白名单字段映射，而不是直接序列化数据库模型。",
    ])

    doc.add_heading("6.5 数据保留、备份与销毁", level=2)
    base.add_table(doc, ["数据", "建议保留", "备份", "销毁/更正"], [
        ["用户账号", "账号存续期＋合理注销缓冲", "每日数据库备份", "部署方应确定注销/匿名化流程（TBD-01）"],
        ["检测与证据", "默认1年，按部署目的调整", "每日数据库备份", "用户更正/删除与公共利益留存冲突需决策"],
        ["报告文件", "与检测记录一致", "与数据库、Chroma作为一致恢复集", "删除记录时清理关联文件；孤儿文件定期巡检"],
        ["系统日志", "建议180天，安全事件按制度延长", "集中或每日备份", "到期安全删除；管理员访问留痕"],
        ["Chroma向量", "与知识项一致", "定期快照；模型/维度切换可重建", "知识删除后删除向量并对账"],
        ["采集任务", "建议180天", "数据库备份", "到期归档/删除，保留重大失败记录"],
    ], widths=[3.1, 4.4, 4.2, 4.8], font_size=8.2)
    base.add_note(doc, "待定事项", "具体保留期限、注销权、更正与删除时限必须由实际部署主体依据处理目的和适用法规批准，当前建议值不自动成为最终政策。", kind="warning")


def build_nfr(doc: Document) -> None:
    doc.add_heading("7 非功能需求", level=1)
    for section_no, title, requirements in NFR_GROUPS:
        doc.add_heading(f"{section_no} {title}", level=2)
        add_req_table(doc, requirements)


def build_environment_constraints(doc: Document) -> None:
    doc.add_heading("8 运行环境与资源需求", level=1)
    doc.add_heading("8.1 计算机资源", level=2)
    base.add_table(doc, ["环境", "最低/建议配置", "说明"], [
        ["开发终端", "4核CPU、16GB内存、20GB可用空间", "运行IDE、后端、前端、MySQL和本地Chroma"],
        ["试运行服务器", "2～4 vCPU、4～8GB内存、50GB SSD", "20并发目标需在此级别或实际目标环境验证"],
        ["数据库", "MySQL 8.x、utf8mb4", "独立最小权限账号、事务与备份"],
        ["向量存储", "Chroma 1.x持久化目录", "容量取决于知识数量和向量维度"],
        ["浏览器", "现代Chrome/Edge", "启用JavaScript；建议1280×720以上"],
        ["网络", "访问模型、Embedding和搜索API", "公开部署必须HTTPS；设置出口和预算限制"],
    ], widths=[3.3, 6.2, 7.0])

    doc.add_heading("8.2 软件资源", level=2)
    base.add_table(doc, ["层次", "软件/版本基线"], [
        ["后端", "Python；FastAPI 0.111.0；Uvicorn 0.30.1；Pydantic随FastAPI依赖"],
        ["数据", "SQLAlchemy 2.0.31；Alembic 1.13.2；PyMySQL 1.1.1；MySQL 8.x"],
        ["安全", "python-jose 3.3.0；passlib 1.7.4；bcrypt 4.0.1"],
        ["向量/报告", "Chroma >=1,<2；Jinja2 3.1.6；xhtml2pdf 0.2.17"],
        ["前端", "Vue 3.5.13；Vite 6.x；Vue Router 4.5.0；Pinia 2.3.0；Axios 1.7.9"],
        ["UI/图表", "Element Plus 2.9.1；ECharts 5.6.0"],
    ], widths=[3.6, 12.9])

    doc.add_heading("8.3 配置需求", level=2)
    base.add_table(doc, ["类别", "关键变量", "要求"], [
        ["应用", "PROJECT_NAME、PROJECT_VERSION、API_PREFIX、APP_ENV、CORS", "环境隔离；生产白名单"],
        ["数据库", "DATABASE_URL或HOST/PORT/USER/PASSWORD/NAME", "密码外置；连接串不写日志"],
        ["认证", "SECRET_KEY、ALGORITHM、ACCESS_TOKEN_EXPIRE_MINUTES", "生产强随机密钥"],
        ["Embedding", "PROVIDER、DIMENSION、MODEL、API_KEY、TIMEOUT", "模型与维度一致；切换重建"],
        ["LLM", "DEEPSEEK_API_KEY、BASE_URL、MODEL、TIMEOUT", "契约兼容；预算与超时"],
        ["搜索/抓取", "WEB_SEARCH_*、BOCHA_*、CRAWL_*、ARTICLE_FETCH_*", "可关闭；SSRF默认拒绝"],
        ["存储", "CHROMA_PATH/PERSIST_DIR、REPORT_DIR", "源码外可写目录，纳入备份"],
        ["限流", "DETECT_RATE_LIMIT_COUNT/WINDOW", "按部署规模调整；多实例共享"],
    ], widths=[2.8, 7.7, 6.0], font_size=8.0)

    doc.add_heading("9 设计和实现约束", level=1)
    base.add_table(doc, ["约束编号", "约束"], [
        ["CON-01", "采用浏览器/REST的前后端分离模式；后端统一/api前缀和JSON响应。"],
        ["CON-02", "后端保持API→Service→CRUD→Model分层；Schema负责边界校验。"],
        ["CON-03", "结构化业务数据以MySQL为事实源，向量数据存Chroma，报告存受控文件目录。"],
        ["CON-04", "数据库结构以Alembic迁移为准，legacy SQL仅作历史参考。"],
        ["CON-05", "风险等级阈值和评分公式集中定义，前端不得独立复制业务判断。"],
        ["CON-06", "外部LLM必须满足结构化契约；供应商可替换但不得改变对外结果字段语义。"],
        ["CON-07", "真实数据库口令、JWT密钥和API Key不得提交到版本库。"],
        ["CON-08", "REPORT_DIR必须在后端源码目录外；非法路径不得生成或下载。"],
        ["CON-09", "正式RAG不得使用hash向量；Embedding模型/维度变更必须重建索引。"],
        ["CON-10", "自动判断必须保留免责声明和人工复核边界。"],
        ["CON-11", "当前范围不引入Redis、消息队列、对象存储、多租户或微服务，除非通过变更控制批准。"],
    ], widths=[2.5, 14.0], font_size=8.7)


def build_qualification(doc: Document) -> None:
    doc.add_heading("10 合格性规定", level=1)
    doc.add_heading("10.1 验证方法", level=2)
    base.add_table(doc, ["代码", "方法", "定义"], [
        ["T", "测试", "通过自动化或人工测试输入、操作和错误场景，比较实际与预期结果。"],
        ["D", "演示", "在代表性环境按操作流程展示功能和用户可见行为。"],
        ["I", "检查", "检查代码、配置、Schema、迁移、日志、文档或部署清单。"],
        ["A", "分析", "通过计算、模型、日志或测试报告证明容量、风险或质量属性。"],
    ], widths=[2.0, 3.0, 11.5])

    doc.add_heading("10.2 合格性矩阵", level=2)
    base.add_table(doc, ["需求组", "主要方法", "合格判据", "证据"], [
        ["FR-AUTH", "T/D/I", "注册登录、禁用、限流、JWT和管理员双重校验全部符合", "auth/admin用户测试、演示"],
        ["FR-DET/RAG/SCORE", "T/D/A", "输入、SSRF、检索、仲裁、公式、降级和持久化用例全部通过", "detect、web、LLM、rule测试"],
        ["FR-HIS/REP", "T/D", "所有权隔离、重新评估、唯一报告和安全下载通过", "history/report测试"],
        ["FR-KB/PRM", "T/D/I", "跨存储补偿、重建、模板契约和唯一默认通过", "knowledge/chroma/prompt测试"],
        ["FR-HR/PUB", "T/D", "三条件公开、字段白名单、审核转换和409冲突通过", "high-risk测试"],
        ["FR-ADM/STAT/LOG", "T/D", "管理权限、统计边界、日志查询与不阻断业务通过", "admin/statistics/log测试"],
        ["NFR-PERF", "T/A", "目标环境性能与20并发指标达到第7.1节", "性能测试报告"],
        ["NFR-SEC/PRI", "T/I/D", "无高危漏洞；密钥/CORS/SSRF/路径/隐私告知和人工复核条件落实", "安全测试与合规清单"],
        ["NFR-REL/OBS", "T/D/I", "备份恢复达到RPO/RTO；关键指标和告警可见", "恢复演练与监控截图"],
        ["NFR-USA/MNT", "D/I/T", "核心任务可完成；测试、构建、迁移和文档门禁通过", "可用性记录、CI/命令输出"],
    ], widths=[3.1, 2.5, 7.4, 3.5], font_size=8.1)

    doc.add_heading("10.3 基线测试与验收门槛", level=2)
    base.add_table(doc, ["项目", "当前基线", "验收门槛"], [
        ["后端自动化测试", "392项通过", "全部通过；阻断/严重缺陷为0"],
        ["前端自动化测试", "14项通过", "全部通过；关键页面交互补充回归"],
        ["前端生产构建", "成功，1,650个模块；有大分块告警", "构建成功；告警修复或风险接受"],
        ["功能验收", "核心闭环已实现", "本文件所有M级I/A需求通过"],
        ["模型效果", "未建立批准的真值阈值", "完成代表性标注集；指标由TBD-02批准"],
        ["性能与恢复", "尚无正式专项报告", "达到7.1与7.2目标并形成报告"],
        ["安全与合规", "已有鉴权、限流、SSRF、审计基础", "生产配置、安全测试、告知与审核流程完成"],
    ], widths=[3.6, 6.2, 6.7], font_size=8.5)


def build_traceability(doc: Document) -> None:
    doc.add_heading("11 需求可追踪性", level=1)
    doc.add_heading("11.1 业务目标到需求", level=2)
    base.add_table(doc, ["业务目标", "需求组", "验证证据"], [
        ["BG-01 降低初筛时间", "FR-DET、FR-RAG、FR-SCORE、NFR-PERF", "检测演示、性能报告"],
        ["BG-02 可解释可追溯", "FR-RAG-009～014、FR-SCORE-008～010、FR-HIS", "契约测试、详情页、数据库检查"],
        ["BG-03 高风险受控公开", "FR-HR、FR-PUB、NFR-PRI-005～006", "高风险API/页面测试"],
        ["BG-04 知识与报告复用", "FR-KB、FR-REP、FR-STAT", "知识同步、报告和统计测试"],
        ["BG-05 工程资产", "NFR-MNT、CON、合格性规定", "迁移、测试、构建与文档审查"],
    ], widths=[4.5, 6.5, 5.5], font_size=8.6)

    doc.add_heading("11.2 需求到实现与测试", level=2)
    base.add_table(doc, ["需求组", "主要实现模块", "主要测试/证据"], [
        ["FR-AUTH", "api/v1/auth.py；services/auth_service.py；core/security.py/deps.py", "test_auth_smoke.py、test_admin_users_api.py、secret_key测试"],
        ["FR-DET", "api/v1/detect.py；schemas/detection.py；web_content_fetcher.py", "test_detect_api.py、test_web_content_fetcher.py"],
        ["FR-RAG/SCORE", "detection_service.py；chroma/embedding/llm/rule/web_search服务", "test_detection_service相关、test_llm、test_chroma、test_rule、test_web_search"],
        ["FR-HIS/REP", "detection_crud.py；report_service.py；api/detect/report", "test_detection_history_api.py、test_report_api/service.py"],
        ["FR-KB", "knowledge_service/crud.py；chroma_service.py；admin_knowledge.py", "test_knowledge_sync.py、test_chroma_service/integration.py"],
        ["FR-PRM", "prompt_service/validator/crud.py；admin_prompts.py", "test_prompt_service.py、test_admin_prompts_api.py"],
        ["FR-HR/PUB", "high_risk_service/crud/utils.py；high_risk/admin_high_risk.py", "test_high_risk_api/service/utils.py"],
        ["FR-ADM/STAT/LOG", "admin_*、statistics/system_log服务与CRUD", "admin_*、statistics、system_log测试"],
        ["FR-CRAWL", "core/scheduler.py；services/web/news_crawler.py；crawl_task.py", "采集服务与Web提取测试；需补充调度集成测试"],
        ["NFR", "core/config/rate_limit/security；迁移；前端路由/工具", "性能/安全/恢复待验收；现有392＋14测试和build"],
    ], widths=[3.0, 7.1, 6.4], font_size=7.9)

    doc.add_heading("11.3 文档链追踪", level=2)
    base.add_table(doc, ["上游/下游文档", "关系"], [
        ["NCE-FS-001 可行性研究报告", "提供项目边界、推荐方案、风险和继续实施条件"],
        ["NCE-SRS-001 本需求分析书", "建立可验证需求基线"],
        ["后续概要设计书", "把需求分配到架构、模块、接口和数据库设计"],
        ["后续详细设计书", "把需求细化为算法、类、流程、异常和数据结构"],
        ["后续验收总结报告", "逐项引用需求编号和合格性证据，给出通过/偏差/遗留"],
    ], widths=[5.2, 11.3])


def build_unresolved_appendices(doc: Document) -> None:
    doc.add_heading("12 尚未解决的问题", level=1)
    base.add_table(doc, ["编号", "待定问题", "需要的决策", "责任建议", "最晚时点"], [
        ["TBD-01", "用户、检测、日志、报告和采集数据保留期限", "批准具体期限、注销/删除/更正SLA", "项目负责人＋合规", "公开部署前"],
        ["TBD-02", "模型效果验收阈值", "确定标注集规模、准确率/召回率/F1、一致性和人工复核标准", "领域负责人＋测试", "验收测试前"],
        ["TBD-03", "目标并发和数据规模", "确认日请求量、峰值并发、知识量和报告量", "项目负责人＋运维", "性能测试前"],
        ["TBD-04", "外部API供应商SLA与预算", "确认模型、向量、搜索供应商、配额、价格和数据条款", "项目负责人", "生产配置前"],
        ["TBD-05", "公开高风险内容的申诉与更正流程", "确定入口、责任人、处理时限和留痕", "内容负责人＋合规", "公开功能启用前"],
        ["TBD-06", "生产部署拓扑", "单机/云主机、HTTPS、共享限流、对象存储和监控方案", "架构/运维", "概要设计批准前"],
        ["TBD-07", "知识材料授权和网页抓取边界", "确认允许来源、摘要长度、robots/条款和删除机制", "内容负责人＋合规", "批量采集前"],
    ], widths=[1.6, 5.1, 5.3, 2.8, 2.2], font_size=7.9)
    base.add_note(doc, "基线规则", "任何TBD若影响M级需求、法律合规、数据库结构或验收指标，必须在相应设计/测试活动开始前关闭；不得由开发人员在代码中静默决定。", kind="warning")

    base.page_break(doc)
    doc.add_heading("附录A 需求统计与索引", level=1)
    total_fr = sum(len(group[2]) for group in REQ_GROUPS)
    total_nfr = sum(len(group[2]) for group in NFR_GROUPS)
    rows = []
    for section_no, title, reqs in REQ_GROUPS + NFR_GROUPS:
        must = sum(1 for r in reqs if r[2] == "M")
        implemented = sum(1 for r in reqs if "I" in r[3])
        rows.append([section_no, title, str(len(reqs)), str(must), str(implemented), f"{reqs[0][0]}～{reqs[-1][0]}"])
    rows.append(["合计", "功能＋非功能", str(total_fr + total_nfr), str(sum(int(r[3]) for r in rows)), str(sum(int(r[4]) for r in rows)), "—"])
    base.add_table(doc, ["章节", "需求组", "条数", "M级", "含I状态", "编号范围"], rows, widths=[1.4, 4.2, 1.2, 1.2, 1.4, 7.1], font_size=7.8)

    doc.add_heading("附录B 权限矩阵", level=1)
    base.add_table(doc, ["能力", "游客", "普通用户", "管理员", "维护人员"], [
        ["注册/登录", "允许", "允许", "允许", "非日常"],
        ["链接预览/新闻检测", "允许/限流", "允许", "允许", "可测试"],
        ["检测历史/详情/重新评估", "不允许", "仅本人", "全站", "仅受控排障"],
        ["报告生成/下载", "不允许", "仅本人", "全站", "文件备份"],
        ["公开高风险浏览", "允许", "允许", "允许", "允许"],
        ["用户/检测/知识/Prompt/高风险管理", "禁止", "禁止", "允许", "仅紧急运维"],
        ["统计/审计日志", "禁止", "禁止", "允许", "日志平台受控访问"],
        ["配置/迁移/备份恢复", "禁止", "禁止", "禁止或只读", "受控允许"],
    ], widths=[5.1, 2.6, 2.8, 2.8, 3.2], font_size=8.2)

    doc.add_heading("附录C 国标内容符合性对照", level=1)
    base.add_table(doc, ["规范内容", "本文件位置", "说明"], [
        ["标识、概述、文档说明和引用", "封面、控制页、第1章", "文档编号、基线、范围、术语和引用完整"],
        ["系统能力需求", "第3～4章", "用例、业务规则和编号化功能需求"],
        ["外部/内部接口需求", "第5章及第9章", "页面、REST、第三方、软件/文件接口和分层约束"],
        ["内部数据需求", "第6章", "实体、关系、状态、校验、保留和备份"],
        ["适应性、安全、保密和环境", "第7～8章", "性能、可靠、安全、隐私、易用、维护、资源和配置"],
        ["设计实现、人员培训和后勤约束", "第2、8、9章", "角色、环境、配置、分层、迁移和运行条件"],
        ["合格性规定", "第10章", "T/D/I/A方法、矩阵和门槛"],
        ["需求可追踪性", "第11章", "目标—需求—实现—测试—文档链"],
        ["尚未解决问题与注解", "第12章、附录", "TBD责任与时点、统计、权限和符合性"],
    ], widths=[5.0, 3.6, 7.9], font_size=8.5)

    doc.add_heading("附录D 需求基线声明", level=1)
    base.add_bullets(doc, [
        "本文件基于2026年6月20日工作区源码、设计资料和codebase知识图谱编制。",
        "知识图谱用于架构、实体、配置和调用链发现；其路由处理器关联不足时，按项目规则对路由装饰器与前端路由做了定向源码核对。",
        "当前实现状态I不代表所有生产质量属性已验收；A和T事项必须在后续概要设计、测试或部署决策中关闭。",
        "需求中的成本、性能、保留期限和模型效果阈值若与批准的可行性报告或部署政策冲突，应通过变更控制解决。",
        "后续设计、测试与验收文档必须引用稳定需求编号，不得只按页面名称或代码文件建立口头对应。",
    ])


def build_document() -> Path:
    configure_base()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    base.configure_styles(doc)
    base.configure_section(doc.sections[0], header_footer=False)
    build_cover(doc)

    front = doc.add_section(WD_SECTION.NEW_PAGE)
    base.configure_section(front, header_footer=True, page_format="lowerRoman")
    build_control(doc)
    build_abstract_toc(doc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    base.configure_section(body, header_footer=True, page_format="decimal")
    build_intro(doc)
    build_overview(doc)
    build_use_cases(doc)
    build_functional(doc)
    build_interfaces(doc)
    build_data(doc)
    build_nfr(doc)
    build_environment_constraints(doc)
    build_qualification(doc)
    build_traceability(doc)
    build_unresolved_appendices(doc)
    base.enable_update_fields(doc)

    doc.core_properties.title = f"《{PROJECT_TITLE}》{DOC_TITLE}"
    doc.core_properties.subject = "软件需求规格说明"
    doc.core_properties.author = f"《{PROJECT_TITLE}》项目组"
    doc.core_properties.keywords = "GB/T 9385-2008, GB/T 8567-2006, 软件需求, RAG, 大语言模型, 新闻真伪鉴别"
    doc.core_properties.comments = "基于源码、codebase知识图谱和当前验证基线编制"
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
