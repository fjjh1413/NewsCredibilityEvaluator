from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_outline_design_docx as base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = Path(r"E:\nan\《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》\开发管理")
OUTPUT_PATH = OUTPUT_DIR / "04_《基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计》详细设计书.docx"

PROJECT_TITLE = "基于 RAG 与大语言模型的网络新闻真伪鉴别系统设计"
SYSTEM_NAME = "智闻辨真"
DOC_TITLE = "详细设计书"
DOC_NO = "NCE-LLD-001"
VERSION = "V1.0"
BASELINE = "LLD-BL-2026-06-20"
DATE_TEXT = "2026 年 06 月 20 日"


# Reuse the typography, table, caption and figure helpers of the approved
# outline-design document so the document family has one visual language.
base.PROJECT_TITLE = PROJECT_TITLE
base.SYSTEM_NAME = SYSTEM_NAME
base.DOC_TITLE = DOC_TITLE
base.DOC_NO = DOC_NO
base.VERSION = VERSION
base.BASELINE = BASELINE
base.DATE_TEXT = DATE_TEXT


ENDPOINTS = [
    ["认证", "POST", "/api/auth/register", "register", "公开；按 IP 限流", "UserCreate", "CurrentUserApiResponse"],
    ["认证", "POST", "/api/auth/login", "login", "公开；按 IP 限流", "LoginRequest", "LoginApiResponse"],
    ["认证", "GET", "/api/auth/me", "read_current_user", "登录", "—", "CurrentUserApiResponse"],
    ["检测", "POST", "/api/detect/news", "detect_news", "游客/登录；限流", "DetectNewsRequest", "DetectNewsApiResponse"],
    ["检测", "POST", "/api/detect/extract-preview", "extract_preview", "游客/登录；限流", "ExtractPreviewRequest", "ExtractPreviewApiResponse"],
    ["检测", "GET", "/api/detect/history", "read_detection_history", "登录且数据隔离", "Query", "DetectionHistoryApiResponse"],
    ["检测", "GET", "/api/detect/{id}", "read_detection_detail", "本人/管理员", "Path", "DetectionDetailApiResponse"],
    ["检测", "POST", "/api/detect/{id}/re-evaluate", "re_evaluate_detection", "本人/管理员；限流", "Path", "DetectNewsApiResponse"],
    ["RAG", "POST", "/api/rag/search", "search_knowledge", "登录", "RagSearchRequest", "RagSearchApiResponse"],
    ["报告", "POST", "/api/report/generate/{detection_id}", "generate_report", "记录所有者/管理员", "Path", "ReportApiResponse"],
    ["报告", "GET", "/api/report/download/{report_id}", "download_report", "记录所有者/管理员", "Path", "FileResponse"],
    ["高风险公开", "GET", "/api/high-risk/public", "read_public_high_risk", "公开", "Query", "PublicHighRiskListApiResponse"],
    ["高风险公开", "GET", "/api/high-risk/ranking", "read_public_high_risk_ranking", "公开", "Query", "PublicHighRiskRankingApiResponse"],
    ["高风险公开", "GET", "/api/high-risk/keywords", "read_public_high_risk_keywords", "公开", "Query", "HighRiskKeywordsApiResponse"],
    ["高风险公开", "GET", "/api/high-risk/category-distribution", "read_public_high_risk_categories", "公开", "—", "CategoryDistributionApiResponse"],
    ["知识管理", "GET/POST", "/api/admin/knowledge", "read/create_knowledge", "管理员", "Query/KnowledgeCreate", "Knowledge*ApiResponse"],
    ["知识管理", "GET/PUT/DELETE", "/api/admin/knowledge/{id}", "read/update/delete_knowledge", "管理员", "Path/KnowledgeUpdate", "Knowledge*ApiResponse"],
    ["知识管理", "POST", "/api/admin/knowledge/{id}/vectorize", "vectorize_knowledge", "管理员", "Path", "KnowledgeItemApiResponse"],
    ["知识管理", "POST", "/api/admin/knowledge/rebuild-index", "rebuild_knowledge_vectors", "管理员", "—", "KnowledgeRebuildIndexApiResponse"],
    ["Prompt", "GET/POST", "/api/admin/prompts", "read/create_prompt", "管理员", "Query/PromptCreate", "Prompt*ApiResponse"],
    ["Prompt", "GET/PUT/DELETE", "/api/admin/prompts/{id}", "read/update/delete_prompt", "管理员", "Path/PromptUpdate", "Prompt*ApiResponse"],
    ["Prompt", "POST", "/api/admin/prompts/{id}/{action}", "enable/disable/set_default", "管理员", "Path", "PromptTemplateItemApiResponse"],
    ["高风险管理", "GET", "/api/admin/high-risk[/{id}]", "read_admin_high_risk*", "管理员", "Query/Path", "AdminHighRisk*ApiResponse"],
    ["高风险管理", "PUT", "/api/admin/high-risk/{id}/{action}", "review/public/remark", "管理员", "Path + Body", "AdminHighRiskDetailApiResponse"],
    ["用户管理", "GET", "/api/admin/users[/{id}]", "read_admin_user*", "管理员", "Query/Path", "AdminUser*ApiResponse"],
    ["用户管理", "POST", "/api/admin/users/{id}/{action}", "enable/disable/role", "管理员", "Path + Body", "AdminUserItemApiResponse"],
    ["检测管理", "GET/DELETE", "/api/admin/detections[/{id}]", "read/delete_admin_detection", "管理员", "Query/Path", "Detection*ApiResponse"],
    ["报告管理", "GET", "/api/admin/reports[/{id}][/download]", "read/download_admin_report", "管理员", "Query/Path", "Report/FileResponse"],
    ["统计", "GET", "/api/admin/statistics/*", "read_statistics_*", "管理员", "Query", "Statistics*ApiResponse"],
    ["日志", "GET", "/api/admin/logs", "read_admin_logs", "管理员", "Query", "SystemLogListApiResponse"],
    ["健康", "GET", "/api/health", "health_check", "公开", "—", "dict"],
]


FRONTEND_ROUTES = [
    ["/", "HomeView", "公开", "系统入口与功能引导"],
    ["/login", "LoginView", "仅游客", "登录并恢复会话"],
    ["/register", "RegisterView", "仅游客", "普通用户注册"],
    ["/detect", "DetectView", "公开", "手工输入或链接提取预览后检测"],
    ["/result/:id", "ResultView", "登录；支持游客临时缓存", "展示分数、证据、质量与建议"],
    ["/history", "HistoryView", "登录", "本人检测历史、筛选与重评"],
    ["/high-risk", "HighRiskView", "公开", "展示审核通过且公开的高风险内容"],
    ["/profile", "ProfileView", "登录", "当前账户信息"],
    ["/admin/dashboard", "AdminDashboardView", "管理员", "后台概览"],
    ["/admin/users", "AdminUsersView", "管理员", "用户启停与角色管理"],
    ["/admin/detections", "AdminDetectionsView", "管理员", "全局检测管理"],
    ["/admin/knowledge", "AdminKnowledgeView", "管理员", "知识及向量同步管理"],
    ["/admin/prompts", "AdminPromptsView", "管理员", "Prompt 生命周期管理"],
    ["/admin/high-risk", "AdminHighRiskView", "管理员", "高风险复核与公开控制"],
    ["/admin/statistics", "AdminStatisticsView", "管理员", "统计图表"],
    ["/admin/reports", "AdminReportsView", "管理员", "报告查询下载"],
    ["/admin/logs", "AdminLogsView", "管理员", "审计日志查询"],
]


PROGRAMS = [
    {
        "id": "P-01",
        "name": "应用启动与配置程序",
        "location": "backend/app/main.py；core/config.py；db/migrate.py；core/scheduler.py",
        "purpose": "建立 FastAPI 应用、装配路由与中间件，校验关键配置，连接数据库，并按配置启动后台采集调度。",
        "inputs": "环境变量、.env 配置、Alembic 版本、启动命令。",
        "outputs": "可服务的 ASGI 应用、API 路由、调度器状态和启动日志。",
        "dependencies": "FastAPI、Uvicorn、SQLAlchemy、Alembic、APScheduler。",
        "steps": ["读取并缓存 Settings；对 SECRET_KEY 等必需项执行启动校验。", "创建应用并配置 CORS、异常处理和统一 /api 路由。", "建立 MySQL 会话工厂；数据库结构以 Alembic head 为准。", "若启用定时采集，则注册周期任务；关闭应用时释放调度器资源。"],
        "exceptions": "关键秘密缺失时拒绝启动；迁移未到 head 时给出明确运维错误；调度任务失败只记录本次 crawl_tasks，不使 Web 服务退出。",
        "tests": "秘密校验、CORS 白名单、数据库迁移、启动/关闭幂等、单实例调度约束。",
        "limits": "当前限流和调度默认面向单实例；多实例部署须引入共享状态与分布式任务锁。",
    },
    {
        "id": "P-02",
        "name": "前端路由与会话程序",
        "location": "frontend/src/router/index.js；stores/user.js；utils/auth.js；utils/request.js",
        "purpose": "管理用户端/管理端页面导航、Token 注入、会话恢复、角色守卫和 401 失效处理。",
        "inputs": "目标路由、localStorage Token、/api/auth/me 返回、Axios 响应。",
        "outputs": "允许/重定向的导航结果、Authorization 请求头、Pinia 用户状态。",
        "dependencies": "Vue Router、Pinia、Axios、Element Plus。",
        "steps": ["首次导航调用 restoreSession，以服务器 /me 结果校验本地 Token。", "requiresAuth 且未登录时跳转登录页并保留 redirect。", "requiresAdmin 且角色非 admin 时提示并返回首页。", "请求拦截器附加 Bearer Token；响应 401 时清理会话并跳转登录。", "afterEach 根据 meta.title 更新浏览器标题。"],
        "exceptions": "网络失败不得把未验证缓存当作已登录；401 统一退出；403 由业务页面提示无权操作。",
        "tests": "无 Token、过期 Token、普通用户访问后台、管理员登录后跳转、刷新恢复、并发 401。",
        "limits": "前端守卫只改善交互，最终授权必须由后端依赖再次执行。",
    },
    {
        "id": "P-03",
        "name": "认证与授权程序",
        "location": "api/v1/auth.py；services/auth_service.py；core/security.py；core/deps.py",
        "purpose": "提供注册、登录、当前用户读取、bcrypt 密码校验、JWT 签发与 user/admin 权限依赖。",
        "inputs": "用户名、密码、可选邮箱、Bearer Token、客户端 IP。",
        "outputs": "用户对象、JWT access_token，或 401/403/409/422/429。",
        "dependencies": "python-jose、passlib/bcrypt、users 表、InMemoryRateLimiter。",
        "steps": ["Pydantic 校验注册或登录载荷，并按 IP 执行滑动窗口限流。", "注册时检查 username/email 唯一性，bcrypt 散列后保存普通用户。", "登录时读取用户、校验状态与密码，签发包含 sub、role、exp 的 JWT。", "受保护接口解码 Token，以 sub 查询实时用户并检查 active。", "管理员接口继续检查 role=admin，避免仅信任客户端或 Token 中陈旧角色。"],
        "exceptions": "凭据错误统一返回认证失败；禁用账户返回 403；重复用户名/邮箱返回冲突；无可用管理员时禁止降权/禁用最后管理员。",
        "tests": "密码不落明文、过期/篡改 Token、禁用用户、普通用户越权、限流边界、唯一约束竞争。",
        "limits": "未实现刷新 Token 与服务端撤销列表；权限变化依靠每次请求读取用户状态即时生效。",
    },
    {
        "id": "P-04",
        "name": "新闻检测 API 编排程序",
        "location": "api/v1/detect.py；schemas/detection.py",
        "purpose": "接收检测、链接提取、历史、详情和重新评估请求，完成参数/权限/限流/错误映射与日志记录。",
        "inputs": "DetectNewsRequest 或 ExtractPreviewRequest、当前用户、分页/路径参数。",
        "outputs": "统一 {code,message,data} 响应、检测审计日志、HTTP 状态码。",
        "dependencies": "detection_service、WebContentFetcher、detection_crud、system_log_service。",
        "steps": ["标题 4～255 字符、正文至少 20 字符；禁止未声明字段。", "新闻检测允许游客，但登录用户绑定 user_id；按客户端 IP 限流。", "链接模式先安全抓取并返回预览，用户可编辑后再提交检测。", "历史和详情按当前用户收敛数据；管理员可访问全局范围。", "重新评估复制历史输入形成新记录，原记录保持不可变以便审计。"],
        "exceptions": "知识检索不可用映射 503；业务输入错误映射 400；抓取/SSRF 拦截映射 422；不存在或越权统一按 404 处理以减少枚举。",
        "tests": "游客/用户双路径、非法长度、extra 字段、限流、越权 ID、历史旧数据重评、SSRF 地址。",
        "limits": "游客完整结果依赖前端临时缓存；服务端历史仅对已绑定用户持久化展示。",
    },
    {
        "id": "P-05",
        "name": "新闻可信度检测核心程序",
        "location": "services/detection_service.py::detect_news_credibility",
        "purpose": "串联文本清洗、关键词、本地检索、联网补证、LLM 分析/仲裁、规则评分、风险分级和持久化。",
        "inputs": "数据库会话、DetectNewsRequest、可选当前用户。",
        "outputs": "DetectNewsResult；detection_records 与 evidence_matches 新记录。",
        "dependencies": "knowledge_service、web_search_service、llm_service、rule_score_service、detection_crud。",
        "steps": ["清洗标题（≤255）和正文（≤12000），抽取最多 8 个关键词。", "本地 Chroma 召回 Top10 并格式化候选；满足触发条件时调用博查联网补证并去重合并。", "为候选分配 kb:<id>/web:<n> 标识；以 SHA-256(seed+candidate_id) 形成可复现且来源中立的输入顺序。", "调用 DeepSeek 取得结构化分析、证据质量与仲裁；验证候选全集覆盖、唯一性、0～100 分数、stance 和理由。", "首次契约失败时执行一次聚焦仲裁重试；仍失败则候选不得参与评分。", "计算规则分、证据分、模型分与最终分，映射四级风险并合并理由、风险点、关键词和相似新闻。", "将有效证据写 evidence_matches，将完整候选、排除理由、质量、契约版本和降级状态写 analysis_payload。"],
        "exceptions": "本地知识检索失败终止并返回 503；联网失败回退本地证据；LLM 失败进入确定性规则降级；持久化失败回滚事务。",
        "tests": "无证据、有证据、联网补证、契约一次成功/重试成功/重试耗尽、LLM 失败、候选遗漏/重复、分数边界、持久化回滚。",
        "limits": "输出是辅助评估；最终质量受知识覆盖、网页质量与模型服务影响。",
    },
    {
        "id": "P-06",
        "name": "证据仲裁与评分程序",
        "location": "detection_service.py::validate_and_apply_llm_ranking；calculate_evidence_score",
        "purpose": "验证模型证据选择，确保只有可追踪、契约有效的证据参与结果、评分和展示。",
        "inputs": "候选证据列表、evidence_arbitration、evidence_quality、LLM/规则分。",
        "outputs": "ranked、rejected、errors、最终分与风险等级。",
        "dependencies": "llm_service 输出契约、risk_level 工具。",
        "steps": ["建立 candidate_id→候选映射。", "逐项校验 ranked/rejected 类型、存在性、唯一性、互斥性和全集覆盖。", "校验 relevance_score/quality_score∈[0,100]、stance∈{support,contradict,neutral}、reason 非空。", "按模型数组顺序重新赋 rank_order；排除证据写入 rejection_reason。", "证据质量分由服务端按 coverage×0.6+consistency×0.4 计算，不信任模型自报聚合值。", "根据证据可用性和降级状态选择评分公式，再映射风险等级。"],
        "exceptions": "任何仲裁结构错误都使本次仲裁整体无效；不能用原始检索相似度绕过仲裁。",
        "tests": "未知/重复/遗漏 candidate_id、非法 stance、布尔值伪装数字、分数越界、空理由、质量字段缺失。",
        "limits": "仲裁重试最多一次；避免不可控延迟与重复成本。",
    },
    {
        "id": "P-07",
        "name": "知识库与向量同步程序",
        "location": "services/knowledge_service.py；crud/knowledge_crud.py；services/chroma_service.py",
        "purpose": "维护 MySQL 知识事实和 Chroma 派生向量的一致性，提供增删改、单条向量化、检索与全量重建。",
        "inputs": "KnowledgeCreate/Update、知识 ID、查询文本、top_k 与过滤条件。",
        "outputs": "知识记录、vector_sync_status、向量 ID、相似证据列表或同步错误。",
        "dependencies": "knowledge_items 表、EmbeddingService、Chroma PersistentClient。",
        "steps": ["组合 title/content/summary/辟谣说明形成 embedding_text。", "先保存 MySQL 事实并置 pending，再计算向量并 upsert 到 knowledge_items 集合。", "成功写回 vector_id=knowledge:<id> 与 synced；失败记录 failed 和错误摘要。", "删除时先尝试删除向量；失败保留 delete_failed 状态，避免静默不一致。", "检索以余弦空间查询并将 distance 转换为 max(0,1-distance)。", "全量重建先重置集合，再逐项向量化并统计成功/失败。"],
        "exceptions": "Chroma 操作失败会清缓存并重试一次；依赖未安装为不可重试错误；状态字段保留补偿入口。",
        "tests": "创建/更新/删除的双存储一致性、重试、维度变化重建、top_k 1～50、过滤组合、失败状态可见。",
        "limits": "MySQL 与 Chroma 无分布式事务，采用事实源+派生索引+可重建的最终一致性。",
    },
    {
        "id": "P-08",
        "name": "Embedding 程序",
        "location": "services/embedding_service.py",
        "purpose": "把文本转换为固定维向量，并统一批量请求、维度校验、归一化和本地演示后备。",
        "inputs": "清洗文本、provider、model、dimension、API Key、timeout。",
        "outputs": "float 向量或批量向量；配置/服务错误。",
        "dependencies": "DashScope text-embedding-v4（正式推荐）、可配置兼容端点、hash fallback。",
        "steps": ["拒绝空文本并解析 provider 与维度。", "正式路径调用远程 embedding 接口，按索引复原批量结果。", "验证结果条数、数值类型、有限性与维度一致。", "hash 后备按 token 哈希到桶并归一化，仅供无 Key 的本地演示。"],
        "exceptions": "HTTP、超时、响应结构、维度不匹配均转换为领域错误；切换 provider/维度后必须全量重建。",
        "tests": "中文分词输入、空文本、批量乱序、非有限值、维度错误、hash 确定性。",
        "limits": "hash 不具语义能力，不得作为正式检索质量结论。",
    },
    {
        "id": "P-09",
        "name": "联网搜索与网页提取程序",
        "location": "services/web/bocha_client.py；web_search_service.py；web_content_fetcher.py；news_crawler.py",
        "purpose": "在本地证据不足时搜索补证，并安全地从用户链接或采集结果提取新闻正文与发布时间。",
        "inputs": "标题、关键词、freshness、count 或 HTTP(S) URL。",
        "outputs": "标准化 web_search 候选，或标题/正文/来源/发布时间预览。",
        "dependencies": "博查搜索 API、HTTP 客户端、BeautifulSoup（可选）。",
        "steps": ["仅在用户允许且本地证据触发条件满足时搜索；组合查询并限制返回数。", "规范化标题、摘要、URL、来源和时间，按 URL/标题与本地候选去重。", "链接提取只接受 HTTP(S)，解析 DNS/IP 并拦截本地、私网、保留地址及危险重定向。", "限制响应体、超时和重定向；优先 article/main 语义区提取正文。", "发布时间按 JSON-LD、meta、time、标题邻域等候选质量排序，输出 date/datetime 精度。"],
        "exceptions": "搜索失败不阻断本地检测；链接抓取失败和 SSRF 拦截返回 422；采集任务记录 partial/failed。",
        "tests": "私网 IPv4/IPv6、DNS 重绑定、重定向、超大响应、非 HTML、乱码、JSON-LD 时间、重复搜索结果。",
        "limits": "外网页面结构和可访问性不可控；提取结果必须先预览再由用户提交检测。",
    },
    {
        "id": "P-10",
        "name": "LLM 与 Prompt 程序",
        "location": "services/llm_service.py；prompt_service.py；prompt_template_validator.py",
        "purpose": "选择启用的默认模板，安全填充上下文，调用 DeepSeek，并把不稳定文本解析为受控结构。",
        "inputs": "新闻、脱敏后的证据候选、Prompt 模板、模型配置。",
        "outputs": "llm_score、risk_level、reason、evidence_quality、evidence_arbitration、similar_news 等。",
        "dependencies": "DeepSeek Chat API、prompt_templates 表。",
        "steps": ["在事务中选取同 type 的启用默认模板；无数据库模板时使用代码默认模板。", "填充标题、正文和候选证据，并附加输出契约版本与只输出 JSON 的约束。", "调用模型并提取 JSON；对兼容文本格式执行有限解析与清洗。", "coverage/consistency 归一化后由服务端计算质量总分。", "模型异常返回带 LLM_ANALYSIS_FAILED 标识的结构化错误，交由核心程序降级。"],
        "exceptions": "模板缺占位符或结构非法时拒绝保存；默认模板切换加行锁；模型超时/非 JSON/空响应不泄露密钥。",
        "tests": "默认模板唯一、并发设默认、JSON 代码块、字段缺失、非法数值、超时、错误体清洗、提示注入隔离。",
        "limits": "模型输出不直接作为数据库命令或权限判断；所有关键字段由服务端再验证。",
    },
    {
        "id": "P-11",
        "name": "确定性规则评分程序",
        "location": "services/rule_score_service.py",
        "purpose": "以可解释规则补充模型判断，识别来源缺失、标题夸张、情绪化、绝对化及证据冲突。",
        "inputs": "标题、正文、来源名、经仲裁的有效证据。",
        "outputs": "rule_score（0～100）、hit_rules 列表。",
        "dependencies": "规则词表和证据 truth_label/stance。",
        "steps": ["初始分 100。", "来源缺失扣分；统计标题/正文中的夸张、情绪化、绝对化表达并按上限扣分。", "根据有效证据中的反驳标签或 contradict stance 扣分。", "分数截断至 0～100，命中原因去重后返回。"],
        "exceptions": "空值按无来源或无命中处理；只接受仲裁有效证据，避免原始召回污染规则分。",
        "tests": "每条规则单独命中、组合上限、大小写/中文标点、无证据、冲突证据、0/100 边界。",
        "limits": "词表规则只能提示风险，不等于事实裁决；需由配置评审后扩充。",
    },
    {
        "id": "P-12",
        "name": "检测持久化程序",
        "location": "crud/detection_crud.py；models/detection_record.py；models/evidence_match.py",
        "purpose": "原子保存检测快照和有效证据，执行用户数据隔离、分页筛选、详情加载与级联删除。",
        "inputs": "DetectionCreate、有效证据、当前用户、筛选条件。",
        "outputs": "DetectionRecord、历史页、详情关系、删除结果。",
        "dependencies": "MySQL、SQLAlchemy Session。",
        "steps": ["创建 detection_records 并 flush 取得主键。", "逐条创建 evidence_matches 快照，联网证据 knowledge_id 为空且 origin=web。", "提交事务并 refresh；异常统一 rollback。", "历史查询由 _apply_user_scope 强制本人范围；管理员才可指定全局用户。", "删除检测由外键级联证据和报告元数据，文件删除由报告服务处理。"],
        "exceptions": "JSON 字段序列化失败、约束冲突或数据库异常均回滚；越权查询不返回对象。",
        "tests": "主子表原子性、游客 NULL user_id、分页排序、关键词筛选、用户隔离、级联删除。",
        "limits": "analysis_payload 为 JSON 文本快照，跨版本读取必须容忍缺失字段。",
    },
    {
        "id": "P-13",
        "name": "报告生成与下载程序",
        "location": "services/report_service.py；templates/reports/detection_report.html；api/v1/report.py",
        "purpose": "从检测快照生成 HTML/PDF，保存元数据，并按所有权安全下载。",
        "inputs": "detection_id、当前用户、模板、REPORT_DIR。",
        "outputs": "reports 记录、HTML/PDF 文件、FileResponse。",
        "dependencies": "Jinja2、WeasyPrint/HTML 渲染能力、reports 表、文件系统。",
        "steps": ["查询检测及证据并校验本人或管理员权限。", "把分数、理由、证据和免责声明渲染为 HTML。", "在受控目录生成确定文件名；转换 PDF。", "按 detection_id 幂等更新或创建报告记录。", "下载前解析真实路径并确认位于 REPORT_DIR 内，设置 PDF 媒体类型。"],
        "exceptions": "检测不存在/越权、模板错误、PDF 依赖或文件缺失分别映射领域错误；数据库失败回滚。",
        "tests": "一检测一报告、重复生成、路径穿越、越权下载、中文字体、文件缺失、渲染失败清理。",
        "limits": "报告是检测时点快照；重新评估产生新 detection_id 和新报告。",
    },
    {
        "id": "P-14",
        "name": "高风险审核与公开程序",
        "location": "services/high_risk_service.py；crud/high_risk_crud.py；api/v1/high_risk.py；admin_high_risk.py",
        "purpose": "对高风险检测执行人工复核、公开控制，并提供只含已审核公开数据的前台榜单。",
        "inputs": "检测分数/等级、review_status、is_public、admin_remark、管理员。",
        "outputs": "审核状态、公开列表/排行/关键词/分类，审计日志。",
        "dependencies": "detection_records、users、system_logs。",
        "steps": ["保存检测时 final_score<40 或 risk_level=高风险谣言则 is_high_risk=true。", "新高风险记录默认 pending 且不公开。", "管理员可 approved/rejected，并记录 reviewed_by/reviewed_at。", "只有 approved 记录才允许 is_public=true；拒绝后强制撤销公开。", "公开查询统一追加 is_high_risk、approved、is_public 三重条件。"],
        "exceptions": "非高风险记录不可进入审核；未通过审核不可公开；并发状态冲突返回 409。",
        "tests": "40 分边界、等级触发、先公开后审核禁止、驳回撤公开、普通用户管理越权、公开查询无泄漏。",
        "limits": "人工审核结论与模型结论分层保存，管理员备注不对公众输出。",
    },
    {
        "id": "P-15",
        "name": "后台管理、统计与审计程序",
        "location": "api/v1/admin_*.py；services/statistics_service.py；system_log_service.py；admin_user_service.py",
        "purpose": "提供用户、检测、Prompt、知识、报告、统计和日志的受控管理能力。",
        "inputs": "管理员身份、分页筛选、状态变更、日期区间。",
        "outputs": "管理列表/详情、聚合统计、审计日志和业务状态变更。",
        "dependencies": "全部业务表、get_current_admin。",
        "steps": ["所有路由先执行实时管理员依赖。", "列表接口统一分页、白名单排序和可选筛选。", "用户启停/角色变更保护当前管理员和最后可用管理员。", "统计按日期、风险、分类、关键词和用户聚合，空数据返回零值结构。", "关键写操作记录 user_id、module、action、description、IP，描述不得含 Token/密码/API Key。"],
        "exceptions": "非法日期/分页返回 422；目标不存在返回 404；业务冲突返回 409；统计不因空集报错。",
        "tests": "管理员权限全覆盖、最后管理员保护、聚合口径、时区边界、日志脱敏、分页上限。",
        "limits": "审计日志当前记录关键业务事件，不替代基础设施访问日志和集中式 SIEM。",
    },
]


def add_code_block(doc: Document, text: str, caption: str | None = None) -> None:
    if caption:
        base.add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    base.set_cell_shading(cell, "F3F5F7")
    base.set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(text.strip().splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        base.set_run_font(run, east_asia="等线", latin="Consolas", size=8.8, color=base.DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    base.configure_section(section, header_footer=False)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("SOFTWARE ENGINEERING DOCUMENT")
    base.set_run_font(run, east_asia="微软雅黑", latin="Arial", size=9, bold=True, color=base.TEAL)
    bar = doc.add_table(rows=1, cols=1)
    base.set_cell_shading(bar.cell(0, 0), base.TEAL)
    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(f"《{PROJECT_TITLE}》")
    base.set_run_font(run, east_asia="黑体", latin="Arial", size=20, bold=True, color=base.NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run(DOC_TITLE)
    base.set_run_font(run, east_asia="微软雅黑", latin="Arial", size=30, bold=True, color=base.BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("DETAILED DESIGN DESCRIPTION")
    base.set_run_font(run, east_asia="微软雅黑", latin="Arial", size=11, bold=True, color=base.GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, (label, value) in enumerate([
        ("文档编号", DOC_NO), ("版本号", VERSION), ("配置基线", BASELINE),
        ("文档状态", "正式版"), ("编制单位", "项目开发组"), ("发布日期", DATE_TEXT),
    ]):
        base.set_cell_shading(table.cell(i, 0), base.PALE_BLUE)
        for j, value_text in enumerate((label, value)):
            cell = table.cell(i, j)
            base.set_cell_margins(cell, top=130, bottom=130)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value_text)
            base.set_run_font(run, east_asia="黑体" if j == 0 else "宋体", size=10.5, bold=(j == 0), color=base.NAVY if j == 0 else base.DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("受控文档 · 未经批准不得擅自修改")
    base.set_run_font(run, size=9.5, color=base.GRAY)


def add_front_matter(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    base.configure_section(section, header_footer=True, page_format="lowerRoman", page_start=1)
    h = base.add_heading(doc, "文档控制信息", 1)
    h.paragraph_format.page_break_before = False
    base.add_body(doc, "本文件是项目详细设计阶段的受控技术文件。它把需求和概要设计落实为程序单元、处理流程、数据结构、接口契约、异常策略和单元测试要点；编码、测试、部署及变更评审应以本文件和当前代码基线共同为准。")
    base.add_heading(doc, "文档审批", 2)
    base.add_table(doc, ["角色", "姓名/单位", "职责", "签字", "日期"], [
        ["编制", "项目开发组", "详细设计、实现核对与排版", "", DATE_TEXT],
        ["审核", "", "程序、接口、数据库、安全与可测试性审查", "", ""],
        ["批准", "", "批准进入实现/测试基线", "", ""],
    ], [2.1, 3.1, 5.5, 2.4, 2.7])
    base.add_heading(doc, "修订记录", 2)
    base.add_table(doc, ["版本", "日期", "修订内容", "修订人", "状态"], [
        [VERSION, DATE_TEXT, "首次发布；依据当前前后端源码、数据模型、Alembic 迁移、测试和前三类项目文档形成详细设计基线。", "项目开发组", "正式"],
    ], [2.0, 2.8, 8.3, 2.5, 2.1])
    base.add_heading(doc, "分发范围", 2)
    base.add_table(doc, ["序号", "接收方", "用途", "介质"], [
        ["1", "项目负责人/指导教师", "详细设计评审与阶段验收", "电子版"],
        ["2", "开发与测试人员", "编码、单元测试、集成测试依据", "电子版"],
        ["3", "配置管理员", "基线、发布包和变更归档", "电子版"],
    ], [1.6, 4.4, 7.6, 3.2])
    h = base.add_heading(doc, "摘要", 1)
    h.paragraph_format.page_break_before = False
    base.add_body(doc, f"本文档描述{SYSTEM_NAME}的低层设计。系统采用 Vue 3 + FastAPI + MySQL + Chroma 的前后端分离结构，使用 DeepSeek 完成结构化可信度分析，使用可重建的向量索引完成 RAG 检索，并以确定性规则、证据质量、人工高风险复核和审计快照形成可解释闭环。")
    base.add_body(doc, "文档按 GB/T 8567—2006《计算机软件文档编制规范》中详细设计说明的核心内容组织：给出软件标识、程序系统组织、每个程序单元的标识/目的/输入/输出/算法/接口/存储/限制/测试，补充数据库物理设计、接口契约、安全、异常、性能与需求追踪。")
    base.add_note(doc, "基线说明", "详细设计以 2026-06-20 工作区实现为事实基线；若代码、迁移或接口发生变化，应同步修订本文档并形成新版本。", "info")
    base.add_note(doc, "责任边界", "系统结果是辅助评估，不替代人工事实核查、官方通报、司法或行政结论。", "warning")
    base.add_toc(doc)


def add_program(doc: Document, index: int, program: dict[str, object]) -> None:
    base.add_heading(doc, f"6.{index} {program['id']} {program['name']}", 2)
    base.add_table(doc, ["设计项", "说明"], [
        ["程序标识", str(program["id"])],
        ["实现位置", str(program["location"])],
        ["目的", str(program["purpose"])],
        ["输入", str(program["inputs"])],
        ["输出", str(program["outputs"])],
        ["依赖/接口", str(program["dependencies"])],
    ], [3.4, 13.7], font_size=9.0, first_col_bold=True)
    base.add_heading(doc, f"6.{index}.1 处理逻辑", 3)
    base.add_numbered(doc, list(program["steps"]))
    base.add_heading(doc, f"6.{index}.2 异常、测试与限制", 3)
    base.add_table(doc, ["类别", "详细设计"], [
        ["异常与恢复", str(program["exceptions"])],
        ["单元/集成测试点", str(program["tests"])],
        ["限制与约束", str(program["limits"])],
    ], [3.4, 13.7], font_size=9.0, first_col_bold=True)


def add_main_content(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    base.configure_section(section, header_footer=True, page_format="decimal", page_start=1)

    base.add_heading(doc, "1 引言", 1)
    base.add_heading(doc, "1.1 软件与文档标识", 2)
    base.add_table(doc, ["项目", "内容"], [
        ["软件名称", f"{SYSTEM_NAME}——{PROJECT_TITLE}"], ["文档名称", DOC_TITLE],
        ["文档编号", DOC_NO], ["版本/基线", f"{VERSION} / {BASELINE}"],
        ["软件形态", "B/S Web 应用；用户端、管理端、REST API、关系库、向量库与报告文件"],
        ["目标读者", "开发、测试、运维、项目管理、评审和验收人员"],
    ], [4.0, 13.1], first_col_bold=True)
    base.add_heading(doc, "1.2 编写目的", 2)
    base.add_body(doc, "把概要设计中的模块边界进一步分解为可编码和可验证的程序单元，明确每个单元的处理步骤、输入输出、数据访问、异常恢复、性能约束和测试要点；消除实现歧义，并为单元测试、集成测试、维护和验收提供追踪依据。")
    base.add_heading(doc, "1.3 范围", 2)
    base.add_body(doc, "范围包括用户注册登录、新闻手工/链接输入、RAG 检索、按需联网补证、LLM 结构化分析、证据仲裁、规则与综合评分、历史/重评、报告、高风险审核、知识/Prompt/用户/报告/日志管理、统计、定时采集及其数据库、向量库和文件存储设计。")
    base.add_heading(doc, "1.4 术语和缩略语", 2)
    base.add_table(doc, ["术语", "含义"], [
        ["RAG", "检索增强生成；先取证据再由模型分析"], ["LLM", "大语言模型，本系统主要对接 DeepSeek"],
        ["Embedding", "将文本编码为固定维浮点向量"], ["证据仲裁", "模型对候选证据的相关性、质量、立场和取舍进行结构化判断"],
        ["候选证据", "本地知识或联网搜索得到、尚未通过仲裁的材料"], ["有效证据", "通过契约校验并被 ranked_evidence 选中的证据"],
        ["派生索引", "可由 MySQL 知识事实重建的 Chroma 向量数据"], ["降级", "模型不可用时仅以确定性规则产生保守结果"],
        ["SSRF", "服务端请求伪造；恶意 URL 诱导服务器访问内网资源"], ["DTO/Schema", "Pydantic 定义的请求/响应数据契约"],
    ], [3.2, 13.9])
    base.add_heading(doc, "1.5 参考资料与编制依据", 2)
    base.add_table(doc, ["编号", "资料", "用途"], [
        ["R-01", "GB/T 8567—2006《计算机软件文档编制规范》", "文档内容和控制要素"],
        ["R-02", "01 可行性研究报告、02 需求分析书、03 概要设计书", "上位项目基线"],
        ["R-03", "docs/01_project_design.md、02_database_design.md、03_api_design.md", "项目设计资料"],
        ["R-04", "backend/app、frontend/src、backend/alembic/versions", "实现与数据库事实基线"],
        ["R-05", "backend/tests、frontend/src/utils/*.test.js", "可测试性与回归基线"],
    ], [1.8, 8.1, 7.2])

    base.add_heading(doc, "2 详细设计原则与约定", 1)
    base.add_heading(doc, "2.1 设计原则", 2)
    base.add_bullets(doc, [
        "事实源唯一：MySQL 保存业务事实，Chroma 是可重建派生索引，报告文件是受控输出。",
        "分层单向依赖：API/视图 → 服务/状态 → CRUD/基础设施；路由不承载核心算法。",
        "零信任模型输出：所有模型分数、候选标识、枚举和 JSON 结构必须在服务端验证。",
        "最小权限：公开、登录、管理员三层权限；前端守卫不代替后端授权。",
        "可审计：保存输入快照、候选/排除证据、契约版本、降级状态和管理员关键操作。",
        "失败可见：外部依赖失败不得伪装成功；允许的降级必须明确标注。",
        "配置与秘密分离：环境变量提供密钥、端点、超时和路径，仓库只保留无秘密模板。",
    ])
    base.add_heading(doc, "2.2 命名和编码约定", 2)
    base.add_table(doc, ["对象", "约定", "示例"], [
        ["Python 模块/函数", "snake_case；领域服务使用动词短语", "detect_news_credibility"],
        ["Python 类/Schema", "PascalCase；请求/响应后缀清晰", "DetectNewsRequest"],
        ["Vue 组件", "PascalCase.vue；页面以 View 结尾", "AdminKnowledgeView.vue"],
        ["JS 变量/函数", "camelCase；API 文件按领域分组", "restoreSession"],
        ["REST 路径", "小写名词和短横线；动作仅用于明确命令", "/admin/knowledge/rebuild-index"],
        ["数据库", "表名复数 snake_case；索引含表/字段语义", "detection_records / idx_detection_created_at"],
        ["时间", "数据库 DATETIME；API 使用 ISO 8601/JSON 序列化", "2026-06-20T10:30:00"],
        ["分数", "服务层 float 0～100；相似度 0～1；数据库 DECIMAL", "final_score=68.50"],
    ], [3.0, 7.5, 6.6])
    base.add_heading(doc, "2.3 统一响应和错误约定", 2)
    add_code_block(doc, '{\n  "code": 200,\n  "message": "success",\n  "data": {}\n}', "统一成功响应")
    base.add_table(doc, ["HTTP", "场景", "处理约定"], [
        ["400", "可识别的业务输入错误", "返回中文可操作信息，不包含堆栈"],
        ["401", "Token 缺失/无效/过期", "WWW-Authenticate: Bearer；前端清会话"],
        ["403", "账户禁用或角色不足", "不执行目标业务"], ["404", "对象不存在或无权访问", "对象级授权失败采用同一结果"],
        ["409", "状态/唯一性冲突", "提示刷新或更换输入"], ["422", "Schema、URL 抓取或日期校验失败", "保留字段级错误语义"],
        ["429", "超过接口限流", "拒绝并提示稍后重试"], ["503", "本地知识检索等关键依赖不可用", "检测不产生伪结果"],
    ], [2.0, 6.0, 9.1])

    base.add_heading(doc, "3 程序系统组织", 1)
    base.add_heading(doc, "3.1 分层和进程", 2)
    base.add_figure(doc, "fig4-2-backend-layered-architecture.png", "图 3-1 后端程序分层与依赖方向", max_width=15.2, max_height=17.0)
    base.add_table(doc, ["运行单元", "主要职责", "状态/存储", "通信"], [
        ["浏览器 SPA", "页面、交互、路由、会话状态、图表", "Pinia/localStorage/临时结果缓存", "HTTPS REST"],
        ["FastAPI 进程", "认证、检测、管理、报告、调度", "进程缓存、限流桶、数据库会话", "HTTP/SQL/外部 API"],
        ["MySQL", "关系业务事实和事务", "8 张核心表", "SQLAlchemy/PyMySQL"],
        ["Chroma", "知识向量索引", "knowledge_items 集合", "本地持久客户端"],
        ["报告目录", "HTML/PDF 文件", "REPORT_DIR", "受控文件 I/O"],
        ["外部服务", "LLM、Embedding、搜索", "供应商侧", "超时受控 HTTPS"],
    ], [3.2, 5.7, 4.5, 3.7])
    base.add_heading(doc, "3.2 后端包组织", 2)
    base.add_table(doc, ["包", "职责", "禁止事项"], [
        ["app/api/v1", "路由、依赖、参数和 HTTP 映射", "不得实现复杂评分算法"],
        ["app/schemas", "请求/响应契约与字段校验", "不得直接访问数据库"],
        ["app/services", "业务规则、编排、外部服务适配", "不得绕过权限暴露对象"],
        ["app/crud", "查询、分页、持久化和事务配合", "不得调用前端或 HTTP"],
        ["app/models", "SQLAlchemy 表和关系", "不得承载请求上下文"],
        ["app/core", "配置、安全、依赖、限流、调度", "不得依赖具体页面"],
        ["app/utils", "纯函数和通用辅助", "不得隐藏业务副作用"],
    ], [3.2, 7.8, 6.1])
    base.add_heading(doc, "3.3 程序单元清单", 2)
    base.add_table(doc, ["标识", "程序单元", "实现位置", "关键存储/外部依赖"], [
        [p["id"], p["name"], p["location"], p["dependencies"]] for p in PROGRAMS
    ], [1.5, 3.7, 7.2, 4.7], font_size=8.2)

    base.add_heading(doc, "4 前端详细设计", 1)
    base.add_heading(doc, "4.1 页面与路由", 2)
    base.add_table(doc, ["路由", "页面组件", "访问控制", "主要职责"], FRONTEND_ROUTES, [3.5, 4.2, 3.2, 6.2], font_size=8.8)
    base.add_heading(doc, "4.2 状态、请求与结果缓存", 2)
    base.add_numbered(doc, [
        "user store 保存 user、restored、sessionVerified；初始化必须调用 /auth/me 验证 Token。",
        "Axios baseURL 默认 /api、超时 15 秒；请求拦截器附加 Authorization。",
        "业务 API 统一消费 response.data；文件下载通过 rawResponse 保留 Blob/响应头。",
        "游客检测结果只写 detectionResultCache，供 /result/:id 临时展示；登录历史以服务端为准。",
        "结果页把 evidence_quality、arbitration_status、candidate/excluded evidence 与最终有效证据区分呈现。",
        "图表数据通过 statisticsCharts 工具转换，不在组件中重复推导统计口径。",
    ])
    base.add_heading(doc, "4.3 关键页面交互状态", 2)
    base.add_table(doc, ["页面", "状态", "触发", "界面行为", "退出条件"], [
        ["检测", "idle", "进入页面", "显示手工输入/链接输入", "提交或提取"],
        ["检测", "extracting", "提交 URL", "禁用重复操作并显示加载", "预览成功/错误"],
        ["检测", "review", "提取成功", "回填可编辑标题、正文、来源和时间", "用户提交"],
        ["检测", "detecting", "提交检测", "展示 AgentSteps 与防重复提交", "成功/失败"],
        ["结果", "quality-warning", "仲裁或质量非 ok", "明确提示证据不足/重试耗尽", "重评或返回"],
        ["管理列表", "loading/empty/error/ready", "筛选/分页/刷新", "骨架、空态、错误或表格", "请求完成"],
    ], [2.4, 2.7, 3.5, 5.6, 3.0])
    base.add_heading(doc, "4.4 前端验证与可用性", 2)
    base.add_bullets(doc, ["表单先做必填、长度和格式提示，但后端仍重复校验。", "危险状态变更采用确认框；提交期间禁用按钮，成功后刷新服务端状态。", "风险不只依赖颜色，必须同时显示文字标签、分数和解释。", "键盘焦点、错误提示和表格空态保持可感知；移动端表格允许横向滚动。", "外部 URL 使用安全链接属性，页面不把模型返回内容作为未转义 HTML 注入。"])

    base.add_heading(doc, "5 数据结构与接口详细设计", 1)
    base.add_heading(doc, "5.1 核心业务对象", 2)
    base.add_table(doc, ["对象", "关键字段", "不变量"], [
        ["DetectNewsRequest", "title/content/category/source/publish_time/enable_web_search", "title≥4；content≥20；extra=forbid"],
        ["EvidenceCandidate", "candidate_id/source_type/title/url/similarity", "candidate_id 在本次候选池唯一"],
        ["EvidenceArbitration", "ranked_evidence/rejected_evidence", "候选全集恰好出现一次"],
        ["EvidenceQuality", "coverage/consistency/score/assessment", "前两项 0～100；score 服务端计算"],
        ["DetectionResult", "三分量/最终分/等级/理由/证据/状态", "分数 0～100；有效证据已仲裁"],
        ["AnalysisPayload", "候选、排除、相似新闻、质量、契约、降级", "允许旧记录缺字段，读取有默认值"],
    ], [3.4, 8.2, 5.2])
    base.add_heading(doc, "5.2 REST 接口程序映射", 2)
    base.add_table(doc, ["域", "方法", "路径", "处理函数", "权限", "输入", "输出"], ENDPOINTS, [2.0, 1.5, 4.4, 3.5, 3.2, 3.0, 3.8], font_size=7.6)
    base.add_heading(doc, "5.3 LLM 输出契约", 2)
    add_code_block(doc, '''{
  "llm_score": 0-100,
  "risk_level": "可信新闻|存疑信息|疑似谣言|高风险谣言",
  "reason": "非空说明",
  "evidence_quality": {
    "coverage": 0-100, "consistency": 0-100,
    "score": "由服务端重算", "assessment": "说明"
  },
  "evidence_arbitration": {
    "ranked_evidence": [{"candidate_id":"kb:1","relevance_score":90,
      "quality_score":85,"stance":"support","reason":"..."}],
    "rejected_evidence": [{"candidate_id":"web:2","reason":"..."}]
  },
  "similar_news": [], "risk_points": [], "keywords": [], "suggestion": "..."
}''', "模型分析 JSON 契约（语义示意）")
    base.add_heading(doc, "5.4 数据库会话与事务边界", 2)
    base.add_table(doc, ["用例", "事务边界", "提交点", "补偿/回滚"], [
        ["注册/用户变更", "单个关系库事务", "用户状态合法后", "异常 rollback；唯一约束兜底"],
        ["检测保存", "记录+有效证据同一事务", "全部快照写入后", "任一失败整体 rollback"],
        ["知识+向量", "MySQL 事务与 Chroma 操作分离", "事实先持久化，状态后更新", "failed/delete_failed；可单条重试/全量重建"],
        ["Prompt 默认切换", "同类型模板加锁", "旧默认清除且新默认设置后", "冲突 rollback"],
        ["报告", "文件生成与元数据分阶段", "文件成功后 upsert 元数据", "失败清理临时文件并 rollback"],
        ["审核", "单条检测状态事务", "状态机校验后", "并发冲突返回 409"],
    ], [3.0, 5.4, 4.3, 4.3])

    base.add_heading(doc, "6 程序单元详细设计", 1)
    base.add_body(doc, "本章按“程序标识—目的—输入—输出—依赖—处理逻辑—异常—测试—限制”的统一模板描述可独立开发和验证的程序单元。程序位置采用当前仓库相对路径；私有辅助函数属于其上级程序单元。")
    for index, program in enumerate(PROGRAMS, 1):
        add_program(doc, index, program)

    base.add_heading(doc, "7 新闻检测算法与流程", 1)
    base.add_heading(doc, "7.1 主流程", 2)
    base.add_figure(doc, "fig4-4-detection-overall-flow.png", "图 7-1 新闻检测主流程", max_width=15.2, max_height=18.0)
    add_code_block(doc, '''function detect(payload, user):
  title, content = clean(payload)
  keywords = extract_keywords(title, content, max=8)
  candidates = chroma_search(title + content, top_k=10)
  if web_enabled and should_trigger(candidates):
      candidates = merge(candidates, web_search(title, keywords))
  assign_candidate_ids(candidates)
  prompt_candidates = neutral_sha256_order(candidates)[:10]
  llm = analyze(prompt_candidates)
  arbitration = validate(llm)
  if invalid(arbitration): arbitration = validate(retry_once())
  effective = arbitration.ranked if valid else []
  rule = calculate_rule_score(title, content, source, effective)
  final = combine(llm, evidence_quality, rule, effective)
  save_detection_and_effective_evidence(final, audit_payload)
  return result_with_disclaimer''', "检测核心伪代码")
    base.add_heading(doc, "7.2 联网触发与证据合并", 2)
    base.add_figure(doc, "fig4-5-local-retrieval-web-supplement.png", "图 7-2 本地检索与联网补证流程", max_width=15.2, max_height=18.0)
    base.add_body(doc, "联网检索是补充而非默认事实源。只有配置开启、用户允许且本地结果不足时触发；联网异常不会抹去本地证据。合并时保留 source_type/source_label/source_url，按稳定键去重，并在候选分配后由模型仲裁，不能仅因来源为 Web 或检索排名靠前而自动成为有效证据。")
    base.add_heading(doc, "7.3 评分公式", 2)
    base.add_table(doc, ["情形", "最终分公式", "设计意图"], [
        ["LLM 失败/降级", "Final = Rule", "不信任未经模型仲裁的检索候选"],
        ["LLM 可用、无有效证据", "Final = 0.6×LLM + 0.4×Rule", "无证据质量项，保留规则制衡"],
        ["LLM 可用、有有效证据", "Final = 0.5×LLM + 0.3×EvidenceQuality + 0.2×Rule", "以模型分析为主，证据质量与确定性规则约束"],
        ["证据质量", "EQ = 0.6×Coverage + 0.4×Consistency", "服务端重算，防止模型伪造聚合分"],
    ], [4.3, 6.0, 7.0])
    base.add_heading(doc, "7.4 风险映射与高风险标记", 2)
    base.add_table(doc, ["分数区间", "风险等级", "高风险标记"], [
        ["80～100", "可信新闻", "否"], ["60～<80", "存疑信息", "否"],
        ["40～<60", "疑似谣言", "否"], ["0～<40", "高风险谣言", "是"],
        ["任意分数但模型等级明确为高风险谣言", "最终仍按 Final 映射；高风险工具可按等级触发", "是"],
    ], [4.4, 6.0, 6.9])
    base.add_note(doc, "分值语义", "final_score 是系统可信度辅助分，不是概率，也不应表达为‘真实性百分比’。", "warning")
    base.add_heading(doc, "7.5 仲裁状态机", 2)
    base.add_table(doc, ["状态", "条件", "评分证据", "用户呈现"], [
        ["ok", "首次或重试契约完全通过", "ranked_evidence", "显示有效/排除证据及质量"],
        ["no_evidence", "候选池为空", "空", "明确无可用证据"],
        ["provider_error", "LLM 服务失败", "空", "显示降级提示"],
        ["retry_exhausted", "两次契约均失败", "空", "显示仲裁失败并允许重新评估"],
        ["unavailable", "旧记录或未形成状态", "空/历史快照", "按兼容逻辑展示不可用"],
    ], [3.0, 5.6, 4.0, 4.7])

    base.add_heading(doc, "8 数据库与存储详细设计", 1)
    base.add_heading(doc, "8.1 概念关系", 2)
    base.add_figure(doc, "fig4-7-database-er.png", "图 8-1 核心数据库 E-R 关系", max_width=15.2, max_height=18.0)
    base.add_heading(doc, "8.2 物理表设计", 2)
    base.add_body(doc, "数据库采用 MySQL/InnoDB/utf8mb4。字段与索引以 Alembic 迁移为最终事实；下列表结构与当前迁移 0001～0007 对齐。所有业务更新时间统一由 updated_at 表示，审计事件另写 system_logs。")
    for index, spec in enumerate(base.DB_TABLES, 1):
        base.add_heading(doc, f"8.2.{index} {spec['title']}（{spec['name']}）", 3)
        base.add_body(doc, str(spec["purpose"]), indent=False)
        base.add_table(doc, ["字段", "类型", "可空", "约束/默认", "说明"], list(spec["fields"]), [3.0, 3.2, 1.7, 4.6, 5.6], font_size=7.9)
        base.add_note(doc, "索引与约束", str(spec["indexes"]), "info")
    base.add_heading(doc, "8.3 关系、删除和完整性", 2)
    base.add_table(doc, ["父对象", "子对象", "关系", "删除策略", "理由"], [
        ["users", "detection_records", "1:N", "SET NULL", "保留匿名化检测审计快照"],
        ["users", "reports", "1:N", "SET NULL", "保留报告元数据"],
        ["users", "prompt_templates", "1:N", "SET NULL", "保留模板历史"],
        ["detection_records", "evidence_matches", "1:N", "CASCADE", "证据快照无独立生命周期"],
        ["detection_records", "reports", "1:0..1", "CASCADE + UNIQUE", "一检测一个报告"],
        ["knowledge_items", "evidence_matches", "1:N", "SET NULL", "知识删除后保留检测时证据快照"],
    ], [3.2, 3.6, 2.1, 3.6, 5.6])
    base.add_heading(doc, "8.4 Chroma 集合", 2)
    base.add_table(doc, ["项目", "设计"], [
        ["集合", "knowledge_items"], ["距离空间", "cosine；similarity=max(0,1-distance)"],
        ["ID", "knowledge:<MySQL knowledge_items.id>"], ["document", "标题、正文、摘要与辟谣说明组成的清洗文本"],
        ["metadata", "knowledge_id、title、category、truth_label、source_name、risk_level"],
        ["top_k", "默认 10，服务层归一化为 1～50"], ["一致性", "MySQL 为事实源；pending/synced/failed/delete_failed；可全量重建"],
    ], [4.2, 12.9], first_col_bold=True)
    base.add_heading(doc, "8.5 报告文件和路径安全", 2)
    base.add_bullets(doc, ["REPORT_DIR 使用部署配置且位于源码目录外；进程账号只授予该目录所需权限。", "文件名由服务端基于报告/检测标识生成，不直接使用用户标题。", "下载前解析规范化绝对路径并验证其仍在 REPORT_DIR 下，拒绝 ../ 和符号链接逃逸。", "数据库只保存受控相对路径；备份需同时覆盖 reports 表和报告目录。", "渲染失败产生的临时文件应清理，不能登记为可下载报告。"])
    base.add_heading(doc, "8.6 迁移、备份与恢复", 2)
    base.add_numbered(doc, ["结构变更只新增 Alembic 版本，不修改已发布迁移。", "部署先备份 MySQL 和报告目录，再执行 python -m app.db.migrate。", "迁移后验证当前 revision=head、核心索引和字段存在。", "Chroma 不作为唯一备份；丢失或维度变更时从 knowledge_items 全量重建。", "恢复顺序：MySQL → 配置/秘密 → 报告目录 → Chroma 重建 → 冒烟测试。"])

    base.add_heading(doc, "9 安全详细设计", 1)
    base.add_heading(doc, "9.1 威胁与控制", 2)
    base.add_table(doc, ["威胁", "入口", "控制", "验证"], [
        ["凭据泄露", "仓库/日志/错误", "环境变量、bcrypt、日志脱敏、启动校验", "秘密扫描与错误响应检查"],
        ["越权访问", "对象 ID/后台路由", "实时 get_current_user/admin、用户范围查询", "跨用户/普通用户用例"],
        ["暴力尝试/资源滥用", "登录/注册/检测/抓取", "按 IP 限流、长度/超时/数量上限", "429 与边界测试"],
        ["SSRF", "链接提取", "协议白名单、DNS/IP 检查、重定向复验、响应限制", "私网/保留地址/重绑定测试"],
        ["Prompt 注入", "新闻/网页/证据文本", "数据与指令分隔、固定输出契约、服务端字段校验", "恶意文本与候选伪造测试"],
        ["路径穿越", "报告下载", "服务端文件名、resolve 后目录边界检查", "../、绝对路径、链接测试"],
        ["未审高风险公开", "公开 API", "approved + is_public + is_high_risk 三条件", "状态组合测试"],
        ["XSS", "模型/网页内容展示", "Vue 默认转义、禁止不可信 v-html、安全链接", "脚本载荷测试"],
    ], [3.1, 3.4, 7.0, 5.1], font_size=8.6)
    base.add_heading(doc, "9.2 认证授权矩阵", 2)
    base.add_table(doc, ["能力", "游客", "普通用户", "管理员"], [
        ["注册/登录/公开高风险", "允许", "允许", "允许"], ["新闻检测/链接预览", "允许（限流）", "允许", "允许"],
        ["本人历史/详情/报告", "不允许持久历史", "允许本人", "允许"], ["独立 RAG", "不允许", "允许", "允许"],
        ["用户/知识/Prompt/审核/统计/日志", "不允许", "不允许", "允许"],
    ], [5.8, 3.5, 4.1, 4.7])
    base.add_heading(doc, "9.3 内容安全与免责声明", 2)
    base.add_body(doc, "新闻正文、网页材料和模型返回均视为不可信数据。系统不执行其中的命令，不把模型结果用于权限或 SQL 构造，不把内部 Prompt、密钥、完整错误体提供给前端。每次检测结果和报告必须包含“仅供辅助评估、不能替代人工与权威核查”的免责声明。")

    base.add_heading(doc, "10 异常处理、恢复与可观测性", 1)
    base.add_heading(doc, "10.1 故障分级", 2)
    base.add_table(doc, ["级别", "示例", "系统动作", "用户结果"], [
        ["A 关键失败", "MySQL/Chroma 本地检索不可用", "终止当前写操作并回滚", "503/明确失败，不生成伪结论"],
        ["B 可降级外部失败", "联网搜索失败", "记录 warning，继续本地 RAG", "结果标注实际来源"],
        ["B 可降级模型失败", "LLM 超时/服务错误", "规则评分降级，保存 provider_error", "显示降级提示和免责声明"],
        ["C 单项补偿", "向量 upsert/delete 失败", "写 failed/delete_failed", "后台可重试/重建"],
        ["D 用户错误", "输入、权限、状态冲突", "拒绝业务，不重试", "4xx 可操作提示"],
    ], [2.6, 5.0, 6.3, 5.2])
    base.add_heading(doc, "10.2 日志设计", 2)
    base.add_table(doc, ["日志", "内容", "禁止内容", "用途"], [
        ["应用日志", "级别、模块、异常类别、耗时和关联 ID", "密码、Token、API Key、完整敏感正文", "开发/运维定位"],
        ["system_logs", "user_id、module、action、description、IP、时间", "秘密、完整认证头", "管理员业务审计"],
        ["crawl_tasks", "查询、数量、失败摘要、起止时间、状态", "外部服务秘密", "采集任务追踪"],
        ["analysis_payload", "候选/排除、质量、契约、降级状态", "外部密钥、系统 Prompt", "检测结果可解释与重放分析"],
    ], [3.0, 7.0, 5.0, 4.1])
    base.add_heading(doc, "10.3 恢复和重试", 2)
    base.add_bullets(doc, ["Chroma 客户端/集合异常清缓存后重试一次；持续失败转领域错误。", "LLM 输出契约错误聚焦重试一次；服务调用错误不进行无界重试。", "HTTP 外部调用必须设置连接/读取总超时；调用方决定降级或失败。", "数据库异常必须 rollback 后才能复用会话。", "重新评估创建新检测，不覆盖历史记录；恢复过程有审计链。"])

    base.add_heading(doc, "11 性能、容量与并发设计", 1)
    base.add_heading(doc, "11.1 性能控制点", 2)
    base.add_table(doc, ["位置", "控制", "目的"], [
        ["请求 Schema", "标题 255、正文服务层 12000、URL 2048、分页≤100", "限制 CPU/内存/Prompt"],
        ["RAG", "Top10，Chroma top_k 1～50", "限制召回与序列化"], ["模型", "最多 10 个中立排序候选；重试 1 次", "控制 Token、费用与时延"],
        ["联网", "按需触发、count/freshness/timeout 可配置", "减少外部依赖"], ["前端", "视图懒加载、请求 15 秒超时", "首屏和故障反馈"],
        ["数据库", "风险、用户、状态、时间等索引；分页查询", "避免全表扫描和大结果"],
    ], [3.3, 8.2, 6.6])
    base.add_heading(doc, "11.2 并发一致性", 2)
    base.add_bullets(doc, ["每请求独立 SQLAlchemy Session；提交/回滚边界由服务明确。", "Prompt 同类型默认切换采用锁和事务，保证至多一个启用默认。", "用户角色/状态变更在事务内再次检查最后管理员约束。", "当前 InMemoryRateLimiter 仅保证单进程；水平扩容时改用 Redis 等共享限流。", "定时采集多实例只能由一个调度节点执行，或使用分布式锁/独立 worker。"])
    base.add_heading(doc, "11.3 初步容量边界", 2)
    base.add_table(doc, ["对象", "设计边界", "扩展措施"], [
        ["检测记录", "按用户/时间分页；正文与 analysis_payload 为主要容量", "按月归档、冷热分层、只读审计库"],
        ["证据快照", "仅保存仲裁有效证据；候选池存 JSON", "压缩/对象存储或 JSON 原生列"],
        ["知识向量", "一知识条目一向量", "分片集合、批量 embedding、异步重建"],
        ["报告", "一检测至多一份", "对象存储、生命周期策略、CDN"],
        ["审计日志", "写多读少", "按时间分区、归档与集中日志平台"],
    ], [3.4, 7.2, 7.5])

    base.add_heading(doc, "12 单元测试与集成测试设计", 1)
    base.add_heading(doc, "12.1 测试层次", 2)
    base.add_table(doc, ["层次", "对象", "方法", "通过准则"], [
        ["单元", "纯函数、Schema、规则、仲裁、格式化", "确定输入输出与边界表", "分支与异常均有断言"],
        ["服务", "检测、知识同步、报告、用户管理", "替身外部 API + 临时数据库/目录", "事务与降级符合设计"],
        ["API", "所有路由、权限和状态码", "FastAPI TestClient", "响应契约、权限、分页正确"],
        ["数据", "Alembic、模型、约束、级联", "升级到 head 与回滚/重建验证", "结构和数据完整"],
        ["前端", "格式化、缓存、质量状态、路由守卫", "Vitest + 构建 + 浏览器冒烟", "状态与展示一致"],
        ["端到端", "注册→检测→历史→报告→审核", "真实 MySQL/Chroma，外部服务可控", "主业务闭环通过"],
    ], [2.7, 5.3, 5.8, 4.3])
    base.add_heading(doc, "12.2 关键测试用例", 2)
    base.add_table(doc, ["编号", "场景", "前置/输入", "预期"], [
        ["LLD-T01", "正常检测", "有效文本、本地证据、合法 LLM 契约", "保存检测与有效证据；公式/等级正确"],
        ["LLD-T02", "仲裁重试", "首次遗漏 candidate_id，第二次合法", "attempts=2，status=ok"],
        ["LLD-T03", "仲裁耗尽", "两次均非法", "有效证据为空；retry_exhausted 可见"],
        ["LLD-T04", "模型失败降级", "LLM 超时", "Final=Rule；provider_error；免责声明存在"],
        ["LLD-T05", "知识检索失败", "Chroma 持续异常", "503；不生成检测记录"],
        ["LLD-T06", "SSRF", "127.0.0.1/私网/保留地址 URL", "422；不发出危险请求"],
        ["LLD-T07", "跨用户详情", "用户 A 请求用户 B id", "404；不泄露记录存在性"],
        ["LLD-T08", "未审公开", "pending 高风险记录", "公开 API 永不返回"],
        ["LLD-T09", "向量失败补偿", "upsert 抛错", "MySQL 状态 failed 且可重试"],
        ["LLD-T10", "报告路径穿越", "篡改数据库路径到目录外", "拒绝下载"],
        ["LLD-T11", "最后管理员", "尝试禁用/降权唯一 active admin", "409；状态不变"],
        ["LLD-T12", "迁移", "空库执行 upgrade head", "8 表、索引、字段和 revision 正确"],
    ], [2.0, 4.0, 6.2, 6.0], font_size=8.4)
    base.add_heading(doc, "12.3 回归命令", 2)
    add_code_block(doc, '''cd backend
python -m unittest discover -s tests -p "test_*.py"
python -m app.db.migrate

cd ../frontend
npm run test -- --run
npm run build

cd ..
git diff --check''')

    base.add_heading(doc, "13 需求—程序—数据追踪", 1)
    base.add_table(doc, ["需求域", "程序单元", "主要接口", "主要数据", "验证"], [
        ["注册登录与权限", "P-02/P-03/P-15", "/auth/*、/admin/users/*", "users/system_logs", "认证、越权、最后管理员测试"],
        ["新闻输入与提取", "P-04/P-09", "/detect/news、extract-preview", "请求快照", "长度、SSRF、预览测试"],
        ["RAG 与补证", "P-05/P-07/P-08/P-09", "/rag/search、knowledge/*", "knowledge_items/Chroma", "召回、同步、重建、联网降级"],
        ["LLM 与评分", "P-05/P-06/P-10/P-11", "/detect/news", "detection_records.analysis_payload", "契约、重试、公式、降级"],
        ["历史与重评", "P-04/P-12", "/detect/history/{id}/re-evaluate", "detection_records/evidence_matches", "隔离、不可变重评"],
        ["报告", "P-13", "/report/*、/admin/reports/*", "reports/文件", "所有权、幂等、路径测试"],
        ["高风险复核", "P-14", "/high-risk/*、/admin/high-risk/*", "detection_records/system_logs", "状态机与公开隔离"],
        ["后台统计审计", "P-15", "/admin/statistics/*、/logs", "全部业务表/system_logs", "口径、空集、脱敏"],
    ], [3.3, 4.0, 5.5, 4.2, 4.0], font_size=8.4)
    base.add_heading(doc, "13.2 设计评审检查表", 2)
    base.add_table(doc, ["检查项", "结论", "说明"], [
        ["程序单元均有标识、目的、输入、输出、逻辑、异常、测试和限制", "满足", "见第 6 章"],
        ["数据库字段、索引、关系、向量与文件存储明确", "满足", "见第 8 章"],
        ["模型输出存在服务端契约与失败隔离", "满足", "见 5.3、7.5"],
        ["权限同时覆盖路由和对象范围", "满足", "见 9.2"],
        ["外部依赖均有超时/降级/失败语义", "满足", "见第 10 章"],
        ["每项关键设计可由测试追踪", "满足", "见第 12、13 章"],
        ["已知限制与扩展路径已记录", "满足", "见各程序限制和第 11 章"],
    ], [9.6, 2.4, 7.9])
    base.add_heading(doc, "13.3 结论", 2)
    base.add_body(doc, "本详细设计已把需求和概要设计落实到 15 个程序单元、31 组 REST 映射、17 个前端路由、8 张关系表、1 个向量集合和受控报告目录，明确了检测算法的证据边界、契约重试、分情形评分、人工复核、异常恢复和测试准则，可作为编码核对、测试设计和项目验收的技术基线。")

    base.add_heading(doc, "附录 A 状态与枚举字典", 1)
    base.add_table(doc, ["对象", "值", "含义/转换约束"], [
        ["users.role", "user/admin", "管理员权限由实时用户记录决定"], ["users.status", "active/disabled", "disabled 不得通过认证依赖"],
        ["risk_level", "可信新闻/存疑信息/疑似谣言/高风险谣言", "由 final_score 映射"],
        ["review_status", "pending/approved/rejected", "只有 approved 可公开"],
        ["vector_sync_status", "pending/synced/failed/delete_failed", "失败可重试或重建"],
        ["arbitration_status", "unavailable/no_evidence/ok/provider_error/retry_exhausted", "决定证据是否可用于评分"],
        ["quality_status", "unavailable/no_evidence/ok", "非 ok 时 evidence_quality 对外为空"],
        ["stance", "support/contradict/neutral", "仅用于已排名证据"],
        ["crawl_tasks.status", "running/success/partial/failed", "由采集统计与错误决定"],
    ], [4.2, 5.7, 9.9])
    base.add_heading(doc, "附录 B 配置项基线", 1)
    base.add_table(doc, ["配置", "用途", "安全/变更要求"], [
        ["DATABASE_URL 或 DATABASE_*", "MySQL 连接", "秘密不入库；最小权限账户"],
        ["SECRET_KEY/ALGORITHM/TOKEN_EXPIRE", "JWT", "强随机；轮换需考虑现有 Token"],
        ["DEEPSEEK_API_KEY/BASE_URL/MODEL", "模型分析", "密钥脱敏；超时受控"],
        ["EMBEDDING_PROVIDER/MODEL/DIMENSION", "向量生成", "provider/维度变化后全量重建"],
        ["DASHSCOPE_API_KEY", "正式语义 Embedding", "不得降级为仓库默认秘密"],
        ["CHROMA_PATH", "向量持久目录", "可重建；进程可读写"],
        ["REPORT_DIR", "报告目录", "源码外；路径边界与备份"],
        ["BOCHA_API_KEY/WEB_SEARCH_*", "联网补证", "可关闭；超时、数量和 freshness 有上限"],
        ["BACKEND_CORS_ORIGINS", "跨域白名单", "生产禁用不受控 *"],
        ["*_RATE_LIMIT_*", "注册/登录/检测限流", "多实例改用共享存储"],
        ["ARTICLE_FETCH_ALLOW_PRIVATE_HOSTS", "链接提取私网例外", "生产保持 false；仅受控测试可开"],
    ], [5.1, 6.0, 8.6], font_size=8.5)
    base.add_heading(doc, "附录 C 源码与文档对应索引", 1)
    base.add_table(doc, ["设计主题", "源码/资料位置"], [
        ["检测与仲裁", "backend/app/services/detection_service.py；llm_service.py；rule_score_service.py"],
        ["RAG/向量", "knowledge_service.py；chroma_service.py；embedding_service.py"],
        ["Web 搜索/提取", "backend/app/services/web/*"], ["认证与权限", "core/security.py；core/deps.py；services/auth_service.py"],
        ["数据库", "backend/app/models/*；backend/alembic/versions/*；docs/02_database_design.md"],
        ["接口", "backend/app/api/v1/*；backend/app/schemas/*；docs/03_api_design.md"],
        ["前端", "frontend/src/router、stores、api、views、components、utils"],
        ["测试", "backend/tests/*；frontend/src/utils/*.test.js"],
    ], [5.0, 14.7])


def set_document_metadata(doc: Document) -> None:
    props = doc.core_properties
    props.title = f"《{PROJECT_TITLE}》{DOC_TITLE}"
    props.subject = "软件工程详细设计、程序单元、算法、数据库、接口与测试设计"
    props.author = "项目开发组"
    props.keywords = "详细设计, GB/T 8567, RAG, 大语言模型, 新闻真伪鉴别, FastAPI, Vue3, MySQL, Chroma"
    props.comments = f"文档编号 {DOC_NO}；版本 {VERSION}；基线 {BASELINE}"
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    base.configure_styles(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_main_content(doc)
    set_document_metadata(doc)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
