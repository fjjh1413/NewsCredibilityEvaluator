from __future__ import annotations

import json
import os
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
MARKDOWN = ROOT / "docs" / "thesis" / "chapter4_system_design.md"
FIGURE_DIR = ROOT / "docs" / "thesis" / "figures" / "chapter4"
FIGURE_MAP = ROOT / "docs" / "thesis" / "chapter4_figure_map.json"
OUTPUT_DIR = ROOT / ".artifacts" / "docs"


def find_source() -> Path:
    explicit_source = os.environ.get("CHAPTER4_SOURCE_DOCX")
    if explicit_source:
        source = Path(explicit_source).expanduser().resolve()
        if not source.is_file() or not zipfile.is_zipfile(source):
            raise FileNotFoundError("CHAPTER4_SOURCE_DOCX must point to a valid DOCX file")
        return source
    candidates = sorted(
        (
            p
            for p in (OUTPUT_DIR / "input").glob("*第3章规范修订版.docx")
            if not p.name.startswith("~$") and zipfile.is_zipfile(p)
        ),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError("未找到第3章规范修订版论文；请放入 .artifacts/docs/input 或设置 CHAPTER4_SOURCE_DOCX")
    return candidates[0]


def output_path(source: Path) -> Path:
    base = source.stem.replace("_第3章规范修订版", "")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / f"{base}_第4章系统总体设计完成版.docx"


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")


def set_border(container, edge: str, val: str, size: str = "8") -> None:
    tag = "w:tcBorders" if container.tag.endswith("}tcPr") else "w:tblBorders"
    borders = container.find(qn(tag))
    if borders is None:
        borders = OxmlElement(tag)
        container.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), val)
    if val != "nil":
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def format_three_line_table(table) -> None:
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    set_border(tbl_pr, "top", "single", "12")
    set_border(tbl_pr, "bottom", "single", "12")
    for edge in ("left", "right", "insideH", "insideV"):
        set_border(tbl_pr, edge, "nil")

    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                set_border(tc_pr, "bottom", "single", "8")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, size=9.5, bold=(row_index == 0))


def move_before_reference(element, reference_paragraph) -> None:
    reference_paragraph._p.addprevious(element)


def add_paragraph_before(doc, reference, text="", style=None, alignment=None):
    paragraph = doc.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    if alignment is not None:
        paragraph.alignment = alignment
    move_before_reference(paragraph._p, reference)
    return paragraph


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells):
            continue
        rows.append(cells)
    return rows


def build() -> tuple[Path, list[dict[str, object]]]:
    source = find_source()
    output = output_path(source)
    doc = Document(source)
    references = next((p for p in doc.paragraphs if p.text.strip() == "参考文献"), None)
    if references is None:
        raise RuntimeError("未找到参考文献标题，无法确定插入位置")

    body_style = "Normal (Web)" if "Normal (Web)" in doc.styles else "Normal"

    page_break = add_paragraph_before(doc, references, style=body_style)
    page_break.add_run().add_break(WD_BREAK.PAGE)

    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    figures: list[dict[str, object]] = []
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()
        if not line:
            i += 1
            continue

        if line.startswith("### "):
            add_paragraph_before(doc, references, line[4:], "Heading 3")
        elif line.startswith("## "):
            add_paragraph_before(doc, references, line[3:], "Heading 2")
        elif line.startswith("# "):
            add_paragraph_before(doc, references, line[2:], "Heading 1")
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", line)
            if not match:
                raise ValueError(f"无效图片标记：{line}")
            caption, rel_path = match.groups()
            number_match = re.search(r"图4-(\d+)", caption)
            if not number_match:
                raise ValueError(f"图片标题缺少图号：{caption}")
            number = int(number_match.group(1))
            placeholder = f"[[CH4_FIGURE_{number}]]"
            image_paragraph = add_paragraph_before(
                doc,
                references,
                placeholder,
                body_style,
                WD_ALIGN_PARAGRAPH.CENTER,
            )
            image_paragraph.paragraph_format.keep_with_next = True
            caption_paragraph = add_paragraph_before(
                doc,
                references,
                caption,
                body_style,
                WD_ALIGN_PARAGRAPH.CENTER,
            )
            caption_paragraph.paragraph_format.keep_with_next = True
            svg_path = (MARKDOWN.parent / rel_path).resolve()
            emf_path = svg_path.with_suffix(".emf")
            png_path = svg_path.with_suffix(".png")
            from PIL import Image

            with Image.open(png_path) as image:
                ratio = image.width / image.height
            figures.append(
                {
                    "placeholder": placeholder,
                    "path": emf_path.relative_to(ROOT).as_posix(),
                    "ratio": ratio,
                    "number": number,
                }
            )
        elif line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            i -= 1
            rows = parse_table(table_lines)
            if not rows:
                raise ValueError("检测到空Markdown表格")
            columns = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=columns)
            for row_index, row in enumerate(rows):
                if len(row) != columns:
                    raise ValueError("Markdown表格列数不一致")
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
            format_three_line_table(table)
            move_before_reference(table._tbl, references)
        else:
            is_caption = bool(re.fullmatch(r"表4-\d+\s+.+", line))
            is_equation = bool(re.search(r"（4-\d+）$", line))
            alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if is_caption or is_equation
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            paragraph = add_paragraph_before(
                doc,
                references,
                line,
                body_style,
                alignment,
            )
            if is_caption:
                paragraph.paragraph_format.keep_with_next = True
            if not is_caption and not is_equation:
                paragraph.paragraph_format.first_line_indent = Pt(24)
        i += 1

    final_break = add_paragraph_before(doc, references, style=body_style)
    final_break.add_run().add_break(WD_BREAK.PAGE)

    doc.core_properties.title = "智闻辨真：基于RAG与大语言模型的新闻可信度评估系统"
    doc.save(output)
    FIGURE_MAP.write_text(
        json.dumps({"docx": output.relative_to(ROOT).as_posix(), "figures": figures}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return output, figures


if __name__ == "__main__":
    output, figures = build()
    print(output)
    print(f"figure_placeholders={len(figures)}")
